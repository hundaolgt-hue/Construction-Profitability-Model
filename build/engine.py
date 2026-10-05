"""Engine sheets 20–30."""
from openpyxl.styles import Alignment
from core import (Sheet, pc, CL, NF, ANN_N, MON_C, MON_L, font, fill, F_BODY, F_BOLD, F_UNIT, F_NOTE, F_AM,
                  NAVY, SEC, INK, BLUE, TEAL, VIOLET, AMBER, BOX, FILL_KEY, LINK_FG, BOTTOM_AMBER, PERIOD_LABELS)
from inputs_def import USES, USES_AM, FLOORS, ELEMENTS, ELEMENTS_AM, EL_KEYS

TAB_ENGINE = "1F5FAD"
TAB_FIN = "2E7D4F"


def C4(i):
    return i <= 4


def Y(i):
    return i - 4 if i > 4 else 0


def build_engine(b):
    X = lambda sh, key, i: b.sheets[sh].c(key, i, sheet=True)
    XS = lambda sh, key: b.sheets[sh].s(key)
    XT = lambda sh, key: b.sheets[sh].tot(key)
    XR = lambda sh, key, a=1, z=None: b.sheets[sh].rng(key, a, z)

    # =================================================================== 20 TIMELINE
    tl = Sheet(b, "20_Timeline", "20 · Timeline & flags", "የጊዜ ሰሌዳና ጠቋሚዎች", TAB_ENGINE, kind="annual")
    b.TL = tl.rows
    T = lambda key, i: f"'20_Timeline'!{pc(i)}{tl.rows[key]}"
    tl.sec("Annual periods: C1–C4 construction years, O1–O10 operating years", "ዓመታዊ ጊዜያት")
    tl.line("T_P", "Period number", lambda i: i, fmt="int", unit="no.", tot=None)
    tl.line("T_LABEL", "Period label", lambda i: "§" + PERIOD_LABELS[i - 1], fmt="gen", unit="", tot=None)
    tl.line("T_C", "Construction-year index", lambda i: i if i <= 4 else 0, fmt="int", unit="no.", tot=None)
    tl.line("T_Y", "Operating-year index", lambda i: Y(i), fmt="int", unit="no.", tot=None)
    tl.line("T_MONTHS", "Months in period", lambda i: f"MAX(0,MIN(12,CONSTR_EFF-12*({tl.c('T_C', i)}-1)))" if C4(i) else 12,
            fmt="int", unit="months")
    tl.line("T_ACTIVE", "Period active flag", lambda i: f"IF({tl.c('T_MONTHS', i)}>0,1,0)", fmt="int", unit="flag", tot=None)
    tl.line("T_CF", "Construction flag", lambda i: f"IF(AND({tl.c('T_C', i)}>0,{tl.c('T_MONTHS', i)}>0),1,0)" if C4(i) else 0,
            fmt="int", unit="flag", tot="sum")
    tl.line("T_OPF", "Operating flag", lambda i: f"IF(AND({tl.c('T_Y', i)}>=1,{tl.c('T_Y', i)}<=OPS_YEARS),1,0)",
            fmt="int", unit="flag", tot="sum")
    tl.line("T_COD", "Completion (COD) period flag",
            lambda i: f"IF(AND({tl.c('T_MONTHS', i)}>0,MIN(12*{tl.c('T_C', i)},CONSTR_EFF)=CONSTR_EFF),1,0)" if C4(i) else 0,
            fmt="int", unit="flag", tot="sum")
    tl.line("T_STARTM", "First month index", lambda i: f"12*({tl.c('T_C', i)}-1)+1" if C4(i) else f"CONSTR_EFF+12*({tl.c('T_Y', i)}-1)+1",
            fmt="int", unit="no.", tot=None)
    tl.line("T_ENDM", "Last month index", lambda i: f"MIN(12*{tl.c('T_C', i)},CONSTR_EFF)" if C4(i) else f"CONSTR_EFF+12*{tl.c('T_Y', i)}",
            fmt="int", unit="no.", tot=None)
    tl.line("T_END", "Period end date", lambda i: f"EDATE(START_DATE,{tl.c('T_ENDM', i)})-1", fmt="date", unit="date", tot=None)
    tl.line("T_YF", "Years from start (end of period)", lambda i: f"({tl.c('T_END', i)}+1-START_DATE)/365", fmt="n2",
            unit="years", tot=None)
    tl.line("T_LEN", "Period length", lambda i: f"{tl.c('T_MONTHS', i)}/12", fmt="n2", unit="years")
    tl.line("T_GC", "Gregorian year (period end)", lambda i: f'"GC "&YEAR({tl.c("T_END", i)})', fmt="gen", unit="", tot=None)
    tl.line("T_EC", "Ethiopian calendar year (period end)",
            lambda i: f'"EC "&(YEAR({tl.c("T_END", i)})-8+IF({tl.c("T_END", i)}>=DATE(YEAR({tl.c("T_END", i)}),9,11),1,0))',
            fmt="gen", unit="", tot=None)
    tl.line("T_DF", "Discount factor @ project rate", lambda i: f"(1+DISC_EFF)^(-{tl.c('T_YF', i)})", fmt="n2", unit="x", tot=None)
    tl.line("T_DFE", "Discount factor @ cost of equity", lambda i: f"(1+COE_EFF)^(-{tl.c('T_YF', i)})", fmt="n2", unit="x", tot=None)
    tl.gap()
    tl.sec("Key timing scalars", "ቁልፍ የጊዜ መረጃዎች")
    tl.scalar("OPS", "Operating horizon (fixed design horizon)", 10, fmt="int", unit="years", name="OPS_YEARS",
              note="Statements always show C1–C4 + O1–O10")
    tl.scalar("CY", "Number of construction periods used", "ROUNDUP(CONSTR_EFF/12,0)", fmt="int", unit="no.", name="CONSTR_Y")
    tl.scalar("CODD", "Completion (COD) date", "EDATE(START_DATE,CONSTR_EFF)-1", fmt="date", unit="date", name="COD_DATE")
    tl.scalar("ENDD", "End of operating horizon (O10)", "EDATE(START_DATE,CONSTR_EFF+12*OPS_YEARS)-1", fmt="date", unit="date",
              name="END_DATE")
    b.TL_PRINT_END = tl.r
    tl.gap()
    tl.sec("Monthly timeline (construction up to 48 months + 240 months of loan life)", "ወርሃዊ የጊዜ ሰሌዳ")
    M = lambda key, i: f"'20_Timeline'!{pc(i)}{tl.rows[key]}"
    tl.line("M_T", "Month index t", lambda i: i, fmt="int", unit="no.", tot=None, n=MON_L)
    tl.line("M_DATE", "Month start date", lambda i: f"EDATE(START_DATE,{tl.c('M_T', i)}-1)", fmt="mon", unit="date", tot=None,
            n=MON_L)
    tl.line("M_CF", "Construction month flag", lambda i: f"IF({tl.c('M_T', i)}<=CONSTR_EFF,1,0)", fmt="int", unit="flag",
            n=MON_L)
    tl.line("M_K", "Months after completion k", lambda i: f"{tl.c('M_T', i)}-CONSTR_EFF", fmt="int", unit="no.", tot=None, n=MON_L)
    tl.line("M_OY", "Operating year (within O1–O10)",
            lambda i: f"IF(AND({tl.c('M_K', i)}>=1,{tl.c('M_K', i)}<=12*OPS_YEARS),ROUNDUP({tl.c('M_K', i)}/12,0),0)",
            fmt="int", unit="no.", tot=None, n=MON_L)
    tl.line("M_LY", "Loan year after completion", lambda i: f"IF({tl.c('M_K', i)}>=1,ROUNDUP({tl.c('M_K', i)}/12,0),0)",
            fmt="int", unit="no.", tot=None, n=MON_L)
    tl.line("M_PIDX", "Annual period index (1–14, 0 = beyond horizon)",
            lambda i: f"IF({tl.c('M_T', i)}<=CONSTR_EFF,ROUNDUP({tl.c('M_T', i)}/12,0),IF({tl.c('M_OY', i)}>0,4+{tl.c('M_OY', i)},0))",
            fmt="int", unit="no.", tot=None, n=MON_L)

    # =================================================================== 21 AREA
    ar = Sheet(b, "21_Area", "21 · Area programme", "የቦታና የወለል ስፋት ፕሮግራም", TAB_ENGINE, kind="free", n=16, width_label=44)
    ar.sec("Floors (from Inputs block 3)", "ፎቆች")
    ar.line("FL_NAME", "Floor", lambda i: "§" + FLOORS[i - 1], fmt="gen", unit="", tot=None)
    ar.line("FL_GROSS", "Gross floor area", lambda i: f"FL_A{i}", fmt="n0", unit="m²")
    ar.line("FL_USE", "Use", lambda i: f"FL_U{i}", fmt="gen", unit="", tot=None)
    ar.line("FL_BASE", "Basement flag", lambda i: 1 if i <= 2 else 0, fmt="int", unit="flag", tot=None)
    ar.line("FL_HAS", "Floor built flag", lambda i: f"IF({ar.c('FL_GROSS', i)}>0,1,0)", fmt="int", unit="flag")
    ar.gap()
    ar.sec("Building summary", "የሕንፃ ማጠቃለያ")
    ar.scalar("GFA", "Gross floor area (GFA)", f"SUM(E{ar.rows['FL_GROSS']}:T{ar.rows['FL_GROSS']})",
              unit="m²", name="GFA", style="key")
    ar.scalar("GFA_B", "Basement gross area", f"SUMPRODUCT(E{ar.rows['FL_GROSS']}:T{ar.rows['FL_GROSS']},E{ar.rows['FL_BASE']}:T{ar.rows['FL_BASE']})",
              unit="m²", name="GFA_BASE")
    ar.scalar("GFA_A", "Above-ground gross area", "GFA-GFA_BASE", unit="m²", name="GFA_ABOVE")
    ar.scalar("N_ABOVE", "Above-ground floors built", f"SUM(G{ar.rows['FL_HAS']}:T{ar.rows['FL_HAS']})", fmt="int", unit="no.",
              name="N_ABOVE")
    ar.scalar("N_BASE", "Basement levels built", f"SUM(E{ar.rows['FL_HAS']}:F{ar.rows['FL_HAS']})", fmt="int", unit="no.", name="N_BASE")
    ar.scalar("BTYPE", "Building configuration", '=IF(N_BASE>0,"B"&N_BASE&"+","")&"G+"&MAX(0,N_ABOVE-1)'[1:], fmt="gen", unit="",
              name="BTYPE_CALC")
    ar.scalar("FOOT", "Ground-floor footprint", f"G{ar.rows['FL_GROSS']}", unit="m²", name="FOOTPRINT")
    ar.scalar("FARA", "Actual FAR (above-ground GFA ÷ plot)", "IF(PLOT_AREA>0,GFA_ABOVE/PLOT_AREA,0)", fmt="n2", unit="x",
              name="FAR_ACT")
    ar.scalar("BTOK", "Building type matches floors entered (1 = yes)",
              "IF(AND(IF(BLDG_TYPE=5,1,IF(N_ABOVE=BLDG_TYPE,1,0))=1,IF(BASEMENT_YN=1,N_BASE>0,N_BASE=0)),1,0)",
              fmt="int", unit="flag", name="BTYPE_OK")
    ar.gap()
    ar.sec("By use – areas and disposition (IAS 2 sold · IAS 40 let · IAS 16 owner-occupied)", "በአገልግሎት")
    U5 = lambda key, i: ar.c(key, i)
    ur = {}
    ar.line("U_NAME", "Use", lambda i: "§" + USES[i - 1] if i <= 5 else None, fmt="gen", unit="", tot=None, n=5)
    ar.line("U_GROSS", "Gross area", lambda i: f"SUMIF($E${ar.rows['FL_USE']}:$T${ar.rows['FL_USE']},{U5('U_NAME', i)},$E${ar.rows['FL_GROSS']}:$T${ar.rows['FL_GROSS']})",
            unit="m²", n=5)
    code = lambda i: USES[i - 1][:3].upper()
    ar.line("U_EFF", "Net-to-gross efficiency", lambda i: f"EFF_{code(i)}", fmt="pct", unit="%", tot=None, n=5)
    ar.line("U_NET", "Net saleable / lettable area", lambda i: f"{U5('U_GROSS', i)}*{U5('U_EFF', i)}", unit="m²", n=5)
    ar.line("U_SOLDF", "Share sold (IAS 2)", lambda i: f"CHOOSE(PROJECT_MODE+1,0,1-OWNUSE_PCT,0,SOLD_{code(i)})", fmt="pct",
            unit="%", tot=None, n=5)
    ar.line("U_OWNF", "Share owner-occupied (IAS 16)",
            lambda i: f"CHOOSE(PROJECT_MODE+1,1-HOME_RENT_SH,OWNUSE_PCT,OWNUSE_PCT,OWNUSE_PCT*(1-{U5('U_SOLDF', i)}))",
            fmt="pct", unit="%", tot=None, n=5)
    ar.line("U_RENTF", "Share let (IAS 40)", lambda i: f"1-{U5('U_SOLDF', i)}-{U5('U_OWNF', i)}", fmt="pct", unit="%", tot=None, n=5)
    ar.line("U_SALE", "Saleable area", lambda i: f"{U5('U_NET', i)}*{U5('U_SOLDF', i)}", unit="m²", n=5)
    ar.line("U_LET", "Lettable area", lambda i: f"{U5('U_NET', i)}*{U5('U_RENTF', i)}", unit="m²", n=5)
    ar.line("U_OWN", "Owner-occupied area", lambda i: f"{U5('U_NET', i)}*{U5('U_OWNF', i)}", unit="m²", n=5)
    ar.line("U_PRICE", "Sale price (today)", lambda i: f"PRICE_{code(i)}", unit="ETB/m²", tot=None, n=5)
    ar.line("U_RENT", "Rent (today)", lambda i: f"RENT_{code(i)}", unit="ETB/m²/mo", tot=None, n=5)
    ar.line("U_SALEVAL", "Sales value (today's prices)", lambda i: f"{U5('U_SALE', i)}*{U5('U_PRICE', i)}", unit="ETB", n=5)
    ar.line("U_RENTVAL", "Potential rent p.a. (today's prices)", lambda i: f"{U5('U_LET', i)}*{U5('U_RENT', i)}*12", unit="ETB",
            n=5)
    ar.gap()
    ar.sec("Area totals and allocation shares", "ጠቅላላ ስፋትና የወጪ ድልድል")
    ar.scalar("NET", "Net area (all uses)", f"$D${ar.rows['U_NET']}", unit="m²", name="NET_AREA")
    ar.scalar("HRS", "Mode 0: share of house rented out", "IF(AND(RENT_HOME_F=1,NET_AREA>0),MIN(1,RENT_HOME_AREA/NET_AREA),0)",
              fmt="pct", unit="%", name="HOME_RENT_SH")
    ar.scalar("SALE", "Saleable area", f"$D${ar.rows['U_SALE']}", unit="m²", name="SALE_AREA", style="key")
    ar.scalar("LET", "Lettable area", f"$D${ar.rows['U_LET']}", unit="m²", name="LET_AREA", style="key")
    ar.scalar("OWN", "Owner-occupied area", f"$D${ar.rows['U_OWN']}", unit="m²", name="OWN_AREA")
    ar.scalar("INV_SH", "Cost share → inventory (IAS 2)", "IF(NET_AREA>0,SALE_AREA/NET_AREA,0)", fmt="pct", unit="%", name="INV_SH")
    ar.scalar("IP_SH", "Cost share → investment property (IAS 40)", "IF(NET_AREA>0,LET_AREA/NET_AREA,0)", fmt="pct", unit="%",
              name="IP_SH")
    ar.scalar("PPE_SH", "Cost share → owner-occupied PPE (IAS 16)", "IF(NET_AREA>0,OWN_AREA/NET_AREA,1)", fmt="pct", unit="%",
              name="PPE_SH")
    ar.scalar("RET_SH", "Cost share retained (IAS 40 + IAS 16)", "IP_SH+PPE_SH", fmt="pct", unit="%", name="RET_SH")
    ar.scalar("SALEVAL", "Sales value at today's prices", f"$D${ar.rows['U_SALEVAL']}", unit="ETB", name="SALEVAL_TODAY")
    ar.scalar("BPRICE", "Blended sale price (today, after scenario & flex)", "IF(SALE_AREA>0,SALEVAL_TODAY/SALE_AREA,0)*D_PRICE_F",
              unit="ETB/m²", name="BLEND_PRICE")
    ar.scalar("GPR0", "Potential rent p.a. (today's prices, before scenario)",
              f"IF(PROJECT_MODE=0,IF(RENT_HOME_F=1,RENT_HOME_MONTHLY*12,0),$D${ar.rows['U_RENTVAL']})", unit="ETB", name="GPR_TODAY")
    ar.scalar("ACHK", "Area reconciliation difference (uses vs floors)", f"$D${ar.rows['U_GROSS']}-GFA", unit="m²", name="AREA_DIFF")
    ar.scalar("SCHK", "Share reconciliation (should be 100%)", "INV_SH+IP_SH+PPE_SH", fmt="pct", unit="%", name="SHARE_SUM")

    # =================================================================== 22 CAPEX (budget)
    cx = Sheet(b, "22_Capex", "22 · Capex – elemental BOQ, soft costs, land", "የካፒታል ወጪ – የግንባታ ዝርዝር", TAB_ENGINE,
               kind="free", n=0, width_label=44)
    cx.sec("Cost drivers", "የወጪ መወሰኛዎች")
    cx.scalar("FIN", "Finish level (1 Basic · 2 Standard · 3 Premium)", "FINISH", fmt="int", unit="code", name="FIN_IX")
    cx.scalar("LOCF", "Location factor", "IF(LOC_TYPE=1,LOCF_ADDIS,LOCF_REG)", fmt="n2", unit="x", name="LOCF")
    cx.scalar("STF", "Storey factor (superstructure)", "1+STOREY_UPLIFT*MAX(0,N_ABOVE-4)", fmt="n2", unit="x", name="STOREY_F")
    cx.scalar("CF", "Construction cost factor (scenario × flex)", "D_COST_F", fmt="n2", unit="x")
    cx.gap()
    cx.sec("Elemental cost plan (budget-date prices)", "የግንባታ ወጪ ዝርዝር (በበጀት ቀን ዋጋ)")
    hdr = cx.r
    cx.r += 1
    el0 = cx.r
    cx.r += 9
    b.EL_ROW = {k: el0 + j for j, k in enumerate(EL_KEYS)}

    def _el(r0):
        ws = cx.ws
        heads = ["", "Element", "Unit", "Budget (ETB)", "Library rate", "Location", "Storey", "Area m²", "Rate × area",
                 "Method", "Lump sum", "Pre-flex"]
        for j, h in enumerate(heads):
            if h:
                c = ws.cell(hdr, j + 1, h)
                c.font = font(9, True, NAVY)
                c.alignment = Alignment(horizontal="center" if j > 2 else "left", wrap_text=True)
        for j, k in enumerate(EL_KEYS):
            r = el0 + j
            ws.cell(r, 2, f"{ELEMENTS[j]}  ·  {ELEMENTS_AM[j]}").font = F_BODY
            ws.cell(r, 3, "ETB").font = F_UNIT
            fs = {5: f"INDEX(RATE_LIB,{j + 1},FIN_IX)", 6: "LOCF", 7: "STOREY_F" if k == "SUPER" else 1, 8: "GFA",
                  9: f"E{r}*F{r}*G{r}*H{r}", 10: f"METH_{k}", 11: f"LUMP_{k}", 12: f"IF(J{r}=2,K{r},I{r})", 4: f"L{r}*D_COST_F"}
            for col, v in fs.items():
                cell = ws.cell(r, col, ("=" + v) if isinstance(v, str) else v)
                cell.number_format = NF["n2"] if col in (6, 7) else (NF["int"] if col == 10 else NF["n0"])
                b.style_cell(cell, v, "key" if col == 4 else None)
            b.name(f"EL_{k}", f"'22_Capex'!$D${r}")
        for col in range(5, 13):
            cx.ws.column_dimensions[CL(col)].width = 12.5
    cx.raw(_el)
    cx.scalar("BASEX", "Basement extra (excavation, retaining, waterproofing)", "GFA_BASE*BASEMENT_RATE*LOCF*D_COST_F", name="BASE_X")
    cx.scalar("COMP", "Compound wall", "COMPOUND_YN*COMPOUND_LEN*COMPOUND_RATE*D_COST_F", name="X_COMPOUND")
    cx.scalar("SEPT", "Septic tank(s)", "SEPTIC_YN*SEPTIC_QTY*SEPTIC_RATE*D_COST_F", name="X_SEPTIC")
    cx.scalar("WTNK", "Water tank(s)", "WTANK_YN*WTANK_QTY*WTANK_RATE*D_COST_F", name="X_WTANK")
    cx.scalar("PAVE", "Paving / landscaping", "EXT_M2*EXT_RATE*D_COST_F", name="X_PAVE")
    cx.scalar("HB", "HARD COST AT BUDGET-DATE PRICES", f"SUM(D{el0}:D{el0 + 8})+BASE_X+X_COMPOUND+X_SEPTIC+X_WTANK+X_PAVE",
              name="HARD_BASE", style="key")
    cx.gap()
    C23 = "23_Construction"
    cx.sec("Total development cost (TDC)", "ጠቅላላ የልማት ወጪ")
    cx.scalar("HE", "Hard cost escalated (S-curve × escalation)", lambda: XT(C23, "HARD"), name="HARD_ESC_TOT")
    cx.scalar("ESCA", "  of which escalation", "HARD_ESC_TOT-HARD_BASE", name="ESC_AMT")
    cx.scalar("CONT", "Contingency", lambda: XT(C23, "CONT"), name="CONT_TOT")
    cx.scalar("DES", "Design fees", lambda: XT(C23, "DESIGN"), name="DESIGN_TOT")
    cx.scalar("SUP", "Supervision", lambda: XT(C23, "SUPERV"), name="SUPERV_TOT")
    cx.scalar("PER", "Permits & fees", lambda: XT(C23, "PERMIT"), name="PERMIT_TOT")
    cx.scalar("FFE", "FF&E / furniture", lambda: XT(C23, "FFE"), name="FFE_TOT")
    cx.scalar("CC", "CONSTRUCTION COST (excl. land & finance)",
              "HARD_ESC_TOT+CONT_TOT+DESIGN_TOT+SUPERV_TOT+PERMIT_TOT+FFE_TOT", name="CONSTR_COST", style="key")
    cx.scalar("LAND", "Land", lambda: XT(C23, "LANDACC"), name="LAND_TOT")
    cx.scalar("ARR", "Arrangement fee (IAS 23)", lambda: XT(C23, "ARRF"), name="ARR_TOT")
    cx.scalar("COM", "Commitment fees (IAS 23)", lambda: XT(C23, "COMF"), name="COMMIT_TOT")
    cx.scalar("IDC", "Interest during construction – capitalised (IAS 23)", lambda: XT(C23, "IDC"), name="IDC_TOT")
    cx.scalar("FINC", "Finance costs capitalised", "ARR_TOT+COMMIT_TOT+IDC_TOT", name="FIN_CAP")
    cx.scalar("TDC", "TOTAL DEVELOPMENT COST (TDC)", "CONSTR_COST+LAND_TOT+FIN_CAP", name="TDC", style="key")
    cx.scalar("TDCC", "Check: TDC vs monthly capex accrual", lambda: f"TDC-{XT(C23, 'ACC')}", name="TDC_DIFF")
    cx.scalar("PGFA", "TDC per gross m²", "IF(GFA>0,TDC/GFA,0)", unit="ETB/m²", name="TDC_M2G")
    cx.scalar("PNET", "TDC per net m²", "IF(NET_AREA>0,TDC/NET_AREA,0)", unit="ETB/m²", name="TDC_M2N")
    cx.scalar("CCM2", "Construction cost per gross m²", "IF(GFA>0,CONSTR_COST/GFA,0)", unit="ETB/m²", name="CC_M2")
    cx.scalar("USD", "Memo: TDC in USD", "IF(FX_USD>0,TDC/FX_USD,0)", unit="USD")
    cx.gap()
    cx.sec("Allocation of TDC (IAS 2 / IAS 40 / IAS 16)", "የወጪ ድልድል")
    cx.scalar("AINV", "Inventory – units for sale (IAS 2)", "TDC*INV_SH", name="TDC_INV")
    cx.scalar("AIP", "Investment property – let (IAS 40, cost model)", "TDC*IP_SH", name="TDC_IP")
    cx.scalar("APPE", "Owner-occupied (IAS 16)", "TDC*PPE_SH", name="TDC_PPE")
    cx.scalar("ACHK", "Allocation difference (should be 0)", "TDC-TDC_INV-TDC_IP-TDC_PPE", name="ALLOC_DIFF")
    cx.gap()
    cx.sec("Depreciation components of retained assets", "የእርጅና ቅናሽ ክፍሎች")
    cx.scalar("LFR", "Land fraction of TDC", "IF(TDC>0,LAND_TOT/TDC,0)", fmt="pct", unit="%", name="LAND_FR")
    cx.scalar("FFR", "FF&E fraction of TDC", "IF(TDC>0,FFE_TOT/TDC,0)", fmt="pct", unit="%", name="FFE_FR")
    cx.scalar("SSH", "Systems share of hard cost (MEP + specialist)", "IF(HARD_BASE>0,(EL_MEP+EL_SPEC)/HARD_BASE,0)", fmt="pct",
              unit="%", name="SYS_SH")
    cx.scalar("SFR", "Systems fraction of TDC", "(1-LAND_FR-FFE_FR)*SYS_SH", fmt="pct", unit="%", name="SYS_FR")
    cx.scalar("STFR", "Structure fraction of TDC", "1-LAND_FR-FFE_FR-SYS_FR", fmt="pct", unit="%", name="STR_FR")
    cx.scalar("RC", "Retained cost (IAS 40 + IAS 16)", "TDC*RET_SH", name="RET_COST")
    cx.scalar("BSTR", "Depreciable base – structure", "RET_COST*STR_FR", name="BASE_STR")
    cx.scalar("BSYS", "Depreciable base – systems", "RET_COST*SYS_FR", name="BASE_SYS")
    cx.scalar("BFFE", "Depreciable base – FF&E", "RET_COST*FFE_FR", name="BASE_FFE")
    cx.scalar("BLND", "Land (not depreciated)", "RET_COST*LAND_FR", name="BASE_LAND")
    cx.gap()
    cx.sec("Cost breakdown for charts (sums to TDC)", "የወጪ ክፍፍል")
    cats = [("CAT1", "Substructure & basement", "EL_SUBSTR+BASE_X"),
            ("CAT2", "Frame & envelope", "EL_SUPER+EL_ENVEL"),
            ("CAT3", "Internal works & finishes", "EL_INTERN+EL_FINISH"),
            ("CAT4", "MEP & specialist systems", "EL_MEP+EL_SPEC"),
            ("CAT5", "Preliminaries & external works", "EL_PRELIM+EL_EXTW+X_COMPOUND+X_SEPTIC+X_WTANK+X_PAVE"),
            ("CAT6", "Escalation & contingency", "ESC_AMT+CONT_TOT"),
            ("CAT7", "Fees & permits", "DESIGN_TOT+SUPERV_TOT+PERMIT_TOT"),
            ("CAT8", "Furniture & FF&E", "FFE_TOT"),
            ("CAT9", "Land", "LAND_TOT"),
            ("CAT10", "Finance costs (fees + IDC)", "FIN_CAP")]
    b.CAT_ROWS = []
    for k, lab, f in cats:
        b.CAT_ROWS.append(cx.scalar(k, lab, f))
    cx.scalar("CATT", "Total (= TDC)", "SUM(D{}:D{})".format(b.CAT_ROWS[0], b.CAT_ROWS[-1]), style="key", name="CAT_TOT")
    cx.gap()
    cx.sec("Contractor payment stages (construction cost incl. escalation & contingency)", "ለኮንትራክተር የክፍያ ደረጃዎች")
    stages = [("ST1", "1 Mobilisation & site set-up", "EL_PRELIM"),
              ("ST2", "2 Foundation / substructure", "EL_SUBSTR+BASE_X"),
              ("ST3", "3 Frame (columns, beams, slabs)", "EL_SUPER"),
              ("ST4", "4 Walls, windows & roof", "EL_ENVEL"),
              ("ST5", "5 MEP & internal works", "EL_MEP+EL_INTERN+EL_SPEC"),
              ("ST6", "6 Finishes & external works", "EL_FINISH+EL_EXTW+X_COMPOUND+X_SEPTIC+X_WTANK+X_PAVE")]
    b.STAGE_ROWS = []
    for k, lab, f in stages:
        b.STAGE_ROWS.append(cx.scalar(k, lab, f"IF(HARD_BASE>0,({f})/HARD_BASE,0)*(HARD_ESC_TOT+CONT_TOT)"))
    cx.scalar("STT", "Total contractor payments", "SUM(D{}:D{})".format(b.STAGE_ROWS[0], b.STAGE_ROWS[-1]), style="key",
              name="STAGE_TOT")

    # =================================================================== 23 CONSTRUCTION MONTHLY
    cm = Sheet(b, C23, "23 · Construction – monthly S-curve, funding waterfall, presales",
               "የግንባታ ወርሃዊ ወጪ፣ የገንዘብ ምንጭና ቅድመ-ሽያጭ", TAB_ENGINE, kind="m48", width_label=50)
    m = lambda key, i: cm.c(key, i)
    prev = lambda key, i: (cm.c(key, i - 1) if i > 1 else "0")
    TT = lambda key, i: f"'20_Timeline'!{pc(i)}{tl.rows[key]}"
    cm.sec("Timing", "ጊዜ")
    cm.line("CF", "Construction month flag", lambda i: TT("M_CF", i), fmt="int", unit="flag")
    cm.line("PIDX", "Annual period index", lambda i: TT("M_PIDX", i), fmt="int", unit="no.", tot=None)
    cm.sec("Spend profile (S-curve) and escalation", "የወጪ ስርጭትና የዋጋ ጭማሪ")
    cm.line("USERP", "User % profile (input row – used if profile = User %)", lambda i: 0, fmt="pct", unit="%", style="input")
    cm.line("IDX", "Rapha Construction Cost Index (input row, optional)", lambda i: 0, fmt="n1", unit="index", style="input", tot=None)
    cm.line("CUM", "Cumulative progress",
            lambda i: (f"IF({TT('M_T', i)}>=CONSTR_EFF,1,IF(SCURVE_TYPE=1,({TT('M_T', i)}/CONSTR_EFF)^SCURVE_A/(({TT('M_T', i)}/CONSTR_EFF)^SCURVE_A+(1-{TT('M_T', i)}/CONSTR_EFF)^SCURVE_A),"
                       f"IF(OR(SCURVE_TYPE=2,USER_TOT<=0),{TT('M_T', i)}/CONSTR_EFF,SUM($E${cm.rows['USERP']}:{m('USERP', i)})/USER_TOT)))"),
            fmt="pct", unit="%", tot=None)
    cm.line("S", "Monthly progress", lambda i: f"{m('CUM', i)}-{prev('CUM', i)}", fmt="pct", unit="%")
    cm.line("ESC", "Escalation factor",
            lambda i: f"IF(AND(INDEX_USE=1,INDEX_BASE>0,{m('IDX', i)}>0),{m('IDX', i)}/INDEX_BASE,(1+ESC_EFF)^(({TT('M_T', i)}-0.5)/12))",
            fmt="n2", unit="x", tot=None)
    cm.sec("Development cost accrual (IAS 2 / 16 / 40 / 23)", "የልማት ወጪ")
    cm.line("HARD", "Hard cost (escalated)", lambda i: f"HARD_BASE*{m('S', i)}*{m('ESC', i)}")
    cm.line("CONT", "Contingency", lambda i: f"{m('HARD', i)}*CONT_PCT")
    cm.line("DESIGN", "Design fees", lambda i: f"DESIGN_PCT*HARD_BASE*(IF({TT('M_T', i)}=1,DESIGN_UPFRONT,0)+(1-DESIGN_UPFRONT)*{m('S', i)})")
    cm.line("SUPERV", "Supervision", lambda i: f"SUPERV_PCT*{m('HARD', i)}")
    cm.line("PERMIT", "Permits & fees", lambda i: "PERMITS" if i == 1 else 0)
    cm.line("FFE", "FF&E / furniture", lambda i: f"IF({TT('M_T', i)}=CONSTR_EFF,FFE,0)")
    cm.line("LANDACC", "Land (recognised month 1)", lambda i: "LAND_VALUE" if i == 1 else 0)
    cm.line("ARRF", "Arrangement fee (on IDC-excluded facility)", lambda i: "ARR_FEE*DEBT_CAP" if i == 1 else 0)
    cm.line("COMF", "Commitment fee on undrawn facility",
            lambda i: f"IF({TT('M_T', i)}<=CONSTR_EFF,MAX(0,DEBT_CAP-{prev('CUMDEBT', i)})*COMMIT_FEE/12,0)")
    cm.line("IDC", "Interest during construction (from loan schedule)", lambda i: X("27_Loan", "CAPC", i))
    cm.line("ACC", "Total capex accrual", lambda i: "+".join(m(k, i) for k in ("HARD", "CONT", "DESIGN", "SUPERV", "PERMIT", "FFE",
                                                                                  "LANDACC", "ARRF", "COMF", "IDC")),
            style="total")
    cm.line("CUMACC", "Cumulative capex accrual", lambda i: f"{prev('CUMACC', i)}+{m('ACC', i)}", tot="last")
    cm.sec("Cash needed during construction", "ለግንባታ የሚያስፈልግ ጥሬ ገንዘብ")
    cm.line("RETW", "Retention withheld from contractor", lambda i: f"RETENTION*{m('HARD', i)}")
    cm.line("LANDC", "Land cash (lease down payment)", lambda i: "IF(LAND_TYPE=1,LAND_VALUE*LEASE_DP,0)" if i == 1 else 0)
    cm.line("DEVC", "Development cash paid (excl. finance costs)",
            lambda i: f"{m('LANDC', i)}+{m('HARD', i)}-{m('RETW', i)}+{m('CONT', i)}+{m('DESIGN', i)}+{m('SUPERV', i)}+{m('PERMIT', i)}+{m('FFE', i)}")
    cm.line("VATC", "Input VAT on capex (commercial modes; reclaimable)", lambda i: f"IF(PROJECT_MODE=0,0,1)*VAT_R*VAT_CAPEX_PCT*({m('HARD', i)}+{m('CONT', i)}+{m('DESIGN', i)}+{m('SUPERV', i)}+{m('FFE', i)})")
    cm.line("LSI", "Land lease instalment (during construction)", lambda i: f"{m('CF', i)}*{X('27_Loan', 'LS_PR', i)}")
    cm.line("LSINT", "Land lease interest (during construction)", lambda i: f"{m('CF', i)}*{X('27_Loan', 'LS_INT', i)}")
    cm.line("NEEDX", "Funding need excl. finance costs",
            lambda i: f"{m('CF', i)}*({m('DEVC', i)}+{m('VATC', i)}+{m('LSI', i)}+{m('LSINT', i)})")
    cm.line("FINC", "Finance costs paid in cash (fees)", lambda i: f"{m('ARRF', i)}+{m('COMF', i)}")
    cm.line("NEED", "TOTAL FUNDING NEED", lambda i: f"{m('NEEDX', i)}+{m('FINC', i)}", style="total")
    cm.line("CUMNEED", "Cumulative funding need", lambda i: f"{prev('CUMNEED', i)}+{m('NEED', i)}", tot="max")
    cm.sec("Presales (contracts signed during construction)", "ቅድመ-ሽያጭ")
    cm.line("BAND", "Construction year (presale band)", lambda i: f"ROUNDUP({TT('M_T', i)}/12,0)", fmt="int", unit="no.", tot=None)
    cm.line("MIB", "Construction months in this band",
            lambda i: f"MAX(1,MIN(12,CONSTR_EFF-12*({m('BAND', i)}-1)))", fmt="int", unit="no.", tot=None)
    cm.line("PSA", "Area presold", lambda i: f"{m('CF', i)}*SALE_AREA*INDEX(PS_EFF,{m('BAND', i)})/{m('MIB', i)}", unit="m²")
    cm.line("PRC", "Sale price (escalated)", lambda i: f"BLEND_PRICE*(1+PESC_EFF)^(({TT('M_T', i)}-1)/12)", unit="ETB/m²", tot=None)
    cm.line("CV", "Contract value signed (net of default allowance)", lambda i: f"{m('PSA', i)}*{m('PRC', i)}*(1-DEFAULT_EFF)")
    cm.line("CUMCV", "Cumulative contract value", lambda i: f"{prev('CUMCV', i)}+{m('CV', i)}", tot="last")
    cm.line("CUMAREA", "Cumulative area presold", lambda i: f"{prev('CUMAREA', i)}+{m('PSA', i)}", unit="m²", tot="last")
    cm.line("DP", "Collections – down payments", lambda i: f"PAY_DP*{m('CV', i)}")
    cm.line("CUMINST", "Cumulative milestone instalments due", lambda i: f"PAY_INST*{m('CUMCV', i)}*{m('CUM', i)}", tot="last")
    cm.line("INST", "Collections – milestone instalments", lambda i: f"{m('CUMINST', i)}-{prev('CUMINST', i)}")
    cm.line("HAND", "Collections – balance at handover", lambda i: f"IF({TT('M_T', i)}=CONSTR_EFF,(1-PAY_DP-PAY_INST)*{m('CUMCV', i)},0)")
    cm.line("COLL", "Total presale collections", lambda i: f"{m('DP', i)}+{m('INST', i)}+{m('HAND', i)}", style="total")
    cm.line("CUMREV", "Cumulative revenue recognised (IFRS 15)",
            lambda i: f"IF(REV_POLICY=2,{m('CUMCV', i)}*{m('CUM', i)},IF({TT('M_T', i)}>=CONSTR_EFF,{m('CUMCV', i)},0))", tot="last")
    cm.line("REV", "Presale revenue recognised", lambda i: f"{m('CUMREV', i)}-{prev('CUMREV', i)}")
    cm.line("CUMCOGS", "Cumulative cost of sales released (IAS 2)",
            lambda i: f"IF(SALE_AREA>0,TDC_INV*{m('CUMAREA', i)}/SALE_AREA*IF(REV_POLICY=2,{m('CUM', i)},IF({TT('M_T', i)}>=CONSTR_EFF,1,0)),0)",
            tot="last")
    cm.line("COGS", "Presale cost of sales", lambda i: f"{m('CUMCOGS', i)}-{prev('CUMCOGS', i)}")
    cm.sec("Funding waterfall: (1) presales → (2) equity → (3) debt → (4) top-up / gap", "የገንዘብ ምንጭ ቅደም ተከተል")
    cm.line("USABLE", "Presale cash usable for construction",
            lambda i: f"{m('CF', i)}*IF(PRESALE_USE=1,PRESALE_USE_PCT,0)*{m('COLL', i)}")
    cm.line("POOLO", "Presale pool – opening", lambda i: prev("POOLC", i), tot=None)
    cm.line("PSUSED", "(1) Presale cash used", lambda i: f"MIN({m('NEED', i)},{m('POOLO', i)}+{m('USABLE', i)})")
    cm.line("POOLC", "Presale pool – closing", lambda i: f"{m('POOLO', i)}+{m('USABLE', i)}-{m('PSUSED', i)}", tot="last")
    cm.line("RESID", "Need after presales", lambda i: f"{m('NEED', i)}-{m('PSUSED', i)}")
    cm.line("EQ1", "(2) Equity draw",
            lambda i: f"MIN({m('RESID', i)}*IF(DRAW_ORDER=1,1,EQ_SHARE),MAX(0,EQ_COMMIT-{prev('CUMEQ', i)}))")
    cm.line("DEBT", "(3) Debt draw", lambda i: f"MIN({m('RESID', i)}-{m('EQ1', i)},MAX(0,DEBT_CAP-{prev('CUMDEBT', i)}))")
    cm.line("EQ2", "(2b) Equity – pro-rata balance",
            lambda i: f"IF(DRAW_ORDER=2,MIN({m('RESID', i)}-{m('EQ1', i)}-{m('DEBT', i)},MAX(0,EQ_COMMIT-{prev('CUMEQ', i)}-{m('EQ1', i)})),0)")
    cm.line("TOPUP", "(4) Sponsor top-up / funding gap", lambda i: f"{m('RESID', i)}-{m('EQ1', i)}-{m('DEBT', i)}-{m('EQ2', i)}")
    cm.line("CUMEQ", "Cumulative committed equity drawn", lambda i: f"{prev('CUMEQ', i)}+{m('EQ1', i)}+{m('EQ2', i)}", tot="last")
    cm.line("CUMDEBT", "Cumulative debt drawn", lambda i: f"{prev('CUMDEBT', i)}+{m('DEBT', i)}", tot="last")
    cm.line("EQT", "Total equity in (incl. top-up)", lambda i: f"{m('EQ1', i)}+{m('EQ2', i)}+{m('TOPUP', i)}")
    cm.line("CASHH", "Presale cash held (pool + escrow), cumulative",
            lambda i: f"{prev('CASHH', i)}+{m('COLL', i)}-{m('PSUSED', i)}", tot="last")
    cm.line("SRCCHK", "Sources − uses (should be 0)",
            lambda i: f"{m('PSUSED', i)}+{m('EQT', i)}+{m('DEBT', i)}-{m('NEED', i)}", tot="sum")
    cm.gap()
    cm.sec("Funding scalars", "የገንዘብ ማጠቃለያ")
    cm.scalar("UTOT", "User profile total (construction months)",
              f"SUMPRODUCT(E{cm.rows['USERP']}:{pc(MON_C)}{cm.rows['USERP']},E{cm.rows['CF']}:{pc(MON_C)}{cm.rows['CF']})",
              fmt="pct", unit="%", name="USER_TOT")
    cm.scalar("FB", "Funding base (need excl. IDC & finance fees)", f"$D${cm.rows['NEEDX']}", name="FUND_BASE")
    cm.scalar("EQC", "Equity commitment", "IF(PROJECT_MODE=0,SAVINGS+FAMILY,EQ_PCT*FUND_BASE)", name="EQ_COMMIT")
    cm.scalar("DC", "Debt facility (cap)",
              "IF(PROJECT_MODE=0,HOME_LOAN,IF(LOAN_CAP>0,MIN(LOAN_CAP,(1-EQ_PCT)*FUND_BASE),(1-EQ_PCT)*FUND_BASE))", name="DEBT_CAP")
    cm.scalar("EQS", "Equity share for pro-rata draws", "IF(EQ_COMMIT+DEBT_CAP>0,EQ_COMMIT/(EQ_COMMIT+DEBT_CAP),1)", fmt="pct",
              unit="%", name="EQ_SHARE")
    cm.scalar("PEAKF", "Peak cumulative funding requirement", f"$D${cm.rows['CUMNEED']}", name="PEAK_FUND", style="key")
    cm.scalar("EQTOT", "Equity drawn (committed)", f"$D${cm.rows['CUMEQ']}", name="EQ_DRAWN")
    cm.scalar("DTOT", "Debt drawn", f"$D${cm.rows['CUMDEBT']}", name="DEBT_DRAWN")
    cm.scalar("GAP", "Sponsor top-up / funding gap", f"$D${cm.rows['TOPUP']}", name="FUND_GAP", style="key")
    cm.scalar("PSU", "Presale cash used for construction", f"$D${cm.rows['PSUSED']}", name="PS_USED")
    cm.scalar("HANDT", "Handover collections at completion", f"$D${cm.rows['HAND']}", name="HAND_TOT")
    cm.scalar("RETT", "Total retention withheld", f"$D${cm.rows['RETW']}", name="RET_TOT")
    cm.scalar("LANDCT", "Land paid in cash in month 1", f"$D${cm.rows['LANDC']}", name="LANDCASH_TOT")
    cm.scalar("SRC", "Sources − uses (should be 0)", f"$D${cm.rows['SRCCHK']}", name="SRC_DIFF")

    # =================================================================== 27 LOAN (monthly, 288)
    ln = Sheet(b, "27_Loan", "27 · Loan – amortization schedule", "የብድር ክፍያ ሰንጠረዥ", TAB_FIN, kind="m288", width_label=48)
    L = lambda key, i: ln.c(key, i)
    lp = lambda key, i: (ln.c(key, i - 1) if i > 1 else "0")
    # annual roll-up block lives on a separate annual sheet section (below) -> use helper sheet rows

    ln.sec("Annual roll-up (C1–C4, O1–O10)", "ዓመታዊ ማጠቃለያ")
    SUMP = lambda key, i: f"SUMIFS($E${ln.rows[key]}:${pc(MON_L)}${ln.rows[key]},$E${ln.rows['PIDX']}:${pc(MON_L)}${ln.rows['PIDX']},{T('T_P', i)})"
    ATM = lambda key, i: f"SUMIFS($E${ln.rows[key]}:${pc(MON_L)}${ln.rows[key]},$E${ln.rows['T']}:${pc(MON_L)}${ln.rows['T']},{T('T_ENDM', i)})"
    ln.line("A_LBL", "Period", lambda i: "§" + PERIOD_LABELS[i - 1], fmt="gen", unit="", tot=None, n=ANN_N)
    ln.line("A_DRAW", "Drawdowns", lambda i: SUMP("DRAW", i), n=ANN_N)
    ln.line("A_CAPC", "IDC capitalised (construction)", lambda i: SUMP("CAPC", i), n=ANN_N)
    ln.line("A_CAPG", "Interest capitalised (grace)", lambda i: SUMP("CAPG", i), n=ANN_N)
    ln.line("A_INTP", "Interest paid", lambda i: SUMP("INTP", i), n=ANN_N)
    ln.line("A_INTE", "Interest expense (P&L)", lambda i: f"{ln.c('A_INTP', i)}+{ln.c('A_CAPG', i)}", n=ANN_N)
    ln.line("A_PRIN", "Scheduled principal", lambda i: SUMP("PRIN", i), n=ANN_N)
    ln.line("A_SWP", "Cash sweep prepayment", lambda i: SUMP("SWP", i), n=ANN_N)
    ln.line("A_DS", "Debt service (interest paid + scheduled principal)", lambda i: f"{ln.c('A_INTP', i)}+{ln.c('A_PRIN', i)}",
            n=ANN_N, style="total")
    ln.line("A_CLOSE", "Closing loan balance", lambda i: ATM("CLOSE", i), n=ANN_N, tot="last", style="total")
    ln.line("A_OPEN", "Opening loan balance", lambda i: ln.c("A_CLOSE", i - 1) if i > 1 else 0, n=ANN_N, tot=None)
    ln.line("A_LSPR", "Land lease principal paid", lambda i: SUMP("LS_PR", i), n=ANN_N)
    ln.line("A_LSINT", "Land lease interest", lambda i: SUMP("LS_INT", i), n=ANN_N)
    ln.line("A_LSREC", "Lease liability recognised", lambda i: SUMP("LS_REC", i), n=ANN_N)
    ln.line("A_LSCL", "Closing lease liability", lambda i: ATM("LS_CLOSE", i), n=ANN_N, tot="last")

    ln.gap()
    ln.sec("Loan scalars and lender summary", "የብድር ማጠቃለያ")
    ln.scalar("FM", "Months between repayments", "CHOOSE(FREQ,1,3,12)", fmt="int", unit="months", name="FREQ_M")
    ln.scalar("TM", "Tenor after completion", "TENOR_Y*12", fmt="int", unit="months", name="TENOR_M")
    ln.scalar("NP", "Number of scheduled repayments", "MAX(0,INT((TENOR_M-GRACE_M)/FREQ_M))", fmt="int", unit="no.", name="N_PAY")
    ln.scalar("LI", "Annual lease instalment", "IF(LEASE_N>0,LAND_VALUE*(1-LEASE_DP)/LEASE_N,0)*IF(LAND_TYPE=1,1,0)",
              name="LEASE_INST")
    ln.scalar("LDR", "Total drawdowns", lambda: f"$D${ln.rows['DRAW']}", name="LN_DRAWS")
    ln.scalar("LIDC", "Total interest capitalised (IDC + grace)", lambda: f"$D${ln.rows['CAPC']}+$D${ln.rows['CAPG']}", name="LN_CAPINT")
    ln.scalar("LINT", "Total interest (all)", lambda: f"$D${ln.rows['INT']}", name="LN_INT_TOT", style="key")
    ln.scalar("LPR", "Total principal repaid (scheduled + sweep)", lambda: f"$D${ln.rows['PRIN']}+$D${ln.rows['SWP']}", name="LN_PRIN_TOT")
    ln.scalar("PEAK", "Peak debt (incl. capitalised interest)", lambda: f"$D${ln.rows['CLOSE']}", name="PEAK_DEBT", style="key")
    ln.scalar("B0", "Balance at start of repayment", lambda: f"$D${ln.rows['B0']}", name="LN_B0")
    ln.scalar("AVL", "Average life after completion", lambda: "IF(LN_PRIN_TOT>0,{}/LN_PRIN_TOT,0)".format(f"$D${ln.rows['KP']}"),
              fmt="n1", unit="years", name="LN_AVG_LIFE")
    ln.scalar("TLAST", "Final repayment month index", lambda: f"$D${ln.rows['TP']}", fmt="int", unit="no.", name="LN_TLAST")
    ln.scalar("FRD", "Final repayment date", "IF(LN_TLAST>0,EDATE(START_DATE,LN_TLAST)-1,0)", fmt="date", unit="date",
              name="LN_FINAL_DATE")
    ln.scalar("TFIRST", "First scheduled repayment month", lambda: f"IF($D${ln.rows['TF']}=999,0,$D${ln.rows['TF']})", fmt="int",
              unit="no.", name="LN_TFIRST")
    ln.scalar("INSTAL", "Instalment at first scheduled repayment",
              lambda: f"IF(LN_TFIRST>0,INDEX(E{ln.rows['DS']}:{pc(MON_L)}{ln.rows['DS']},LN_TFIRST),0)", name="LN_INSTAL", style="key")
    ln.scalar("INSTM", "Monthly-equivalent instalment", "LN_INSTAL/FREQ_M", name="LN_INSTAL_M", style="key")
    ln.scalar("BTEN", "Balance outstanding at end of tenor",
              lambda: f"SUMIFS(E{ln.rows['CLOSE']}:{pc(MON_L)}{ln.rows['CLOSE']},E{ln.rows['K']}:{pc(MON_L)}{ln.rows['K']},MIN(TENOR_M,{MON_L}-CONSTR_EFF))",
              name="LN_BAL_TENOR")
    ln.scalar("BEND", "Balance outstanding at end of O10", lambda: f"SUMIFS(E{ln.rows['CLOSE']}:{pc(MON_L)}{ln.rows['CLOSE']},E{ln.rows['K']}:{pc(MON_L)}{ln.rows['K']},12*OPS_YEARS)",
              name="LN_BAL_O10")
    ln.scalar("MINDS", "Min DSCR (from ratios)", lambda: XS("43_Ratios", "MIN"), fmt="x", unit="x")
    ln.scalar("AVGDS", "Average DSCR", lambda: XS("43_Ratios", "AVG"), fmt="x", unit="x")
    ln.scalar("LLCR", "LLCR at completion", lambda: XS("43_Ratios", "LLCR1"), fmt="x", unit="x")
    ln.scalar("MICR", "Min interest cover (ICR)", lambda: XS("43_Ratios", "MINICR"), fmt="x", unit="x")
    ln.scalar("MGR", "Max gearing", lambda: XS("43_Ratios", "MAXG"), fmt="pct", unit="%")
    ln.gap()
    b.LOAN_PRINT_END = ln.r
    ln.sec("Timing and rate", "ጊዜና የወለድ መጠን")
    ln.line("T", "Month t", lambda i: TT("M_T", i), fmt="int", unit="no.", tot=None)
    ln.line("CF", "Construction flag", lambda i: TT("M_CF", i), fmt="int", unit="flag")
    ln.line("K", "Months after completion k", lambda i: TT("M_K", i), fmt="int", unit="no.", tot=None)
    ln.line("LY", "Loan year after completion", lambda i: TT("M_LY", i), fmt="int", unit="no.", tot=None)
    ln.line("PIDX", "Annual period index", lambda i: TT("M_PIDX", i), fmt="int", unit="no.", tot=None)
    ln.line("RATE", "Interest rate (p.a.)",
            lambda i: f"MAX(0,(IF(AND(RATE_TYPE=2,{L('LY', i)}>0),INDEX(RATE_PATH,MIN(10,MAX(1,{L('LY', i)}))),RATE)+D_RATE_ADD)*D_RATE_F)",
            fmt="pct2", unit="%", tot=None)
    ln.line("GR", "Grace flag", lambda i: f"IF(AND({L('K', i)}>=1,{L('K', i)}<=GRACE_M),1,0)", fmt="int", unit="flag")
    ln.line("EV", "Repayment event flag",
            lambda i: f"IF(AND({L('K', i)}>GRACE_M,{L('K', i)}<=GRACE_M+N_PAY*FREQ_M,MOD({L('K', i)}-GRACE_M,FREQ_M)=0),1,0)",
            fmt="int", unit="flag")
    ln.line("REM", "Remaining payments incl. this one",
            lambda i: f"IF({L('EV', i)}=1,N_PAY-({L('K', i)}-GRACE_M)/FREQ_M+1,0)", fmt="int", unit="no.", tot=None)
    ln.sec("Loan balance", "የብድር ቀሪ ሂሳብ")
    ln.line("OPEN", "Opening balance", lambda i: lp("CLOSE", i), tot=None)
    ln.line("DRAW", "Drawdowns", lambda i: X(C23, "DEBT", i) if i <= MON_C else 0)
    ln.line("INT", "Interest accrued", lambda i: f"({L('OPEN', i)}+{L('DRAW', i)})*{L('RATE', i)}/12")
    ln.line("CAPC", "Interest capitalised – construction (IDC)", lambda i: f"IF({L('CF', i)}=1,{L('INT', i)},0)")
    ln.line("CAPG", "Interest capitalised – grace", lambda i: f"IF(AND({L('GR', i)}=1,GRACE_TYPE=2),{L('INT', i)},0)")
    ln.line("INTP", "Interest paid in cash", lambda i: f"{L('INT', i)}-{L('CAPC', i)}-{L('CAPG', i)}")
    ln.line("B0", "Balance at start of repayment",
            lambda i: f"IF({L('K', i)}=GRACE_M+1,{L('OPEN', i)},{lp('B0', i)})", tot="max")
    ln.line("RP", "Periodic rate", lambda i: f"{L('RATE', i)}*FREQ_M/12", fmt="pct2", unit="%", tot=None)
    ln.line("BFV", "Balloon amount", lambda i: f"IF(REPAY_TYPE=3,BALLOON_PCT*{L('B0', i)},0)", tot=None)
    ln.line("PRIN", "Scheduled principal",
            lambda i: (f"IF({L('EV', i)}=1,IF({L('REM', i)}<=1,{L('OPEN', i)}+{L('CAPG', i)},MAX(0,MIN({L('OPEN', i)}+{L('CAPG', i)},"
                       f"IF(REPAY_TYPE=2,({L('OPEN', i)}-{L('BFV', i)})/{L('REM', i)},PMT({L('RP', i)},{L('REM', i)},-{L('OPEN', i)},{L('BFV', i)})-{L('OPEN', i)}*{L('RP', i)})))),0)"))
    ln.line("SWB", "Cash sweep base (sales proceeds)",
            lambda i: (f"IF({L('K', i)}=1,HAND_TOT,0)+IF(AND({L('K', i)}>=12,MOD({L('K', i)},12)=0,{L('K', i)}<=12*OPS_YEARS),"
                       f"INDEX(SALES_COLL_OPS,MAX(1,MIN(10,{L('K', i)}/12))),0)"))
    ln.line("SWP", "Cash sweep prepayment",
            lambda i: f"MIN(MAX(0,{L('OPEN', i)}+{L('DRAW', i)}+{L('CAPC', i)}+{L('CAPG', i)}-{L('PRIN', i)}),SWEEP*{L('SWB', i)})")
    ln.line("CLOSE", "Closing balance",
            lambda i: f"{L('OPEN', i)}+{L('DRAW', i)}+{L('CAPC', i)}+{L('CAPG', i)}-{L('PRIN', i)}-{L('SWP', i)}", tot="max",
            style="total")
    ln.line("DS", "Scheduled debt service (interest paid + principal)", lambda i: f"{L('INTP', i)}+{L('PRIN', i)}")
    ln.line("TP", "Helper: month index if principal paid", lambda i: f"IF({L('PRIN', i)}+{L('SWP', i)}>0,{L('T', i)},0)",
            fmt="int", unit="no.", tot="max")
    ln.line("TF", "Helper: month index of scheduled payment (else 999)", lambda i: f"IF({L('PRIN', i)}>0,{L('T', i)},999)",
            fmt="int", unit="no.", tot="min")
    ln.line("KP", "Helper: principal × years after completion", lambda i: f"({L('PRIN', i)}+{L('SWP', i)})*MAX(0,{L('K', i)})/12",
            fmt="n0")
    ln.sec("Land lease liability", "የሊዝ ዕዳ")
    ln.line("LS_REC", "Lease liability recognised", lambda i: "IF(LAND_TYPE=1,LAND_VALUE*(1-LEASE_DP),0)" if i == 1 else 0)
    ln.line("LS_OPEN", "Opening lease liability", lambda i: lp("LS_CLOSE", i), tot=None)
    ln.line("LS_FL", "Instalment month flag",
            lambda i: f"IF(AND(LAND_TYPE=1,{L('T', i)}>1,MOD({L('T', i)}-1,12)=0,({L('T', i)}-1)/12<=LEASE_N),1,0)", fmt="int", unit="flag")
    ln.line("LS_PR", "Lease instalment (principal)", lambda i: f"{L('LS_FL', i)}*MIN({L('LS_OPEN', i)},LEASE_INST)")
    ln.line("LS_INT", "Lease interest", lambda i: f"{L('LS_FL', i)}*{L('LS_OPEN', i)}*LEASE_INT")
    ln.line("LS_CLOSE", "Closing lease liability", lambda i: f"{L('LS_OPEN', i)}+{L('LS_REC', i)}-{L('LS_PR', i)}", tot="last")
    return tl, ar, cx, cm, ln


def build_engine2(b):
    X = lambda sh, key, i: b.sheets[sh].c(key, i, sheet=True)
    XS = lambda sh, key: b.sheets[sh].s(key)
    XT = lambda sh, key: b.sheets[sh].tot(key)
    T = lambda key, i: f"'20_Timeline'!{pc(i)}{b.TL[key]}"
    C23 = "23_Construction"
    cmr = b.sheets[C23].rows
    SUMC = lambda key, i: f"SUMIFS('{C23}'!$E${cmr[key]}:${pc(MON_C)}${cmr[key]},'{C23}'!$E${cmr['PIDX']}:${pc(MON_C)}${cmr['PIDX']},{T('T_P', i)})"
    LN = "27_Loan"

    # =================================================================== 24 REVENUE
    rv = Sheet(b, "24_Revenue", "24 · Revenue – sales, rent, other income", "ገቢ – ሽያጭ፣ ኪራይና ሌሎች", TAB_ENGINE, banner=True)
    r = lambda key, i: rv.c(key, i)
    rp = lambda key, i: (rv.c(key, i - 1) if i > 1 else "0")
    rv.sec("Presale bands (effective after scenario, capped at 100%)", "የቅድመ-ሽያጭ ድርሻ")
    for n in range(1, 5):
        rv.scalar(f"PSR{n}", f"Presale band {n} – raw",
                  f"PRESALE_{n}*D_PRESALE_F*IF({T('T_MONTHS', n)}>0,1,0)", fmt="pct", unit="%")
    rv.scalar("PSRT", "Presales raw total", "SUM(D{}:D{})".format(rv.scal["PSR1"], rv.scal["PSR4"]), fmt="pct", unit="%")
    for n in range(1, 5):
        rv.scalar(f"PSE{n}", f"Presale band {n} – effective",
                  f"D{rv.scal[f'PSR{n}']}*IF(D{rv.scal['PSRT']}>1,1/D{rv.scal['PSRT']},1)", fmt="pct", unit="%")
    b.name("PS_EFF", f"'24_Revenue'!$D${rv.scal['PSE1']}:$D${rv.scal['PSE4']}")
    rv.scalar("PRES", "Total presold share", "SUM(PS_EFF)", fmt="pct", unit="%", name="PRESOLD")
    rv.scalar("UNS", "Unsold area at completion", "SALE_AREA*(1-PRESOLD)", unit="m²", name="UNSOLD_AREA")
    for n in range(1, 5):
        rv.scalar(f"ABE{n}", f"Absorption O{n} – effective", f"MIN(1,ABS_{n}*ABS_F)", fmt="pct", unit="%")
    rv.gap()
    rv.sec("Sales – presales (from monthly sheet)", "ቅድመ-ሽያጭ")
    rv.line("PS_AREA", "Area presold", lambda i: SUMC("PSA", i), unit="m²")
    rv.line("PS_CV", "Contract value signed", lambda i: SUMC("CV", i))
    rv.line("PS_COLL", "Presale collections", lambda i: SUMC("COLL", i))
    rv.line("PS_REV", "Presale revenue recognised", lambda i: SUMC("REV", i))
    rv.line("PS_COGS", "Presale cost of sales", lambda i: SUMC("COGS", i))
    rv.sec("Sales – after completion (point in time, cash on sale)", "ከግንባታ በኋላ ሽያጭ")
    absr = lambda n: f"$D${rv.scal[f'ABE{n}']}"
    rv.line("PO_AREA", "Area sold",
            lambda i: 0 if i <= 4 else (f"MIN({rp('PO_REM', i) if i > 5 else 'UNSOLD_AREA'},UNSOLD_AREA*{absr(Y(i))})" if Y(i) <= 4
                                        else (rp("PO_REM", i) if Y(i) == 5 else 0)), unit="m²")
    rv.line("PO_REM", "Unsold area remaining", lambda i: 0 if i <= 4 else f"{rp('PO_REM', i) if i > 5 else 'UNSOLD_AREA'}-{r('PO_AREA', i)}",
            unit="m²", tot="last")
    rv.line("PO_PRICE", "Sale price (escalated)", lambda i: 0 if i <= 4 else f"BLEND_PRICE*(1+PESC_EFF)^(CONSTR_EFF/12+{T('T_Y', i)}-0.5)",
            unit="ETB/m²", tot=None)
    rv.line("PO_REV", "Sales revenue = collections", lambda i: f"{r('PO_AREA', i)}*{r('PO_PRICE', i)}*(1-DEFAULT_EFF)")
    rv.line("PO_COGS", "Cost of sales", lambda i: f"IF(SALE_AREA>0,TDC_INV*{r('PO_AREA', i)}/SALE_AREA,0)")
    rv.sec("Sales totals and contract balances (IFRS 15)", "የሽያጭ ድምር")
    rv.line("SALES_REV", "Sales revenue recognised", lambda i: f"{r('PS_REV', i)}+{r('PO_REV', i)}", style="total")
    rv.line("SALES_COLL", "Sales cash collected", lambda i: f"{r('PS_COLL', i)}+{r('PO_REV', i)}")
    rv.line("COGS", "Cost of sales (IAS 2)", lambda i: f"{r('PS_COGS', i)}+{r('PO_COGS', i)}")
    rv.line("CVS", "Contract value signed in period", lambda i: f"{r('PS_CV', i)}+{r('PO_REV', i)}")
    rv.line("COMM", "Sales commission", lambda i: f"COMM_PCT*{r('SALES_REV', i)}")
    rv.line("MKT", "Marketing", lambda i: f"MKT_PCT*{r('CVS', i)}")
    rv.line("CUMCOLL", "Cumulative collections", lambda i: f"{rp('CUMCOLL', i)}+{r('SALES_COLL', i)}", tot="last")
    rv.line("CUMREV", "Cumulative revenue", lambda i: f"{rp('CUMREV', i)}+{r('SALES_REV', i)}", tot="last")
    rv.line("CBAL", "Contract balance (+ liability / − asset)", lambda i: f"{r('CUMCOLL', i)}-{r('CUMREV', i)}", tot="last")
    rv.line("CLIAB", "Contract liability", lambda i: f"MAX(0,{r('CBAL', i)})", tot="last")
    rv.line("CASSET", "Contract asset", lambda i: f"MAX(0,-{r('CBAL', i)})", tot="last")
    rv.line("SC_OPS", "Sales collections in operating year (for sweep)", lambda i: 0 if i <= 4 else r("SALES_COLL", i))
    b.name("SALES_COLL_OPS", f"'24_Revenue'!$I${rv.rows['SC_OPS']}:$R${rv.rows['SC_OPS']}")
    rv.sec("Rent (IFRS 16 lessor)", "ኪራይ")
    rv.line("GPR", "Gross potential rent",
            lambda i: 0 if i <= 4 else f"GPR_TODAY*D_RENT_F*(1+RESC_EFF)^(CONSTR_EFF/12+{T('T_Y', i)}-1)")
    rv.line("OCC", "Occupancy",
            lambda i: 0 if i <= 4 else f"MIN(1,MAX(0,(CHOOSE(MIN({T('T_Y', i)},3),OCC_O1,OCC_O2,OCC_STAB)+D_OCC_ADD)*D_OCC_F))",
            fmt="pct", unit="%", tot=None)
    rv.line("RENT", "Rental income", lambda i: f"{r('GPR', i)}*{r('OCC', i)}", style="total")
    rv.line("KEYR", "Key money received", lambda i: f"IF(PROJECT_MODE=0,0,KEY_MONTHS*{r('RENT', i)}/12)" if i == 5 else 0)
    rv.line("KEYA", "Key money released to income",
            lambda i: 0 if i <= 4 else f"IF({T('T_Y', i)}<=MAX(1,KEY_AMORT_Y),$D${rv.rows['KEYR']}/MAX(1,KEY_AMORT_Y),0)")
    rv.line("KEYD", "Deferred key money (liability)", lambda i: f"{rp('KEYD', i)}+{r('KEYR', i)}-{r('KEYA', i)}", tot="last")
    rv.line("SCR", "Service charge recoveries", lambda i: f"SC_RECOV*{X('25_Opex', 'PROP', i)}")
    rv.line("OTH", "Other income",
            lambda i: 0 if i <= 4 else f"IF(PROJECT_MODE=0,0,OTHER_INC*(1+OPEX_ESC)^({T('T_Y', i)}-1))")
    rv.line("TOTREV", "TOTAL REVENUE", lambda i: "+".join(r(k, i) for k in ("SALES_REV", "RENT", "KEYA", "SCR", "OTH")),
            style="grand")
    rv.line("NOI", "Net operating income of let property (NOI)",
            lambda i: f"{r('RENT', i)}+{r('SCR', i)}+{r('OTH', i)}-{X('25_Opex', 'PROP', i)}-{X('25_Opex', 'MGMT', i)}-{X('25_Opex', 'INS', i)}-{X('25_Opex', 'PTAX', i)}-{X('25_Opex', 'MAINT', i)}-{X('25_Opex', 'BAD', i)}",
            style="key")

    # =================================================================== 25 OPEX
    ox = Sheet(b, "25_Opex", "25 · Operating costs", "የሥራ ማስኬጃ ወጪዎች", TAB_ENGINE, banner=True)
    o = lambda key, i: ox.c(key, i)
    RV = lambda key, i: X("24_Revenue", key, i)
    ox.sec("Operating costs (positive = cost)", "ወጪዎች")
    ox.line("ESC", "Opex escalation index (from today)", lambda i: 0 if i <= 4 else f"(1+OPEX_ESC)^(CONSTR_EFF/12+{T('T_Y', i)}-1)",
            fmt="n2", unit="x", tot=None)
    ox.line("PROP", "Property operating costs",
            lambda i: 0 if i <= 4 else f"IF(OPEX_METHOD=1,OPEX_PCT*{RV('RENT', i)},OPEX_M2*LET_AREA*{o('ESC', i)})*D_OPEX_F")
    ox.line("MGMT", "Management fee", lambda i: f"MGMT_PCT*{RV('RENT', i)}")
    ox.line("INS", "Insurance", lambda i: 0 if i <= 4 else f"INS_PCT*HARD_ESC_TOT*RET_SH*{o('ESC', i)}*D_OPEX_F")
    ox.line("PTAX", "Property tax", lambda i: 0 if i <= 4 else f"PROPTAX_PCT*(HARD_ESC_TOT+LAND_TOT)*RET_SH*{o('ESC', i)}")
    ox.line("MAINT", "Maintenance reserve", lambda i: 0 if i <= 4 else f"MAINT_PCT*HARD_ESC_TOT*RET_SH*{o('ESC', i)}*D_OPEX_F")
    ox.line("MKT", "Marketing", lambda i: RV("MKT", i))
    ox.line("OVH", "Owner overheads", lambda i: 0 if i <= 4 else f"IF(PROJECT_MODE=0,0,OVERHEAD*{o('ESC', i)}*D_OPEX_F)")
    ox.line("BAD", "Bad debts", lambda i: f"BAD_DEBT*{RV('RENT', i)}")
    ox.line("TOT", "TOTAL OPERATING COSTS", lambda i: "+".join(o(k, i) for k in ("PROP", "MGMT", "INS", "PTAX", "MAINT", "MKT", "OVH", "BAD")),
            style="total")
    ox.line("VATB", "Opex base carrying input VAT", lambda i: f"OPEX_VAT_PCT*({o('PROP', i)}+{o('MGMT', i)}+{o('MAINT', i)}+{o('MKT', i)})")

    # =================================================================== 26 ASSETS
    asx = Sheet(b, "26_Assets", "26 · Assets & depreciation (IAS 2 / 16 / 40)", "ንብረትና የእርጅና ቅናሽ", TAB_ENGINE, banner=True)
    a = lambda key, i: asx.c(key, i)
    ap = lambda key, i: (asx.c(key, i - 1) if i > 1 else "0")
    asx.sec("Capex accrual and allocation", "የወጪ ድልድል")
    asx.line("CAPEX", "Capex accrual (incl. capitalised finance)", lambda i: SUMC("ACC", i))
    asx.line("IDC", "  of which IDC (non-cash)", lambda i: SUMC("IDC", i))
    asx.line("CUMCAPEX", "Cumulative capex", lambda i: f"{ap('CUMCAPEX', i)}+{a('CAPEX', i)}", tot="last")
    asx.line("ADD_INV", "Additions – inventory (IAS 2)", lambda i: f"INV_SH*{a('CAPEX', i)}")
    asx.line("ADD_IP", "Additions – investment property (IAS 40)", lambda i: f"IP_SH*{a('CAPEX', i)}")
    asx.line("ADD_PPE", "Additions – owner-occupied (IAS 16)", lambda i: f"PPE_SH*{a('CAPEX', i)}")
    asx.sec("Accounting depreciation – components, straight line from O1", "የሂሳብ እርጅና ቅናሽ")
    dep = lambda base, life, key: (lambda i: 0 if i <= 4 else f"MIN({base}/MAX(1,{life}),{base}-{ap(key + 'C', i)})")
    asx.line("DSTR", "Structure", dep("BASE_STR", "LIFE_BLD", "DSTR"))
    asx.line("DSTRC", "  cumulative", lambda i: f"{ap('DSTRC', i)}+{a('DSTR', i)}", tot="last")
    asx.line("DSYS", "Systems (MEP, lifts)", dep("BASE_SYS", "LIFE_SYS", "DSYS"))
    asx.line("DSYSC", "  cumulative", lambda i: f"{ap('DSYSC', i)}+{a('DSYS', i)}", tot="last")
    asx.line("DFFE", "FF&E (life capped)", dep("BASE_FFE", "MIN(LIFE_FFE,FFE_LIFE_CAP)", "DFFE"))
    asx.line("DFFEC", "  cumulative", lambda i: f"{ap('DFFEC', i)}+{a('DFFE', i)}", tot="last")
    asx.line("DEP", "Total depreciation", lambda i: f"{a('DSTR', i)}+{a('DSYS', i)}+{a('DFFE', i)}", style="total")
    asx.line("DEPC", "Cumulative depreciation", lambda i: f"{ap('DEPC', i)}+{a('DEP', i)}", tot="last")
    asx.line("DEP_IP", "  allocated to investment property", lambda i: f"IF(RET_SH>0,{a('DEP', i)}*IP_SH/RET_SH,0)")
    asx.line("DEP_PPE", "  allocated to owner-occupied PPE", lambda i: f"{a('DEP', i)}-{a('DEP_IP', i)}")
    asx.sec("Carrying amounts", "የመዝገብ ዋጋ")
    asx.line("NBV_IP", "Investment property (cost model)", lambda i: f"IP_SH*{a('CUMCAPEX', i)}-{ap('NBVIPDC', i)}-{a('DEP_IP', i)}",
             tot="last")
    asx.line("NBVIPDC", "  cumulative depreciation IP", lambda i: f"{ap('NBVIPDC', i)}+{a('DEP_IP', i)}", tot="last")
    asx.line("NBV_PPE", "Owner-occupied PPE", lambda i: f"PPE_SH*{a('CUMCAPEX', i)}-{ap('NBVPPDC', i)}-{a('DEP_PPE', i)}",
             tot="last")
    asx.line("NBVPPDC", "  cumulative depreciation PPE", lambda i: f"{ap('NBVPPDC', i)}+{a('DEP_PPE', i)}", tot="last")
    asx.line("INV", "Inventory (IAS 2)", lambda i: f"INV_SH*{a('CUMCAPEX', i)}-{ap('INVC', i)}-{X('24_Revenue', 'COGS', i)}", tot="last")
    asx.line("INVC", "  cumulative cost of sales", lambda i: f"{ap('INVC', i)}+{X('24_Revenue', 'COGS', i)}", tot="last")
    asx.line("DEPBASE", "Depreciable base (excl. land)", lambda i: "BASE_STR+BASE_SYS+BASE_FFE", tot="last")
    asx.sec("IAS 36 impairment test (value in use vs carrying amount)", "የንብረት ዋጋ መቀነስ ፈተና")
    asx.line("VIU", "Value in use (PV of NOI + exit value)",
             lambda i: 0 if i <= 4 else ("TV_GROSS" if i == ANN_N else
                                         f"({a('VIU', i + 1)}+{X('24_Revenue', 'NOI', i + 1)})/(1+DISC_EFF)"),
             tot=None)
    asx.line("CA", "Carrying amount (IP + PPE)", lambda i: f"{a('NBV_IP', i)}+{a('NBV_PPE', i)}", tot=None)
    asx.line("HEAD", "Headroom", lambda i: 0 if i <= 4 else f"{a('VIU', i)}-{a('CA', i)}", tot="min")
    asx.line("IMPF", "Impairment indicator (1 = review)", lambda i: 0 if i <= 4 else f"IF(AND({a('CA', i)}>1,{a('HEAD', i)}<0),1,0)",
             fmt="int", unit="flag")

    # =================================================================== 28 TAX
    tx = Sheet(b, "28_Tax", "28 · Tax – CIT, losses, deferred tax, UPT", "ግብር – የገቢ ግብር፣ ኪሳራ፣ የዘገየ ግብር", TAB_FIN, banner=True)
    t = lambda key, i: tx.c(key, i)
    tp = lambda key, i: (tx.c(key, i - 1) if i > 1 else "0")
    PL = lambda key, i: X("40_PL", key, i)
    AS = lambda key, i: X("26_Assets", key, i)
    tx.sec("Taxable profit", "ግብር የሚከፈልበት ትርፍ")
    tx.line("PBT", "Profit before tax", lambda i: PL("PBT", i))
    tx.line("ADEP", "Add: accounting depreciation", lambda i: AS("DEP", i))
    tx.line("TDB", "Less: tax depreciation – buildings (straight line)",
            lambda i: 0 if i <= 4 else f"MIN(TAXDEP_BLD*BASE_STR,BASE_STR-{tp('TDBC', i)})")
    tx.line("TDBC", "  cumulative", lambda i: f"{tp('TDBC', i)}+{t('TDB', i)}", tot="last")
    tx.line("POOLO", "Plant pool – opening WDV", lambda i: 0 if i <= 4 else (tp("POOLC", i) if i > 5 else "BASE_SYS+BASE_FFE"),
            tot=None)
    tx.line("TDP", "Less: tax depreciation – plant pool (diminishing)", lambda i: f"TAXDEP_PLANT*{t('POOLO', i)}")
    tx.line("POOLC", "Plant pool – closing WDV", lambda i: f"{t('POOLO', i)}-{t('TDP', i)}", tot="last")
    tx.line("TDPC", "  cumulative plant allowances", lambda i: f"{tp('TDPC', i)}+{t('TDP', i)}", tot="last")
    tx.line("TI", "Taxable income before losses", lambda i: f"{t('PBT', i)}+{t('ADEP', i)}-{t('TDB', i)}-{t('TDP', i)}",
            style="total")
    tx.line("LNEW", "Tax loss arising", lambda i: f"MAX(0,-{t('TI', i)})")
    tx.sec("Loss carry-forward by vintage (FIFO, expiry after LCF_YEARS)", "ኪሳራ በዓመት")
    # remaining-at-start matrix (vintage v in rows, period p in cols)
    remkeys, usekeys = [], []
    for v in range(1, ANN_N + 1):
        k = f"REM{v}"
        remkeys.append(k)
        tx.line(k, f"Vintage {PERIOD_LABELS[v - 1]} – available at start",
                (lambda v, k: lambda i: 0 if i <= v else (f"IF({T('T_P', i)}-{T('T_P', v)}>LCF_YEARS,0,{t('LNEW', v)})" if i == v + 1
                                                           else f"IF({T('T_P', i)}-{T('T_P', v)}>LCF_YEARS,0,{tx.c(k, i - 1)}-{tx.c('USE' + str(v), i - 1)})"))(v, k),
                tot=None)
    tx.line("AVAIL", "Losses available at start", lambda i: "+".join(t(k, i) for k in remkeys), tot=None)
    tx.line("USAGE", "Losses used", lambda i: f"MIN({t('AVAIL', i)},LCF_CAP*MAX(0,{t('TI', i)}))")
    for v in range(1, ANN_N + 1):
        k = f"USE{v}"
        usekeys.append(k)
        prior = (lambda v: (lambda i: "+".join(t(f"USE{u}", i) for u in range(1, v)) if v > 1 else "0"))(v)
        tx.line(k, f"Used from vintage {PERIOD_LABELS[v - 1]}",
                (lambda v, prior: lambda i: f"MIN({t(f'REM{v}', i)},MAX(0,{t('USAGE', i)}-({prior(i)})))")(v, prior))
    tx.line("LCF", "Losses carried forward (end)", lambda i: f"{t('AVAIL', i)}-{t('USAGE', i)}+{t('LNEW', i)}", tot="last")
    tx.line("LEXP", "Losses expired", lambda i: f"MAX(0,{tp('LCF', i)}-{t('AVAIL', i)})" if i > 1 else 0)
    tx.sec("Current tax", "የዘመኑ ግብር")
    tx.line("TAXABLE", "Taxable profit after losses", lambda i: f"MAX(0,{t('TI', i)})-{t('USAGE', i)}")
    tx.line("CIT", "Corporate income tax", lambda i: f"CIT*{t('TAXABLE', i)}", style="total")
    tx.sec("Deferred tax (IAS 12)", "የዘገየ ግብር")
    tx.line("TD", "Temporary difference (tax allowances − book depreciation)",
            lambda i: f"{t('TDBC', i)}+{t('TDPC', i)}-{AS('DEPC', i)}", tot="last")
    tx.line("DTL", "Deferred tax liability on assets", lambda i: f"CIT*{t('TD', i)}", tot="last")
    tx.line("DTA", "Deferred tax asset on losses", lambda i: f"CIT*{t('LCF', i)}", tot="last")
    tx.line("DTN", "Net deferred tax liability / (asset)", lambda i: f"{t('DTL', i)}-{t('DTA', i)}", tot="last")
    tx.line("DTE", "Deferred tax expense / (credit)", lambda i: f"{t('DTN', i)}-{tp('DTN', i)}")
    tx.sec("Undistributed profits tax (closed form – no circularity)", "ያልተከፋፈለ ትርፍ ግብር")
    tx.line("PATPRE", "Profit after CIT & deferred tax", lambda i: f"{t('PBT', i)}-{t('CIT', i)}-{t('DTE', i)}")
    tx.line("UK", "Undistributed share k = (1 − reserve %)(1 − payout)", lambda i: "(1-LR_PCT)*(1-DIV_PAYOUT)", fmt="pct",
            unit="%", tot=None)
    tx.line("UPT", "UPT = r·k·PAT / (1 + r·k)",
            lambda i: f"IF(AND(PROJECT_MODE<>0,{t('PATPRE', i)}>0),UPT_R*{t('UK', i)}*{t('PATPRE', i)}/(1+UPT_R*{t('UK', i)}),0)")
    tx.sec("Tax payable and cash", "የሚከፈል ግብር")
    tx.line("CHG", "Current tax charge (CIT + UPT)", lambda i: f"{t('CIT', i)}+{t('UPT', i)}")
    tx.line("PAID", "Tax paid (following period)", lambda i: tp("CHG", i))
    tx.line("PAY", "Tax payable (end)", lambda i: f"{tp('PAY', i)}+{t('CHG', i)}-{t('PAID', i)}", tot="last")
    tx.line("WHT", "Memo: dividend WHT (borne by shareholders)", lambda i: f"DIV_WHT*{X('42_CF', 'DIV', i)}")
    tx.line("LNEG", "Check: negative loss balances", lambda i: f"IF(MIN({t('LCF', i)},{t('AVAIL', i)})<-0.5,1,0)", fmt="int",
            unit="flag")

    # =================================================================== 29 VAT
    vt = Sheet(b, "29_VAT", "29 · VAT – output, input, refunds", "የተጨማሪ እሴት ታክስ", TAB_FIN, banner=True)
    v = lambda key, i: vt.c(key, i)
    vp = lambda key, i: (vt.c(key, i - 1) if i > 1 else "0")
    vt.line("OUT_R", "Output VAT – rent & recoveries", lambda i: f"VAT_R*RENT_VAT_PCT*({RV('RENT', i)}+{RV('SCR', i)})")
    vt.line("OUT_S", "Output VAT – sales", lambda i: f"VAT_R*VAT_SALES_PCT*{RV('SALES_COLL', i)}")
    vt.line("OUT", "Total output VAT", lambda i: f"{v('OUT_R', i)}+{v('OUT_S', i)}", style="total")
    vt.line("IN_C", "Input VAT – capex", lambda i: SUMC("VATC", i))
    vt.line("IN_O", "Input VAT – opex", lambda i: f"VAT_R*{X('25_Opex', 'VATB', i)}")
    vt.line("IN", "Total input VAT", lambda i: f"{v('IN_C', i)}+{v('IN_O', i)}", style="total")
    vt.line("OPEN", "Opening VAT credit (receivable)", lambda i: vp("CLOSE", i), tot=None)
    vt.line("REF", "Refund received (prior credit)", lambda i: f"IF(VAT_REFUND_F=1,{v('OPEN', i)},0)")
    vt.line("PAYT", "VAT paid to authority", lambda i: f"MAX(0,{v('OUT', i)}-{v('IN', i)}-({v('OPEN', i)}-{v('REF', i)}))")
    vt.line("CLOSE", "Closing VAT credit (receivable)", lambda i: f"MAX(0,{v('OPEN', i)}-{v('REF', i)}+{v('IN', i)}-{v('OUT', i)})",
            tot="last", style="total")
    vt.line("CHK", "Check: roll-forward difference",
            lambda i: f"{v('OPEN', i)}-{v('REF', i)}+{v('IN', i)}-{v('OUT', i)}+{v('PAYT', i)}-{v('CLOSE', i)}", tot="sum")

    # =================================================================== 30 WORKING CAPITAL
    wc = Sheet(b, "30_WorkingCap", "30 · Working capital", "የሥራ ማስኬጃ ካፒታል", TAB_FIN, banner=True)
    w = lambda key, i: wc.c(key, i)
    wp = lambda key, i: (wc.c(key, i - 1) if i > 1 else "0")
    wc.scalar("RRP", "Retention release period index",
              "IF(RET_REL_M=0,ROUNDUP(CONSTR_EFF/12,0),4+MIN(OPS_YEARS,ROUNDUP(RET_REL_M/12,0)))", fmt="int", unit="no.",
              name="RET_REL_P")
    wc.line("AR", "Rent receivables", lambda i: f"RECV_DAYS/365*({RV('RENT', i)}-{X('25_Opex', 'BAD', i)})", tot="last")
    wc.line("AP", "Opex payables", lambda i: f"PAY_DAYS/365*({X('25_Opex', 'TOT', i)}-{X('25_Opex', 'BAD', i)})", tot="last")
    wc.line("RETW", "Retention withheld", lambda i: SUMC("RETW", i))
    wc.line("RETR", "Retention released", lambda i: f"IF({T('T_P', i)}=RET_REL_P,RET_TOT,0)")
    wc.line("RET", "Retention payable", lambda i: f"{wp('RET', i)}+{w('RETW', i)}-{w('RETR', i)}", tot="last")
    wc.line("VAT", "VAT receivable", lambda i: X("29_VAT", "CLOSE", i), tot="last")
    wc.line("CL", "Contract liability", lambda i: RV("CLIAB", i), tot="last")
    wc.line("CA", "Contract asset", lambda i: RV("CASSET", i), tot="last")
    wc.line("KEY", "Deferred key money", lambda i: RV("KEYD", i), tot="last")
    wc.line("NWC", "Net working capital (assets − liabilities)",
            lambda i: f"{w('AR', i)}+{w('VAT', i)}+{w('CA', i)}-{w('AP', i)}-{w('RET', i)}-{w('CL', i)}-{w('KEY', i)}", tot="last",
            style="total")
