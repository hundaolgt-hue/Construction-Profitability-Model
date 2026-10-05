"""Brute-force evaluation of every Excel Data Table in the workbook.

LibreOffice's MULTIPLE.OPERATIONS import is unreliable for multi-column tables, so each
table cell is computed by setting the table's input cell(s) to the row/column value,
recalculating the whole model, and reading the output formula cell(s)."""
import os
import sys
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries, get_column_letter as CL
from openpyxl.worksheet.formula import DataTableFormula

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "qa"))
from harness import run_many, values  # noqa: E402


def find_tables(path):
    wb = load_workbook(path)
    out = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, DataTableFormula):
                    d = c.value
                    c1, r1, c2, r2 = range_boundaries(d.ref)
                    t = dict(sheet=ws.title, c1=c1, r1=r1, c2=c2, r2=r2, two=bool(d.dt2D),
                             in1=d.r1, in2=d.r2)
                    if t["two"]:
                        t["rowvals"] = [ws.cell(r1 - 1, cc).value for cc in range(c1, c2 + 1)]
                        t["colvals"] = [ws.cell(rr, c1 - 1).value for rr in range(r1, r2 + 1)]
                        t["outcell"] = f"{CL(c1 - 1)}{r1 - 1}"
                    else:
                        t["colvals"] = [ws.cell(rr, c1 - 1).value for rr in range(r1, r2 + 1)]
                        t["outcells"] = [f"{CL(cc)}{r1 - 1}" for cc in range(c1, c2 + 1)]
                    out.append(t)
    return out


def compute(path, workdir, workers=4):
    tables = find_tables(path)
    combos = {}

    def key(d):
        return tuple(sorted((k, round(float(v), 9)) for k, v in d.items()))

    for t in tables:
        sh = t["sheet"]
        if t["two"]:
            for j, rv in enumerate(t["rowvals"]):
                for i, cv in enumerate(t["colvals"]):
                    combos.setdefault(key({(sh, t["in1"]): rv, (sh, t["in2"]): cv}), [])
        else:
            for i, cv in enumerate(t["colvals"]):
                combos.setdefault(key({(sh, t["in1"]): cv}), [])
    specs = []
    tags = {}
    for n, k in enumerate(sorted(combos)):
        tag = f"dt{n:03d}"
        tags[k] = tag
        specs.append((tag, dict(cells={c: v for c, v in k}, strip_dt=True)))
    print(f"data tables: {len(tables)}, distinct recalculations: {len(specs)}", flush=True)
    res = run_many(path, specs, workers=workers, keep_dir=workdir)
    cache = {}
    for k, tag in tags.items():
        r, p = res[tag]
        if "error" in r:
            raise RuntimeError(f"{tag}: {r['error']}")
        cache[k] = values(p)
    vals = {}  # (sheet, coord) -> value
    for t in tables:
        sh = t["sheet"]
        if t["two"]:
            for j, rv in enumerate(t["rowvals"]):
                for i, cv in enumerate(t["colvals"]):
                    wb = cache[key({(sh, t["in1"]): rv, (sh, t["in2"]): cv})]
                    vals[(sh, f"{CL(t['c1'] + j)}{t['r1'] + i}")] = wb[sh][t["outcell"]].value
        else:
            for i, cv in enumerate(t["colvals"]):
                wb = cache[key({(sh, t["in1"]): cv})]
                for j, oc in enumerate(t["outcells"]):
                    vals[(sh, f"{CL(t['c1'] + j)}{t['r1'] + i}")] = wb[sh][oc].value
    return tables, vals
