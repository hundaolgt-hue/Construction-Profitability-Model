"""Core builder framework for the Rapha Development Profitability Model.

Rows are *reserved* when declared and their formulas are generated lazily at
render time, so any sheet can reference rows on sheets declared later.
"""
from openpyxl import Workbook
from openpyxl.utils import get_column_letter as CL
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.formatting.rule import FormulaRule

# ---------------------------------------------------------------- palette
NAVY, AMBER = "0F2A44", "D98E2B"
BLUE, TEAL, VIOLET, ROSE = "1F5FAD", "2A9D8F", "7A5195", "C44E52"
CANVAS, CARD, BORDER, INK, SEC = "F4F6F8", "FFFFFF", "D9DEE4", "1B2430", "5B6675"
GOOD_FG, GOOD_BG = "2E7D4F", "E3F2E8"
WARN_FG, WARN_BG = "B7791F", "FDF1DC"
CRIT_FG, CRIT_BG = "B42318", "FDE7E5"
INPUT_FG, INPUT_BG = "1F4E9C", "FFF8E6"
LINK_FG = "1E7B34"
KEY_BG = "FFF2B3"
FONT = "Segoe UI"
AMFONT = "Nyala"

NF = {
    "n0": '#,##0;(#,##0);"–"',
    "n1": '#,##0.0;(#,##0.0);"–"',
    "n2": '#,##0.00;(#,##0.00);"–"',
    "pct": '0.0%;(0.0%);"–"',
    "pct2": '0.00%;(0.00%);"–"',
    "x": '0.00"x";(0.00"x");"–"',
    "date": "dd-mmm-yyyy;;\"–\"",
    "mon": "mmm-yy",
    "int": "0",
    "txt": "@",
    "gen": "General",
    "k": '#,##0,;(#,##0,);"–"',
    "yr": '0.0" yrs";(0.0" yrs");"–"',
    "mo": '0" mo";(0" mo");"–"',
    "m": '#,##0.0,,"m";(#,##0.0,,"m");"–"',
}

thin = Side(style="thin", color=BORDER)
hair = Side(style="hair", color=BORDER)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
TOPLINE = Border(top=Side(style="thin", color=INK))
DTOP = Border(top=Side(style="thin", color=INK), bottom=Side(style="double", color=INK))
BOTTOM_AMBER = Border(bottom=Side(style="medium", color=AMBER))


def axis_font(ax, size=700, color="5B6675"):
    """set tick-label font size (hundredths of a point) on a chart axis"""
    from openpyxl.chart.text import RichText
    from openpyxl.drawing.text import Paragraph, ParagraphProperties, CharacterProperties, RichTextProperties
    cp = CharacterProperties(sz=size, solidFill=color)
    ax.txPr = RichText(bodyPr=RichTextProperties(), p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def fill(c):
    return PatternFill("solid", start_color=c, end_color=c)


def font(size=10, bold=False, color=INK, italic=False, name=FONT, underline=None):
    return Font(name=name, size=size, bold=bold, color=color, italic=italic, underline=underline)


F_BODY = font()
F_BOLD = font(bold=True)
F_INPUT = font(color=INPUT_FG)
F_LINK = font(color=LINK_FG)
F_NOTE = font(9, color=SEC, italic=True)
F_UNIT = font(9, color=SEC)
F_AM = font(10, color=SEC, name=AMFONT)
FILL_INPUT = fill(INPUT_BG)
FILL_KEY = fill(KEY_BG)
FILL_NAVY = fill(NAVY)
FILL_AMBER = fill(AMBER)
FILL_CANVAS = fill(CANVAS)
FILL_CARD = fill(CARD)

ANN_N = 14          # C1..C4 + O1..O10
MON_C = 48          # construction months (max)
MON_L = 288         # loan / lifetime monthly timeline (48 + 240)
FIRST_COL = 5       # column E is period 1 / month 1


def pc(p):
    """column letter of annual period p (1..14) or month p."""
    return CL(FIRST_COL + p - 1)


PERIOD_LABELS = [f"C{i}" for i in range(1, 5)] + [f"O{i}" for i in range(1, 11)]


class Row:
    __slots__ = ("sheet", "key", "r", "label", "f", "fmt", "unit", "tot", "style", "n", "am", "note")

    def __init__(self, **kw):
        for k in self.__slots__:
            setattr(self, k, kw.get(k))


class Sheet:
    def __init__(self, b, name, title, am, tab, kind="annual", n=None, banner=False, width_label=46):
        self.b, self.name, self.title, self.am, self.kind = b, name, title, am, kind
        self.ws = b.wb.create_sheet(name)
        self.ws.sheet_properties.tabColor = tab
        self.ws.sheet_view.showGridLines = False
        self.rows = {}
        self.scal = {}
        self.items = []      # deferred renders
        self.banner = banner
        self.n = n if n is not None else (ANN_N if kind == "annual" else (MON_C if kind == "m48" else (MON_L if kind == "m288" else 0)))
        self.r = 10
        self.width_label = width_label
        b.sheets[name] = self

    # ------------------------------------------------------------ refs
    @property
    def q(self):
        return f"'{self.name}'!"

    def c(self, key, i, absr=False, sheet=False):
        r = self.rows[key]
        s = self.q if sheet else ""
        return f"{s}{pc(i)}{'$' if absr else ''}{r}"

    def rng(self, key, a=1, z=None, sheet=True, absc=True):
        z = z or self.n
        r = self.rows[key]
        s = self.q if sheet else ""
        d = "$" if absc else ""
        return f"{s}{d}{pc(a)}${r}:{d}{pc(z)}${r}"

    def s(self, key, sheet=True):
        r = self.scal[key]
        return f"{self.q if sheet else ''}$D${r}"

    def tot(self, key, sheet=True):
        r = self.rows[key]
        return f"{self.q if sheet else ''}$D${r}"

    # ------------------------------------------------------------ layout
    def gap(self, k=1):
        self.r += k

    def sec(self, title, am=None):
        r = self.r
        self.items.append(("sec", r, title, am))
        self.r += 1
        return r

    def raw(self, fn, nrows=0):
        """fn(r0) is called at render; reserves nrows rows starting at current row."""
        r0 = self.r
        self.items.append(("raw", fn, r0))
        self.r += nrows
        return r0

    def note(self, text):
        self.items.append(("note", self.r, text))
        self.r += 1

    def line(self, key, label, f, fmt="n0", unit="ETB", tot="sum", style=None, am=None, note=None, n=None):
        """f: callable(i) -> formula string (no '=') | number | None"""
        assert key not in self.rows, (self.name, key)
        row = Row(sheet=self, key=key, r=self.r, label=label, f=f, fmt=fmt, unit=unit, tot=tot,
                  style=style, n=n or self.n, am=am, note=note)
        self.rows[key] = self.r
        self.items.append(("line", row))
        self.r += 1
        return row

    def scalar(self, key, label, f, fmt="n0", unit="ETB", style=None, name=None, note=None, am=None):
        assert key not in self.scal, (self.name, key)
        self.scal[key] = self.r
        self.items.append(("scalar", self.r, key, label, f, fmt, unit, style, name, note, am))
        self.r += 1
        return self.r - 1

    # ------------------------------------------------------------ render
    def render(self):
        ws = self.ws
        b = self.b
        for it in self.items:
            kind = it[0]
            if kind == "sec":
                _, r, title, am = it
                c = ws.cell(r, 2, title)
                c.font = font(11, True, NAVY)
                c.border = BOTTOM_AMBER
                if am:
                    a = ws.cell(r, 3, am)
                    a.font = font(10, False, SEC, name=AMFONT)
                last = 4 + max(self.n, 1) if self.kind != "free" else 10
                for cc in range(3 if not am else 4, min(last, 4 + 14) + 1):
                    ws.cell(r, cc).border = BOTTOM_AMBER
                ws.row_dimensions[r].height = 20
            elif kind == "note":
                _, r, text = it
                c = ws.cell(r, 2, text)
                c.font = F_NOTE
            elif kind == "raw":
                it[1](it[2])
            elif kind == "line":
                self._render_line(it[1])
            elif kind == "scalar":
                _, r, key, label, f, fmt, unit, style, name, note, am = it
                ws.cell(r, 2, label).font = F_BODY
                if am:
                    ws.cell(r, 2).value = f"{label}  ·  {am}"
                ws.cell(r, 3, unit).font = F_UNIT
                v = f() if callable(f) else f
                cell = ws.cell(r, 4, ("=" + v) if isinstance(v, str) else v)
                cell.number_format = NF.get(fmt, fmt)
                b.style_cell(cell, v, style)
                if note:
                    ws.cell(r, 5, note).font = F_NOTE
                if name:
                    b.name(name, f"{self.q}$D${r}")

    def _render_line(self, row):
        ws, b = self.ws, self.b
        r = row.r
        lab = ws.cell(r, 2, row.label)
        lab.font = F_BOLD if row.style in ("total", "key", "grand") else F_BODY
        if row.style in ("total", "grand"):
            lab.font = F_BOLD
        if row.am:
            lab.value = f"{row.label}  ·  {row.am}"
        ws.cell(r, 3, row.unit).font = F_UNIT
        n = row.n
        for i in range(1, n + 1):
            v = row.f(i) if callable(row.f) else row.f
            if v is None:
                continue
            if isinstance(v, str) and v.startswith("§"):
                cell = ws.cell(r, 4 + i, v[1:])
                cell.font = font(9, True, NAVY)
                cell.alignment = Alignment(horizontal="right")
                continue
            cell = ws.cell(r, 4 + i, ("=" + v) if isinstance(v, str) else v)
            cell.number_format = NF.get(row.fmt, row.fmt)
            b.style_cell(cell, v, row.style)
        a, z = pc(1), pc(n)
        t = row.tot
        tf = None
        if t == "sum":
            tf = f"SUM({a}{r}:{z}{r})"
        elif t == "max":
            tf = f"MAX({a}{r}:{z}{r})"
        elif t == "min":
            tf = f"MIN({a}{r}:{z}{r})"
        elif t == "last":
            tf = f"{z}{r}"
        elif t == "first":
            tf = f"{a}{r}"
        elif isinstance(t, str) and t.startswith("="):
            tf = t[1:]
        if tf:
            cell = ws.cell(r, 4, "=" + tf)
            cell.number_format = NF.get(row.fmt, row.fmt)
            cell.font = F_BOLD
            cell.fill = fill("EEF1F5")
        if row.style in ("total", "grand"):
            for cc in range(2, 5 + n - 1 + 1):
                ws.cell(r, cc).border = DTOP if row.style == "grand" else TOPLINE
        if row.note:
            pass


import re as _re
_LINK_RE = _re.compile(r"^(?:'[^']+'!\$?[A-Z]{1,3}\$?\d+|[A-Z][A-Z0-9_]*)$")


class Builder:
    def __init__(self):
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.sheets = {}
        self.names = {}

    def name(self, nm, ref):
        assert nm not in self.names, nm
        self.names[nm] = ref

    def sh(self, name):
        return self.sheets[name]

    def style_cell(self, cell, v, style):
        if style == "input":
            cell.font = F_INPUT
            cell.fill = FILL_INPUT
            cell.border = BOX
            return
        is_formula = isinstance(v, str)
        link = is_formula and bool(_LINK_RE.match(v.strip()))
        bold = style in ("total", "grand", "key")
        color = LINK_FG if link else INK
        cell.font = Font(name=FONT, size=10, bold=bold, color=color)
        if style == "key":
            cell.fill = FILL_KEY

    def finalize_names(self):
        for nm, ref in self.names.items():
            self.wb.defined_names[nm] = DefinedName(nm, attr_text=ref)


# ------------------------------------------------------------------ header
def header(b, sh, ncols_vis=None, start_link=True, period_hdr=True):
    ws = sh.ws
    last = ncols_vis or (4 + min(max(sh.n, 1), 14))
    if sh.kind == "free":
        last = ncols_vis or 12
    for r in (1, 2, 3):
        for c in range(1, last + 1):
            ws.cell(r, c).fill = FILL_NAVY
    for c in range(1, last + 1):
        ws.cell(4, c).fill = FILL_AMBER
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 18
    ws.row_dimensions[4].height = 4
    t = ws.cell(1, 2, sh.title)
    t.font = font(20, True, "FFFFFF")
    t.alignment = Alignment(vertical="center")
    a = ws.cell(2, 2, sh.am)
    a.font = font(13, False, "F3D7AE", name=AMFONT)
    p = ws.cell(3, 2, '=PROJ_NAME&"   |   Scenario: "&SCEN_NAME&"   |   Mode: "&MODE_NAME')
    p.font = font(9, False, "D6DEE8")
    if start_link:
        lc = ws.cell(1, last, '=HYPERLINK("#\'00_Start\'!A1","◂ Start")')
        lc.font = font(10, True, "FFFFFF", underline="single")
        lc.alignment = Alignment(horizontal="right", vertical="center")
    if sh.banner:
        bc = ws.cell(5, 2, '=IF(PROJECT_MODE=0,"Not needed for personal projects — switch mode to Commercial (1–3) on the Start sheet to see profitability.  ·  ለግል ፕሮጀክት አያስፈልግም — ትርፋማነትን ለማየት በመጀመሪያ ገጽ ላይ ዓይነቱን ወደ ንግድ (1–3) ይቀይሩ።","")')
        bc.font = font(10, True, CRIT_FG)
        rng = f"B5:{CL(last)}5"
        ws.conditional_formatting.add(rng, FormulaRule(formula=["PROJECT_MODE=0"], fill=fill(CRIT_BG)))
    # column widths
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = sh.width_label
    if sh.kind == "free":
        ws.column_dimensions["D"].width = 18
    if sh.kind == "free" and sh.n:
        ws.column_dimensions["C"].width = 9
        ws.column_dimensions["D"].width = 15
        for i in range(1, sh.n + 1):
            ws.column_dimensions[pc(i)].width = 12.5
    if sh.kind != "free":
        ws.column_dimensions["C"].width = 9
        ws.column_dimensions["D"].width = 15
        w = 12.5 if sh.kind == "annual" else 11.5
        for i in range(1, sh.n + 1):
            ws.column_dimensions[pc(i)].width = w
    if period_hdr and sh.kind == "annual":
        hdr_cells = [(6, "T_LABEL", None, True), (7, "T_END", "mon", False), (8, "T_EC", None, False)]
        ws.cell(6, 2, "Period").font = font(9, True, SEC)
        ws.cell(7, 2, "Period end (Gregorian)").font = font(9, False, SEC)
        ws.cell(8, 2, "Ethiopian calendar year").font = font(9, False, SEC)
        ws.cell(6, 4, "Total").font = font(9, True, NAVY)
        ws.cell(6, 4).alignment = Alignment(horizontal="right")
        for r, key, fmt, bold in hdr_cells:
            for i in range(1, sh.n + 1):
                v = f"='20_Timeline'!{pc(i)}{b.TL[key]}" if sh.name != "20_Timeline" else f"={pc(i)}{b.TL[key]}"
                c = ws.cell(r, 4 + i, v)
                c.font = font(9, bold, NAVY if bold else SEC)
                c.alignment = Alignment(horizontal="right")
                if fmt:
                    c.number_format = NF[fmt]
                if bold:
                    c.border = Border(bottom=Side(style="thin", color=NAVY))
        ws.freeze_panes = "E9"
    elif period_hdr and sh.kind in ("m48", "m288"):
        ws.cell(6, 2, "Month").font = font(9, True, SEC)
        ws.cell(7, 2, "Month start").font = font(9, False, SEC)
        ws.cell(8, 2, "Annual period").font = font(9, False, SEC)
        ws.cell(6, 4, "Total").font = font(9, True, NAVY)
        for i in range(1, sh.n + 1):
            pre = "" if sh.name == "20_Timeline" else "'20_Timeline'!"
            c = ws.cell(6, 4 + i, f'="M"&{pre}{pc(i)}{b.TL["M_T"]}')
            c.font = font(9, True, NAVY)
            c.alignment = Alignment(horizontal="right")
            c.border = Border(bottom=Side(style="thin", color=NAVY))
            c2 = ws.cell(7, 4 + i, f"={pre}{pc(i)}{b.TL['M_DATE']}")
            c2.number_format = NF["mon"]
            c2.font = font(9, False, SEC)
            c3 = ws.cell(8, 4 + i, f'=IF({pre}{pc(i)}{b.TL["M_PIDX"]}=0,"–",INDEX({pre}$E${b.TL["T_LABEL"]}:$R${b.TL["T_LABEL"]},{pre}{pc(i)}{b.TL["M_PIDX"]}))')
            c3.font = font(9, False, SEC)
            c3.alignment = Alignment(horizontal="right")
        ws.freeze_panes = "E9"
    else:
        ws.freeze_panes = "A6"
    # print setup
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.6
    ws.oddHeader.left.text = "&\"Segoe UI,Bold\"Rapha Development Profitability Model"
    ws.oddHeader.right.text = "&A"
    ws.oddFooter.left.text = "Planning tool – not investment, tax or legal advice"
    ws.oddFooter.right.text = "Page &P of &N"
    if sh.kind == "annual":
        ws.print_title_rows = "1:8"
    if sh.kind in ("m48", "m288"):
        ws.print_title_rows = "1:8"
        ws.print_title_cols = "A:D"
        ws.page_setup.fitToWidth = 4 if sh.kind == "m48" else 1
        ws.page_setup.fitToHeight = 0
