"""Full production pipeline for one deliverable workbook.

1. build the workbook with openpyxl (formulas, charts, styles, static links)
2. recalculate a copy with LibreOffice
3. inject every computed value as the cached result into the built file (formulas and charts are kept), so the file
   shows numbers immediately in Excel, LibreOffice, Google Sheets and file previews
4. scan the result for errors and print the check summary

usage: python finalize.py --out ../output/X.xlsx [--blank] [--track N --pdata N --mode N --data N --scen N]
"""
import argparse
import json
import os
import sys
import tempfile

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(HERE, "qa"))
import main as builder  # noqa: E402
import postprocess  # noqa: E402
from harness import make_variant, calc, values  # noqa: E402

ERRS = ("#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#NULL!", "#NUM!", "#N/A")


def run(out, blank=False, track=None, pdata=None, mode=None, data=None, scen=None, inputs=None, work=None):
    work = work or tempfile.mkdtemp(prefix="rapha_fin_")
    os.makedirs(work, exist_ok=True)
    built = os.path.join(work, "built.xlsx")
    b = builder.build(blank)
    b.wb.save(built)
    src = built
    if any(v is not None for v in (track, pdata, mode, data, scen)) or inputs:
        src = os.path.join(work, "built_settings.xlsx")
        set_start(built, src, track, pdata, mode, data, scen, inputs)
    calc_copy = os.path.join(work, "calc.xlsx")
    make_variant(src, calc_copy)
    r = calc(calc_copy, 900)
    if "error" in r:
        raise RuntimeError(r["error"])
    wbv = values(calc_copy)
    vals = {}
    for ws in wbv.worksheets:
        d = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    d[c.coordinate] = c.value
        vals[ws.title] = d
    stats = postprocess.inject(src, out, vals, [])
    errs = []
    for sh, d in vals.items():
        for co, v in d.items():
            if isinstance(v, str) and v in ERRS:
                errs.append(f"{sh}!{co}={v}")
    chk = {k: nm(wbv, k) for k in ("TRACK", "CHK_PASS", "CHK_FAIL", "CHK_WARN", "CHK_TOTAL", "VERDICT", "EIRR", "TDC",
                                   "PCHK_PASS", "PCHK_FAIL", "PCHK_WARN", "P_TOTAL")}
    rep = {"out": out, "recalc_total_formulas": r.get("total_formulas"), "recalc_errors": r.get("total_errors"),
           "value_errors": len(errs), "error_cells": errs[:20], "injected": stats, "checks": chk}
    print(json.dumps(rep, indent=1, default=str))
    return rep


def model_hash(path):
    """hash of everything that can change a value: all formulas, numbers and the input-sheet text
    (labels, styles and charts are excluded so cosmetic edits do not invalidate the data-table cache)"""
    import hashlib
    from openpyxl.worksheet.formula import DataTableFormula
    h = hashlib.sha256()
    wb = load_workbook(path)
    text_sheets = {"00_Start", "10_Inputs", "11_Scenarios", "Lists", "52_Sensitivity"}
    for ws in wb.worksheets:
        h.update(ws.title.encode())
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v is None:
                    continue
                if isinstance(v, DataTableFormula):
                    v = f"DT{v.ref}{v.r1}{v.r2}{v.dt2D}"
                elif isinstance(v, str) and not v.startswith("=") and ws.title not in text_sheets:
                    continue
                h.update(f"{c.coordinate}={v}|".encode())
    for nm, dn in sorted(wb.defined_names.items()):
        h.update(f"{nm}:{dn.attr_text}".encode())
    return h.hexdigest()


def render_copy(src, dst, strip_hf=False):
    """copy without fullCalcOnLoad so LibreOffice shows the injected cached values"""
    import re
    import zipfile
    zin = zipfile.ZipFile(src)
    zout = zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED)
    for item in zin.infolist():
        d = zin.read(item.filename)
        if item.filename == "xl/workbook.xml":
            d = re.sub(rb'\s*fullCalcOnLoad="1"', b"", d)
        if item.filename.startswith("xl/worksheets/sheet"):
            d = re.sub(rb'<f t="dataTable"[^>]*/>', b"", d)
        if strip_hf and item.filename.startswith("xl/worksheets/sheet"):
            d = re.sub(rb"<headerFooter>.*?</headerFooter>", b"", d, flags=re.S)
        zout.writestr(item, d)
    zout.close()


def nm(wb, n):
    d = wb.defined_names[n]
    sh, ref = list(d.destinations)[0]
    return wb[sh][ref.replace("$", "").split(":")[0]].value


def set_start(src, dst, track=None, pdata=None, mode=None, data=None, scen=None, inputs=None):
    """Change Start selections / inputs by editing the sheet XML directly (keeps charts)."""
    import re
    import zipfile
    from harness import MODE, DATA, SCEN, PDATA, TRACK
    from xml.sax.saxutils import escape
    z = zipfile.ZipFile(src)
    files = postprocess._sheet_files(z)
    edits = {}
    for v, cell, mp in ((track, "D9", TRACK), (pdata, "D15", PDATA), (mode, "I15", MODE), (data, "I17", DATA), (scen, "I19", SCEN)):
        if v is not None:
            edits.setdefault("00_Start", {})[cell] = mp[v]
    if inputs:
        wb = load_workbook(src, read_only=False)
        for k, v in inputs.items():
            dn = wb.defined_names[k]
            sh, ref = list(dn.destinations)[0]
            row = int("".join(ch for ch in ref if ch.isdigit()))
            col = {1: "E", 2: "F", 3: "G"}[data or 2]
            edits.setdefault(sh, {})[f"{col}{row}"] = v
    zout = zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED)
    inv = {v: k for k, v in files.items()}
    for item in z.infolist():
        data_ = z.read(item.filename)
        sh = inv.get(item.filename)
        if sh in edits:
            xml = data_.decode()
            for co, v in edits[sh].items():
                if isinstance(v, str):
                    new = f'<c r="{co}"\\g<1> t="inlineStr"><is><t>{escape(v)}</t></is></c>'
                else:
                    new = f'<c r="{co}"\\g<1> t="n"><v>{v}</v></c>'
                pat = re.compile(rf'<c r="{co}"((?:\s+s="\d+")?)(?:\s+t="[^"]*")?\s*(?:/>|>.*?</c>)')
                xml, n = pat.subn(new, xml, count=1)
                assert n == 1, (sh, co)
            data_ = xml.encode()
        zout.writestr(item, data_)
    zout.close()


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", required=True)
    ap.add_argument("--blank", action="store_true")
    ap.add_argument("--track", type=int)
    ap.add_argument("--pdata", type=int)
    ap.add_argument("--mode", type=int)
    ap.add_argument("--data", type=int)
    ap.add_argument("--scen", type=int)
    ap.add_argument("--inputs", default=None, help="json dict of input overrides")
    ap.add_argument("--work", default=None)
    a = ap.parse_args()
    run(a.out, a.blank, a.track, a.pdata, a.mode, a.data, a.scen, json.loads(a.inputs) if a.inputs else None, a.work)
