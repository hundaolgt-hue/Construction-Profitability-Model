import re, sys
from openpyxl import load_workbook
wb = load_workbook(sys.argv[1])
names = set(wb.defined_names.keys())
funcs = set("""SUM SUMIFS SUMIF SUMPRODUCT IF AND OR NOT MAX MIN INDEX MATCH CHOOSE IFERROR ROUNDUP ROUND EDATE YEAR DATE PMT MOD INT
ABS COUNT COUNTIF AVERAGE ISNUMBER ISERROR XIRR IRR LARGE LN TEXT HYPERLINK ROWS COLUMN VALUE LEFT TRUE FALSE""".split())
bad = {}
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if isinstance(v, str) and v.startswith("="):
                s = re.sub(r'"[^"]*"', '', v)
                s = re.sub(r"'[^']*'!", '', s)
                for tok in re.findall(r"(?<![A-Za-z0-9_$.])([A-Za-z_][A-Za-z0-9_.]*)(?![A-Za-z0-9_]*!)", s):
                    if tok.upper() in funcs: continue
                    if re.fullmatch(r"\$?[A-Z]{1,3}\$?\d+", tok): continue
                    if re.fullmatch(r"[A-Z]{1,3}", tok): continue
                    if tok in names: continue
                    bad.setdefault(tok, []).append(f"{ws.title}!{c.coordinate}")
for k, v in sorted(bad.items()):
    print(k, len(v), v[:3])
