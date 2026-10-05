import sys
from openpyxl import load_workbook
v = load_workbook(sys.argv[1], data_only=True)
f = load_workbook(sys.argv[2] if len(sys.argv) > 2 else sys.argv[1].replace("_calc", ""))
n = int(sys.argv[3]) if len(sys.argv) > 3 else 4
for ws in v.worksheets:
    errs = []
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("#"):
                errs.append(c.coordinate)
    if errs:
        print(f"== {ws.title}: {len(errs)}")
        for co in errs[:n]:
            print("   ", co, v[ws.title][co].value, "|", str(f[ws.title][co].value)[:220])
