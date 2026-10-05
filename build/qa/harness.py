"""QA harness: build variants of the model (switch mode / data set / scenario / inputs),
recalculate with LibreOffice in parallel and read results."""
import glob
import os
import shutil
import sys
import tempfile
from concurrent.futures import ThreadPoolExecutor

from openpyxl import load_workbook
from openpyxl.worksheet.formula import DataTableFormula

sys.path.insert(0, glob.glob("/root/.claude/skills/synced/*/xlsx/scripts")[0])
from recalc import recalc  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
from inputs_def import LISTS  # noqa: E402

MODE = {m: LISTS["L_MODE"][m] for m in range(4)}
DATA = {d: LISTS["L_DATASET"][d - 1] for d in (1, 2, 3)}
SCEN = {s: LISTS["L_SCEN"][s - 1] for s in (1, 2, 3, 4)}


def input_cell(wb, name, col="E"):
    d = wb.defined_names[name]
    sh, ref = list(d.destinations)[0]
    row = int("".join(ch for ch in ref if ch.isdigit()))
    return sh, f"{col}{row}"


def make_variant(src, out, mode=None, data=None, scen=None, inputs=None, flex=None, strip_dt=True, col=None, cells=None,
                 paste=None):
    wb = load_workbook(src)
    st = wb["00_Start"]
    if mode is not None:
        st["D10"] = MODE[mode]
    if data is not None:
        st["D12"] = DATA[data]
    if scen is not None:
        st["D14"] = SCEN[scen]
    dcol = col or {1: "E", 2: "F", 3: "G"}.get(data or 3, "G")
    for k, v in (inputs or {}).items():
        sh, c = input_cell(wb, k, dcol)
        wb[sh][c] = v
    for k, v in (flex or {}).items():
        d = wb.defined_names[k]
        sh, ref = list(d.destinations)[0]
        wb[sh][ref.replace("$", "")] = v
    for (sh, co), v in (cells or {}).items():
        wb[sh][co] = v
    for (sh, co), v in (paste or {}).items():
        wb[sh][co] = v
    if strip_dt:
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, DataTableFormula):
                        c.value = None
    wb.save(out)
    return out


def calc(path, timeout=600):
    r = recalc(path, timeout)
    return r


def values(path):
    return load_workbook(path, data_only=True)


def nm(wb, n):
    d = wb.defined_names[n]
    sh, ref = list(d.destinations)[0]
    return wb[sh][ref.replace("$", "").split(":")[0]].value


def run_many(src, specs, workers=4, keep_dir=None):
    """specs: list of (tag, kwargs). returns {tag: (recalc_result, path)}"""
    tmp = keep_dir or tempfile.mkdtemp(prefix="rapha_qa_")
    os.makedirs(tmp, exist_ok=True)

    def one(spec):
        tag, kw = spec
        p = os.path.join(tmp, f"{tag}.xlsx")
        make_variant(src, p, **kw)
        r = calc(p)
        return tag, (r, p)

    out = {}
    with ThreadPoolExecutor(workers) as ex:
        for tag, res in ex.map(one, specs):
            out[tag] = res
    return out
