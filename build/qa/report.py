import sys
from openpyxl import load_workbook
wb = load_workbook(sys.argv[1], data_only=True)
def nm(n):
    d = wb.defined_names[n]
    sh, ref = list(d.destinations)[0]
    c = wb[sh][ref.replace("$", "").split(":")[0]]
    return c.value
ws = wb["60_Checks"]
for r in range(9, 40):
    if ws.cell(r, 6).value in ("PASS", "FAIL", "WARN"):
        print(f"{ws.cell(r,6).value:5} {ws.cell(r,2).value[:70]:70} {ws.cell(r,4).value}")
keys = ["PROJECT_MODE","DATA_SET","SCEN","CONSTR_EFF","GFA","NET_AREA","SALE_AREA","LET_AREA","HARD_BASE","HARD_ESC_TOT","CONSTR_COST","LAND_TOT","IDC_TOT","FIN_CAP","TDC","HOME_TOTAL",
        "EQ_COMMIT","DEBT_CAP","EQ_DRAWN","DEBT_DRAWN","FUND_GAP","PS_USED","PEAK_DEBT","PEAK_FUND","LN_INSTAL","LN_INSTAL_M","LN_BAL_TENOR",
        "GDV","POC","PNPV","PIRR","ENPV","EIRR","EMOIC","MIN_DSCR","AVG_DSCR","LLCR_O1","MIN_ICR","MAX_GEAR","TV_NET","EQ_EXIT","VERDICT",
        "PB_P","PB_PD","PB_E","PB_ED","CHK_PASS","CHK_FAIL","CHK_WARN","AFFORD_PCT","CASH_MIN_OPS","N_IMPAIR","BLEND_PRICE",
        "BE_PRICE_NPV","BE_PRICE_POC","BE_OCC","BE_COST","BE_RATE","BE_DELAY"]
for k in keys:
    try:
        v = nm(k)
    except Exception as e:
        v = f"ERR {e}"
    print(f"{k:14} {v:,.4f}" if isinstance(v, (int, float)) else f"{k:14} {v}")
