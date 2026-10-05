"""Independently recompute a sample of Data Table cells (fresh model recalcs) and compare with the shipped values."""
import sys
from openpyxl import load_workbook
from harness import run_many, values, nm
deliv = sys.argv[1]
wb = load_workbook(deliv, data_only=True)
se = wb["52_Sensitivity"]
def find(label):
    for r in range(1, se.max_row + 1):
        if str(se.cell(r, 2).value or "").startswith(label):
            return r
cases = []
r = find("One-way: Sale price"); cases.append(("price0.9", {"FLX_PRICE": 0.9}, [("EIRR", se.cell(r + 5, 4).value), ("MIN_DSCR", se.cell(r + 5, 5).value), ("POC", se.cell(r + 5, 6).value)]))
r = find("One-way: Construction cost"); cases.append(("cost1.2", {"FLX_COST": 1.2}, [("EIRR", se.cell(r + 8, 4).value), ("POC", se.cell(r + 8, 6).value)]))
r = find("One-way: construction delay"); cases.append(("delay6", {"FLX_DELAY": 6}, [("EIRR", se.cell(r + 5, 4).value)]))
r = find("Two-way"); cases.append(("p1.1c0.9", {"FLX_PRICE": 1.1, "FLX_COST": 0.9}, [("EIRR", se.cell(r + 3, 7).value)]))
r = find("Scenario grid"); cases.append(("scen3", {"SCEN_OVR": 3}, [("EIRR", se.cell(r + 5, 4).value), ("MIN_DSCR", se.cell(r + 5, 7).value)]))
res = run_many(deliv, [(t, dict(flex=f)) for t, f, _ in cases], workers=4, keep_dir="/tmp/claude-0/qa_dtspot")
bad = 0
for t, f, exp in cases:
    w = values(res[t][1])
    for k, v in exp:
        got = nm(w, k)
        ok = abs(got - v) < 1e-6
        bad += not ok
        print(f"{t:10} {k:9} table={v:.6f} fresh={got:.6f} {'OK' if ok else 'BAD'}")
print("DATA TABLE SPOT CHECK", "PASSED" if bad == 0 else f"FAILED {bad}")
