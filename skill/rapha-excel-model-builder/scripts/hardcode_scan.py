"""List numeric literals inside formulas (excluding structural constants)."""
import re, sys
from collections import Counter
from openpyxl import load_workbook
ALLOWED = {"0", "1", "2", "3", "4", "5", "12", "100", "365", "1000", "1000000", "0.5", "999", "99", "9", "11", "8", "10", "-1",
           "0.0001", "0.001", "1.25", "48", "288", "20", "0.2", "0.15", "-0.05", "0.05"}
wb = load_workbook(sys.argv[1])
cnt = Counter(); where = {}
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if isinstance(v, str) and v.startswith("="):
                s = re.sub(r'"[^"]*"', '', v)
                s = re.sub(r"'[^']*'!", '', s)
                s = re.sub(r"\$?[A-Z]{1,3}\$?\d+", "", s)
                s = re.sub(r"[A-Za-z_][A-Za-z0-9_.]*", "", s)
                for n in re.findall(r"(?<![\w.])\d+(?:\.\d+)?", s):
                    if n not in ALLOWED:
                        cnt[n] += 1
                        where.setdefault(n, f"{ws.title}!{c.coordinate}: {v[:120]}")
for n, k in cnt.most_common():
    print(f"{n:>10} x{k:<5} e.g. {where[n]}")
