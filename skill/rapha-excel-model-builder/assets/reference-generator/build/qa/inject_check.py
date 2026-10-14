"""Verify every cached value in the deliverable equals the LibreOffice/brute-force value for that same cell."""
import sys, json
from openpyxl import load_workbook
deliv, pasted = sys.argv[1], sys.argv[2]
a = load_workbook(deliv, data_only=True); b = load_workbook(pasted, data_only=True)
f = load_workbook(deliv)
bad = 0; n = 0
for ws in f.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                n += 1
                va, vb = a[ws.title][c.coordinate].value, b[ws.title][c.coordinate].value
                if isinstance(va, float) and isinstance(vb, (int, float)):
                    ok = abs(va - vb) <= 1e-6 * max(1, abs(vb))
                else:
                    ok = (va == vb) or (va in (None, "") and vb in (None, ""))
                if not ok:
                    bad += 1
                    if bad < 10: print("MISMATCH", ws.title, c.coordinate, va, vb)
print(f"formula cells compared: {n}, mismatches: {bad}")
