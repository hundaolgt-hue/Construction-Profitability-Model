"""Statements (40–43), analysis (50–52) and checks (60)."""
from openpyxl.styles import Alignment, Font
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from core import (Sheet, pc, CL, NF, ANN_N, MON_C, MON_L, font, fill, F_BODY, F_BOLD, F_UNIT, F_NOTE, FILL_INPUT, F_INPUT,
                  BOX, NAVY, SEC, INK, AMBER, GOOD_FG, GOOD_BG, WARN_FG, WARN_BG, CRIT_FG, CRIT_BG, FILL_KEY,
                  PERIOD_LABELS, BOTTOM_AMBER)

TAB_ST = "0F2A44"
TAB_AN = "7A5195"
TAB_CTL = "B42318"


def Y(i):
    return i - 4 if i > 4 else 0


def build_fin(b):
    X = lambda sh, key, i: b.sheets[sh].c(key, i, sheet=True)
    XS = lambda sh, key: b.sheets[sh].s(key)
    XT = lambda sh, key: b.sheets[sh].tot(key)
    XR = lambda sh, key, a=1, z=None: b.sheets[sh].rng(key, a, z)
    T = lambda key, i: f"'20_Timeline'!{pc(i)}{b.TL[key]}"
    C23 = "23_Construction"
    cmr = b.sheets[C23].rows
    SUMC = lambda key, i: f"SUMIFS('{C23}'!$E${cmr[key]}:${pc(MON_C)}${cmr[key]},'{C23}'!$E${cmr['PIDX']}:${pc(MON_C)}${cmr['PIDX']},{T('T_P', i)})"
    RV = lambda k, i: X("24_Revenue", k, i)
    OX = lambda k, i: X("25_Opex", k, i)
    AS = lambda k, i: X("26_Assets", k, i)
    LN = lambda k, i: X("27_Loan", k, i)
    TX = lambda k, i: X("28_Tax", k, i)
    WC = lambda k, i: X("30_WorkingCap", k, i)

    # =================================================================== 40 P&L
    pl = Sheet(b, "40_PL", "40 · Profit & loss", "የትርፍና ኪሳራ መግለጫ", TAB_ST, banner=True)
    p = lambda k, i: pl.c(k, i)
    pp = lambda k, i: (pl.c(k, i - 1) if i > 1 else "0")
    U = "=UNIT_LBL"
    pl.sec("Revenue (IFRS 15 / IFRS 16)", "ገቢ")
    pl.line("SALES", "Sales of units (IFRS 15)", lambda i: RV("SALES_REV", i), fmt="k", unit=U)
    pl.line("RENT", "Rental income (IFRS 16)", lambda i: RV("RENT", i), fmt="k", unit=U)
    pl.line("KEY", "Key money released", lambda i: RV("KEYA", i), fmt="k", unit=U)
    pl.line("SCR", "Service charge recoveries", lambda i: RV("SCR", i), fmt="k", unit=U)
    pl.line("OTH", "Other income", lambda i: RV("OTH", i), fmt="k", unit=U)
    pl.line("REV", "Total revenue", lambda i: "+".join(p(k, i) for k in ("SALES", "RENT", "KEY", "SCR", "OTH")), fmt="k",
            unit=U, style="total")
    pl.line("COGS", "Cost of sales – units (IAS 2)", lambda i: f"-{RV('COGS', i)}", fmt="k", unit=U)
    pl.line("COMM", "Sales commission", lambda i: f"-{RV('COMM', i)}", fmt="k", unit=U)
    pl.line("GP", "Gross profit", lambda i: f"{p('REV', i)}+{p('COGS', i)}+{p('COMM', i)}", fmt="k", unit=U, style="total")
    pl.sec("Operating expenses", "የሥራ ማስኬጃ ወጪ")
    for k, lab in (("PROP", "Property operating costs"), ("MGMT", "Management fee"), ("INS", "Insurance"),
                   ("PTAX", "Property tax"), ("MAINT", "Maintenance"), ("MKT", "Marketing"), ("OVH", "Owner overheads"),
                   ("PREOP", "Pre-opening & launch marketing"), ("BAD", "Bad debts")):
        pl.line(k, lab, (lambda k: lambda i: f"-{OX(k, i)}")(k), fmt="k", unit=U)
    pl.line("OPEX", "Total operating expenses",
            lambda i: "+".join(p(k, i) for k in ("PROP", "MGMT", "INS", "PTAX", "MAINT", "MKT", "OVH", "PREOP", "BAD")), fmt="k", unit=U,
            style="total")
    pl.line("EBITDA", "EBITDA", lambda i: f"{p('GP', i)}+{p('OPEX', i)}", fmt="k", unit=U, style="key")
    pl.line("DIP", "Depreciation – investment property", lambda i: f"-{AS('DEP_IP', i)}", fmt="k", unit=U)
    pl.line("DPPE", "Depreciation – owner-occupied PPE", lambda i: f"-{AS('DEP_PPE', i)}", fmt="k", unit=U)
    pl.line("EBIT", "Operating profit (EBIT)", lambda i: f"{p('EBITDA', i)}+{p('DIP', i)}+{p('DPPE', i)}", fmt="k", unit=U,
            style="total")
    pl.line("INTL", "Finance cost – bank loan (post-completion)", lambda i: f"-{LN('A_INTE', i)}", fmt="k", unit=U)
    pl.line("INTLS", "Finance cost – land lease", lambda i: f"-{LN('A_LSINT', i)}", fmt="k", unit=U)
    pl.line("PBT", "Profit before tax", lambda i: f"{p('EBIT', i)}+{p('INTL', i)}+{p('INTLS', i)}", fmt="k", unit=U,
            style="total")
    pl.line("CIT", "Current income tax", lambda i: f"-{TX('CIT', i)}", fmt="k", unit=U)
    pl.line("DT", "Deferred tax", lambda i: f"-{TX('DTE', i)}", fmt="k", unit=U)
    pl.line("UPT", "Undistributed profits tax", lambda i: f"-{TX('UPT', i)}", fmt="k", unit=U)
    pl.line("PAT", "PROFIT AFTER TAX", lambda i: f"{p('PBT', i)}+{p('CIT', i)}+{p('DT', i)}+{p('UPT', i)}", fmt="k", unit=U,
            style="grand")
    pl.line("PAT_NCI", "  attributable to non-controlling interest", lambda i: f"{p('PAT', i)}*NCI_PCT", fmt="k", unit=U)
    pl.line("PAT_OWN", "  attributable to owners", lambda i: f"{p('PAT', i)}-{p('PAT_NCI', i)}", fmt="k", unit=U)
    pl.sec("Appropriations", "የትርፍ ድልድል")
    pl.line("LRT", "Transfer to legal reserve (Commercial Code)",
            lambda i: f"IF({p('PAT', i)}>0,MAX(0,MIN(LR_PCT*{p('PAT', i)},LR_CAP*({X('41_BS', 'SC', i - 1) if i > 1 else '0'}+{X('42_CF', 'EQ', i)})-{pp('LRB', i)})),0)",
            fmt="k", unit=U)
    pl.line("LRB", "Legal reserve balance", lambda i: f"{pp('LRB', i)}+{p('LRT', i)}", fmt="k", unit=U, tot="last")
    pl.line("DIVD", "Dividends paid (incl. special)", lambda i: f"-{X('42_CF', 'DIV', i)}", fmt="k", unit=U)
    pl.line("RETOT", "Retained earnings (total, before NCI split)",
            lambda i: f"{pp('RETOT', i)}+{p('PAT', i)}-{p('LRT', i)}-{p('DIVD', i)}", fmt="k", unit=U, tot="last")
    pl.line("MARG", "EBITDA margin", lambda i: f"IF({p('REV', i)}>0,{p('EBITDA', i)}/{p('REV', i)},0)", fmt="pct", unit="%", tot=None)

    # =================================================================== 42 CASH FLOW (declared before BS so BS can link)
    cf = Sheet(b, "42_CF", "42 · Cash flow statement (IAS 7, indirect)", "የጥሬ ገንዘብ ፍሰት መግለጫ", TAB_ST, banner=True)
    c = lambda k, i: cf.c(k, i)
    cp = lambda k, i: (cf.c(k, i - 1) if i > 1 else "0")
    D = lambda sh, k, i: f"({X(sh, k, i)}-{X(sh, k, i - 1) if i > 1 else '0'})"
    cf.line("DEVC", "Memo: development cash paid (capex − IDC − non-cash land − Δretention)",
            lambda i: f"{AS('CAPEX', i)}-{AS('IDC', i)}-{'(LAND_TOT-LANDCASH_TOT)' if i == 1 else '0'}-{D('30_WorkingCap', 'RET', i)}",
            fmt="k", unit=U)
    cf.sec("Operating activities", "ከሥራ እንቅስቃሴ")
    cf.line("PBT", "Profit before tax", lambda i: X("40_PL", "PBT", i), fmt="k", unit=U)
    cf.line("DEP", "Add: depreciation", lambda i: AS("DEP", i), fmt="k", unit=U)
    cf.line("COGS", "Add: cost of sales released (non-cash)", lambda i: RV("COGS", i), fmt="k", unit=U)
    cf.line("FIN", "Add: finance costs (shown in financing)", lambda i: f"{LN('A_INTE', i)}+{LN('A_LSINT', i)}", fmt="k", unit=U)
    cf.line("DCL", "Increase in contract liability / (asset)", lambda i: D("24_Revenue", "CBAL", i), fmt="k", unit=U)
    cf.line("DKEY", "Increase in deferred key money", lambda i: D("24_Revenue", "KEYD", i), fmt="k", unit=U)
    cf.line("DAR", "(Increase) in receivables", lambda i: f"-{D('30_WorkingCap', 'AR', i)}", fmt="k", unit=U)
    cf.line("DAP", "Increase in payables", lambda i: D("30_WorkingCap", "AP", i), fmt="k", unit=U)
    cf.line("DVAT", "(Increase) in VAT receivable", lambda i: f"-{D('29_VAT', 'CLOSE', i)}", fmt="k", unit=U)
    cf.line("TAX", "Income tax paid", lambda i: f"-{TX('PAID', i)}", fmt="k", unit=U)
    cf.line("DINV", "Development costs paid – units for sale (IAS 2)", lambda i: f"-INV_SH*{c('DEVC', i)}", fmt="k", unit=U)
    cf.line("CFO", "Net cash from operating activities",
            lambda i: "+".join(c(k, i) for k in ("PBT", "DEP", "COGS", "FIN", "DCL", "DKEY", "DAR", "DAP", "DVAT", "TAX", "DINV")),
            fmt="k", unit=U, style="total")
    cf.sec("Investing activities", "ከኢንቨስትመንት")
    cf.line("DRET", "Development costs paid – retained property (IAS 40 / 16)", lambda i: f"-RET_SH*{c('DEVC', i)}", fmt="k", unit=U)
    cf.line("CFI", "Net cash from investing activities", lambda i: c("DRET", i), fmt="k", unit=U, style="total")
    cf.sec("Financing activities", "ከፋይናንስ")
    cf.line("EQ", "Equity injected (incl. top-up)", lambda i: SUMC("EQT", i), fmt="k", unit=U)
    cf.line("DRAW", "Loan drawdowns", lambda i: LN("A_DRAW", i), fmt="k", unit=U)
    cf.line("PRIN", "Loan principal repaid (incl. sweep)", lambda i: f"-({LN('A_PRIN', i)}+{LN('A_SWP', i)})", fmt="k", unit=U)
    cf.line("INTP", "Loan interest paid", lambda i: f"-{LN('A_INTP', i)}", fmt="k", unit=U)
    cf.line("LSPR", "Land lease instalments", lambda i: f"-{LN('A_LSPR', i)}", fmt="k", unit=U)
    cf.line("LSINT", "Land lease interest", lambda i: f"-{LN('A_LSINT', i)}", fmt="k", unit=U)
    cf.line("DSRA", "(Funding) / release of DSRA", lambda i: f"-({c('DSRAB', i)}-{cp('DSRAB', i)})", fmt="k", unit=U)
    cf.line("CFFPRE", "Net financing before dividends",
            lambda i: "+".join(c(k, i) for k in ("EQ", "DRAW", "PRIN", "INTP", "LSPR", "LSINT", "DSRA")), fmt="k", unit=U,
            style="total")
    cf.line("CASHPRE", "Cash available before dividends",
            lambda i: f"{cp('CLOSE', i)}+{c('CFO', i)}+{c('CFI', i)}+{c('CFFPRE', i)}", fmt="k", unit=U, tot=None)
    cf.line("LOCK", "Dividend lock-up flag (DSCR below covenant)",
            lambda i: f"IF(AND({LN('A_CLOSE', i)}>0,{X('43_Ratios', 'DSCRN', i)}<DSCR_COV),1,0)", fmt="int", unit="flag")
    cf.line("DIVP", "Dividends (positive)",
            lambda i: 0 if i <= 4 else (f"IF({c('LOCK', i)}=0,MAX(0,MIN(DIV_PAYOUT*MAX(0,{X('40_PL', 'PAT', i)}-{X('40_PL', 'LRT', i)}),"
                                        f"{c('CASHPRE', i)}-MIN_CASH,{X('40_PL', 'RETOT', i - 1)}+{X('40_PL', 'PAT', i)}-{X('40_PL', 'LRT', i)})),0)"),
            fmt="k", unit=U)
    cf.line("DIVX", "Special dividend of surplus cash (positive)",
            lambda i: 0 if i <= 4 else (f"IF(AND(DIST_SURPLUS=1,{c('LOCK', i)}=0),MAX(0,MIN({c('CASHPRE', i)}-{c('DIVP', i)}-{c('CRP', i)}-MIN_CASH,"
                                        f"{X('40_PL', 'RETOT', i - 1)}+{X('40_PL', 'PAT', i)}-{X('40_PL', 'LRT', i)}-{c('DIVP', i)})),0)"),
            fmt="k", unit=U)
    cf.line("DIV", "Dividends paid", lambda i: f"-({c('DIVP', i)}+{c('DIVX', i)})", fmt="k", unit=U)
    cf.line("CRP", "Return of surplus capital (positive)",
            lambda i: 0 if i <= 4 else (f"IF(AND(DIST_SURPLUS=1,{c('LOCK', i)}=0),MAX(0,MIN({c('CASHPRE', i)}-{c('DIVP', i)}-MIN_CASH,"
                                        f"{X('41_BS', 'SC', i - 1)}+{c('EQ', i)}-IF(LAND_TYPE=2,LAND_TOT,0))),0)"),
            fmt="k", unit=U)
    cf.line("CR", "Capital returned to shareholders", lambda i: f"-{c('CRP', i)}", fmt="k", unit=U)
    cf.line("CFF", "Net cash from financing activities", lambda i: f"{c('CFFPRE', i)}+{c('DIV', i)}+{c('CR', i)}", fmt="k", unit=U,
            style="total")
    cf.line("NET", "Net change in cash", lambda i: f"{c('CFO', i)}+{c('CFI', i)}+{c('CFF', i)}", fmt="k", unit=U, style="total")
    cf.line("OPEN", "Opening cash", lambda i: cp("CLOSE", i), fmt="k", unit=U, tot="first")
    cf.line("CLOSE", "CLOSING CASH", lambda i: f"{c('OPEN', i)}+{c('NET', i)}", fmt="k", unit=U, tot="last", style="grand")
    cf.sec("Debt service reserve account (DSRA)", "የዕዳ ክፍያ መጠባበቂያ")
    cf.line("DSRAB", "DSRA balance (restricted cash)",
            lambda i: 0 if i <= 4 else f"IF({LN('A_CLOSE', i)}>0,DSRA_M/12*{LN('A_DS', i + 1) if i < ANN_N else LN('A_DS', i)},0)",
            fmt="k", unit=U, tot="last")

    # =================================================================== 41 BALANCE SHEET
    bs = Sheet(b, "41_BS", "41 · Balance sheet", "የሀብትና ዕዳ መግለጫ", TAB_ST, banner=True)
    s = lambda k, i: bs.c(k, i)
    sp = lambda k, i: (bs.c(k, i - 1) if i > 1 else "0")
    bs.sec("Assets", "ሀብት")
    bs.line("IP", "Investment property (IAS 40, cost)", lambda i: AS("NBV_IP", i), fmt="k", unit=U, tot="last")
    bs.line("PPE", "Property, plant & equipment (IAS 16)", lambda i: AS("NBV_PPE", i), fmt="k", unit=U, tot="last")
    bs.line("DTA", "Deferred tax asset", lambda i: f"MAX(0,-{TX('DTN', i)})", fmt="k", unit=U, tot="last")
    bs.line("NCA", "Non-current assets", lambda i: f"{s('IP', i)}+{s('PPE', i)}+{s('DTA', i)}", fmt="k", unit=U, tot="last",
            style="total")
    bs.line("INV", "Inventory – units for sale (IAS 2)", lambda i: AS("INV", i), fmt="k", unit=U, tot="last")
    bs.line("AR", "Trade receivables", lambda i: WC("AR", i), fmt="k", unit=U, tot="last")
    bs.line("VAT", "VAT receivable", lambda i: WC("VAT", i), fmt="k", unit=U, tot="last")
    bs.line("CA", "Contract asset (IFRS 15)", lambda i: WC("CA", i), fmt="k", unit=U, tot="last")
    bs.line("DSRA", "Restricted cash – DSRA", lambda i: X("42_CF", "DSRAB", i), fmt="k", unit=U, tot="last")
    bs.line("CASH", "Cash and cash equivalents", lambda i: X("42_CF", "CLOSE", i), fmt="k", unit=U, tot="last")
    bs.line("CUR", "Current assets", lambda i: "+".join(s(k, i) for k in ("INV", "AR", "VAT", "CA", "DSRA", "CASH")), fmt="k",
            unit=U, tot="last", style="total")
    bs.line("TA", "TOTAL ASSETS", lambda i: f"{s('NCA', i)}+{s('CUR', i)}", fmt="k", unit=U, tot="last", style="grand")
    bs.sec("Liabilities", "ዕዳ")
    bs.line("LOAN", "Bank loan", lambda i: LN("A_CLOSE", i), fmt="k", unit=U, tot="last")
    bs.line("LEASE", "Land lease liability", lambda i: LN("A_LSCL", i), fmt="k", unit=U, tot="last")
    bs.line("DTL", "Deferred tax liability", lambda i: f"MAX(0,{TX('DTN', i)})", fmt="k", unit=U, tot="last")
    bs.line("KEY", "Deferred income – key money", lambda i: WC("KEY", i), fmt="k", unit=U, tot="last")
    bs.line("RET", "Retention payable", lambda i: WC("RET", i), fmt="k", unit=U, tot="last")
    bs.line("AP", "Trade payables", lambda i: WC("AP", i), fmt="k", unit=U, tot="last")
    bs.line("CL", "Contract liability – customer advances (IFRS 15)", lambda i: WC("CL", i), fmt="k", unit=U, tot="last")
    bs.line("TAXP", "Income tax payable", lambda i: TX("PAY", i), fmt="k", unit=U, tot="last")
    bs.line("TL", "TOTAL LIABILITIES", lambda i: "+".join(s(k, i) for k in ("LOAN", "LEASE", "DTL", "KEY", "RET", "AP", "CL", "TAXP")),
            fmt="k", unit=U, tot="last", style="total")
    bs.sec("Equity", "የባለቤት ሀብት")
    bs.line("SC", "Share capital / owner contributions",
            lambda i: f"{sp('SC', i)}+{X('42_CF', 'EQ', i)}+{'IF(LAND_TYPE=2,LAND_TOT,0)' if i == 1 else '0'}-{X('42_CF', 'CRP', i)}", fmt="k", unit=U,
            tot="last")
    bs.line("LR", "Legal reserve", lambda i: X("40_PL", "LRB", i), fmt="k", unit=U, tot="last")
    bs.line("RE", "Retained earnings", lambda i: f"(1-NCI_PCT)*{X('40_PL', 'RETOT', i)}", fmt="k", unit=U, tot="last")
    bs.line("EQO", "Equity attributable to owners", lambda i: f"{s('SC', i)}+{s('LR', i)}+{s('RE', i)}", fmt="k", unit=U,
            tot="last", style="total")
    bs.line("NCI", "Non-controlling interest (in-kind land partner)",
            lambda i: f"IF(LAND_TYPE=3,LAND_TOT,0)+NCI_PCT*{X('40_PL', 'RETOT', i)}", fmt="k", unit=U, tot="last")
    bs.line("TE", "TOTAL EQUITY", lambda i: f"{s('EQO', i)}+{s('NCI', i)}", fmt="k", unit=U, tot="last", style="total")
    bs.line("TLE", "TOTAL LIABILITIES AND EQUITY", lambda i: f"{s('TL', i)}+{s('TE', i)}", fmt="k", unit=U, tot="last",
            style="grand")
    bs.line("CHK", "Balance check (assets − liabilities − equity)", lambda i: f"ROUND({s('TA', i)}-{s('TLE', i)},2)", fmt="n2",
            unit=U, tot="sum", style="key")
    bs.gap()
    bs.scalar("NCIP", "NCI share (land in-kind ÷ total equity)",
              "IF(AND(LAND_TYPE=3,LAND_TOT+EQ_DRAWN+FUND_GAP>0),LAND_TOT/(LAND_TOT+EQ_DRAWN+FUND_GAP),0)", fmt="pct", unit="%",
              name="NCI_PCT")

    # =================================================================== 43 RATIOS
    ra = Sheet(b, "43_Ratios", "43 · Ratios & lender pack", "የአፈጻጸም ምጣኔዎችና የባንክ መረጃ", TAB_ST, banner=True)
    q = lambda k, i: ra.c(k, i)
    qn = lambda k, i: (ra.c(k, i + 1) if i < ANN_N else "0")
    ra.line("CFADS", "Cash flow available for debt service (CFADS)",
            lambda i: f"{X('42_CF', 'CFO', i)}+{X('42_CF', 'CFI', i)}+{X('42_CF', 'LSPR', i)}+{X('42_CF', 'LSINT', i)}", fmt="k",
            unit=U)
    ra.line("DS", "Debt service", lambda i: LN("A_DS", i), fmt="k", unit=U)
    ra.line("DSCR", "DSCR", lambda i: 0 if i <= 4 else f'IF({q("DS", i)}>0,{q("CFADS", i)}/{q("DS", i)},"")', fmt="x", unit="x",
            tot=None, style="key")
    ra.line("DSCRN", "DSCR (numeric helper, 99 = no debt service)",
            lambda i: f"IF({q('DS', i)}>0,{q('CFADS', i)}/{q('DS', i)},99)", fmt="x", unit="x", tot=None)
    ra.line("COV", "Covenant", lambda i: "DSCR_COV", fmt="x", unit="x", tot=None)
    ra.line("PVC", "PV of CFADS over remaining loan life",
            lambda i: 0 if i <= 4 else f"IF({LN('A_OPEN', i)}>0,({q('CFADS', i)}+{qn('PVC', i)})/(1+RATE_EFF),0)", fmt="k", unit=U,
            tot=None)
    ra.line("LLCR", "LLCR", lambda i: 0 if i <= 4 else f'IF({LN("A_OPEN", i)}>0,{q("PVC", i)}/{LN("A_OPEN", i)},"")', fmt="x",
            unit="x", tot=None)
    ra.line("ICR", "Interest cover (EBITDA ÷ finance cost)",
            lambda i: 0 if i <= 4 else f'IF({LN("A_INTE", i)}+{LN("A_LSINT", i)}>0,{X("40_PL", "EBITDA", i)}/({LN("A_INTE", i)}+{LN("A_LSINT", i)}),"")',
            fmt="x", unit="x", tot=None)
    ra.line("GEAR", "Gearing (loan ÷ (loan + equity))",
            lambda i: f"IF({X('41_BS', 'LOAN', i)}+{X('41_BS', 'TE', i)}>0,{X('41_BS', 'LOAN', i)}/({X('41_BS', 'LOAN', i)}+{X('41_BS', 'TE', i)}),0)",
            fmt="pct", unit="%", tot="max")
    ra.line("NDE", "Net debt ÷ EBITDA",
            lambda i: f'IF({X("40_PL", "EBITDA", i)}>0,({X("41_BS", "LOAN", i)}-{X("41_BS", "CASH", i)})/{X("40_PL", "EBITDA", i)},"")',
            fmt="x", unit="x", tot=None)
    ra.line("CASHO", "Cash balance in operating years", lambda i: 0 if i <= 4 else X("42_CF", "CLOSE", i), fmt="k", unit=U, tot=None)
    ra.line("DSCRF", "Covenant breach flag", lambda i: f'IF(AND({q("DS", i)}>0,{q("DSCRN", i)}<DSCR_COV),1,0)', fmt="int",
            unit="flag")
    ra.line("CASHF", "Below minimum cash flag (operating years)",
            lambda i: 0 if i <= 4 else f"IF({X('42_CF', 'CLOSE', i)}<MIN_CASH-0.5,1,0)", fmt="int", unit="flag")
    ra.gap()
    ra.sec("Summary", "ማጠቃለያ")
    rr = lambda k: f"$I${ra.rows[k]}:$R${ra.rows[k]}"
    ra.scalar("MIN", "Minimum DSCR", f"IF(COUNT({rr('DSCR')})>0,MIN({rr('DSCR')}),0)", fmt="x", unit="x", name="MIN_DSCR",
              style="key")
    ra.scalar("AVG", "Average DSCR", f"IF(COUNT({rr('DSCR')})>0,AVERAGE({rr('DSCR')}),0)", fmt="x", unit="x", name="AVG_DSCR")
    ra.scalar("NDS", "Years with debt service", f"COUNT({rr('DSCR')})", fmt="int", unit="no.", name="N_DS_YEARS")
    ra.scalar("LLCR1", "LLCR at first operating year", f"IF(ISNUMBER($I${ra.rows['LLCR']}),$I${ra.rows['LLCR']},0)", fmt="x",
              unit="x", name="LLCR_O1")
    ra.scalar("MINICR", "Minimum ICR", f"IF(COUNT({rr('ICR')})>0,MIN({rr('ICR')}),0)", fmt="x", unit="x", name="MIN_ICR")
    ra.scalar("MAXG", "Maximum gearing", f"$D${ra.rows['GEAR']}", fmt="pct", unit="%", name="MAX_GEAR")
    ra.scalar("BRE", "Years below DSCR covenant", f"$D${ra.rows['DSCRF']}", fmt="int", unit="no.", name="N_BREACH")
    ra.scalar("CMIN", "Minimum cash in operating years", f"MIN({rr('CASHO')})", unit=U, name="CASH_MIN_OPS")
    ra.scalar("CBF", "Years below minimum cash", f"$D${ra.rows['CASHF']}", fmt="int", unit="no.", name="N_LOWCASH")

    # =================================================================== 50 RETURNS
    rt = Sheet(b, "50_Returns", "50 · Returns – NPV, IRR, MOIC, margins", "ተመላሽ – NPV፣ IRR፣ ትርፋማነት", TAB_AN, banner=True)
    r = lambda k, i: rt.c(k, i)
    rt.sec("Terminal value at O10", "የመጨረሻ ዋጋ")
    rt.scalar("NOI10", "NOI in O10", lambda: f"'24_Revenue'!$R${b.sheets['24_Revenue'].rows['NOI']}", name="NOI_O10")
    rt.scalar("NOI3", "Stabilised NOI (O3)", lambda: f"'24_Revenue'!$K${b.sheets['24_Revenue'].rows['NOI']}", name="NOI_STAB")
    rt.scalar("TVG", "Terminal value (gross)",
              "MAX(0,IF(EXIT_METHOD=1,NOI_O10*(1+RESC_EFF)/EXIT_CAP_EFF,NOI_O10*EXIT_MULT))", name="TV_GROSS")
    rt.scalar("TVN", "Terminal value net of exit costs", "TV_GROSS*(1-EXIT_COST)", name="TV_NET", style="key")
    rt.scalar("EQX", "Equity exit value (owners' book equity + TV uplift)",
              lambda: f"'41_BS'!$R${b.sheets['41_BS'].rows['EQO']}+(TV_NET-'41_BS'!$R${b.sheets['41_BS'].rows['IP']}-'41_BS'!$R${b.sheets['41_BS'].rows['PPE']})*(1-NCI_PCT)",
              name="EQ_EXIT")
    rt.gap()
    rt.sec("Project cash flow (unlevered, after tax paid)", "የፕሮጀክት የገንዘብ ፍሰት")
    rt.line("PRE", "Operating + investing cash flow", lambda i: f"{X('42_CF', 'CFO', i)}+{X('42_CF', 'CFI', i)}", fmt="k", unit=U)
    rt.line("FEES", "Add back: financing fees paid", lambda i: SUMC("FINC", i), fmt="k", unit=U)
    rt.line("LAND", "Less: land lease payments", lambda i: f"{X('42_CF', 'LSPR', i)}+{X('42_CF', 'LSINT', i)}", fmt="k", unit=U)
    rt.line("TV", "Terminal value", lambda i: "TV_NET" if i == ANN_N else 0, fmt="k", unit=U)
    rt.line("PCF", "PROJECT CASH FLOW", lambda i: f"{r('PRE', i)}+{r('FEES', i)}+{r('LAND', i)}+{r('TV', i)}", fmt="k", unit=U,
            style="total")
    rt.line("PDCF", "Discounted @ project rate", lambda i: f"{r('PCF', i)}*{T('T_DF', i)}", fmt="k", unit=U)
    rt.sec("Equity cash flow", "የባለቤት የገንዘብ ፍሰት")
    rt.line("EIN", "Equity injected", lambda i: f"-{X('42_CF', 'EQ', i)}", fmt="k", unit=U)
    rt.line("EDIV", "Dividends + capital returned", lambda i: f"-{X('42_CF', 'DIV', i)}+{X('42_CF', 'CRP', i)}", fmt="k", unit=U)
    rt.line("EEX", "Equity exit value", lambda i: "EQ_EXIT" if i == ANN_N else 0, fmt="k", unit=U)
    rt.line("ECF", "EQUITY CASH FLOW", lambda i: f"{r('EIN', i)}+{r('EDIV', i)}+{r('EEX', i)}", fmt="k", unit=U, style="total")
    rt.line("EDCF", "Discounted @ cost of equity", lambda i: f"{r('ECF', i)}*{T('T_DFE', i)}", fmt="k", unit=U)
    rt.line("DATE", "Cash flow date (period end)", lambda i: T("T_END", i), fmt="date", unit="date", tot=None)
    rng = lambda k: f"$E${rt.rows[k]}:$R${rt.rows[k]}"
    dts = rng("DATE")
    rt.gap()
    rt.sec("Return metrics", "የተመላሽ መለኪያዎች")
    rt.scalar("PV", "Project cash flow valid for IRR (has + and −)", f"IF(AND(MIN({rng('PCF')})<0,MAX({rng('PCF')})>0),1,0)",
              fmt="int", unit="flag", name="PCF_OK")
    rt.scalar("EV", "Equity cash flow valid for IRR", f"IF(AND(MIN({rng('ECF')})<0,MAX({rng('ECF')})>0),1,0)", fmt="int",
              unit="flag", name="ECF_OK")
    rt.scalar("PNPV", "Project NPV @ discount rate", f"$D${rt.rows['PDCF']}", name="PNPV", style="key")
    rt.scalar("PIRR", "Project IRR (XIRR, dated)", f"IF(PCF_OK=1,IFERROR(XIRR({rng('PCF')},{dts},0.15),IFERROR(XIRR({rng('PCF')},{dts},-0.05),0)),0)", fmt="pct", unit="%",
              name="PIRR", style="key")
    rt.scalar("ENPV", "Equity NPV @ cost of equity", f"$D${rt.rows['EDCF']}", name="ENPV", style="key")
    rt.scalar("EIRR", "Equity IRR (XIRR, dated)", f"IF(ECF_OK=1,IFERROR(XIRR({rng('ECF')},{dts},0.2),IFERROR(XIRR({rng('ECF')},{dts},-0.05),0)),0)", fmt="pct", unit="%",
              name="EIRR", style="key")
    rt.scalar("PIRRP", "Memo: project IRR per period (IRR on columns)", f"IF(PCF_OK=1,IFERROR(IRR({rng('PCF')},0.15),0),0)",
              fmt="pct", unit="%", name="PIRR_PER")
    rt.scalar("EMOIC", "Equity multiple (MOIC)",
              f"IF(SUMIF({rng('ECF')},\"<0\")<0,SUMIF({rng('ECF')},\">0\")/-SUMIF({rng('ECF')},\"<0\"),0)", fmt="x", unit="x",
              name="EMOIC", style="key")
    rt.scalar("PMOIC", "Project multiple",
              f"IF(SUMIF({rng('PCF')},\"<0\")<0,SUMIF({rng('PCF')},\">0\")/-SUMIF({rng('PCF')},\"<0\"),0)", fmt="x", unit="x",
              name="PMOIC")
    rt.scalar("SNOM", "Sales revenue (nominal, net)", lambda: XT("24_Revenue", "SALES_REV"), name="SALES_NOM")
    rt.scalar("RVAL", "Value of retained property at stabilisation (NOI O3 ÷ exit cap)",
              "MAX(0,NOI_STAB/EXIT_CAP_EFF)", name="RET_VALUE")
    rt.scalar("GDV", "Gross development value (GDV)", "SALES_NOM+RET_VALUE", name="GDV", style="key")
    rt.scalar("POC", "Profit on cost", "IF(TDC>0,(GDV-TDC)/TDC,0)", fmt="pct", unit="%", name="POC", style="key")
    rt.scalar("DM", "Development margin (profit ÷ GDV)", "IF(GDV>0,(GDV-TDC)/GDV,0)", fmt="pct", unit="%", name="DEV_MARGIN")
    rt.scalar("YOC", "Yield on cost (stabilised NOI ÷ retained cost)", "IF(RET_COST>0,NOI_STAB/RET_COST,0)", fmt="pct", unit="%",
              name="YOC")
    rt.scalar("GT", "GDV ÷ TDC", "IF(TDC>0,GDV/TDC,0)", fmt="x", unit="x", name="GDV_TDC")
    rt.scalar("PATS", "Total profit after tax (C1–O10)", lambda: XT("40_PL", "PAT"), name="PAT_TOT")
    rt.scalar("PEQ", "Peak equity (committed + top-up)", "EQ_DRAWN+FUND_GAP", name="PEAK_EQ")
    rt.scalar("DOK", "DSCR test passed (or no debt service)", "IF(N_DS_YEARS=0,1,IF(MIN_DSCR>=DSCR_COV,1,0))", fmt="int",
              unit="flag", name="DSCR_OK")
    rt.scalar("VERD", "VERDICT",
              '=IF(TDC<=0,"NO DATA",IF(AND(EIRR>=HURDLE,POC>0,DSCR_OK=1),"VIABLE",IF(AND(EIRR>=HURDLE-MARGIN_BAND,GDV>=TDC),"MARGINAL","NOT VIABLE")))'[1:],
              fmt="gen", unit="", name="VERDICT", style="key")
    rt.scalar("VERDAM", "ውሳኔ", '=IF(VERDICT="VIABLE","አዋጭ",IF(VERDICT="MARGINAL","በድንበር ላይ",IF(VERDICT="NOT VIABLE","አዋጭ አይደለም","መረጃ የለም")))'[1:],
              fmt="gen", unit="", name="VERDICT_AM")
    rt.gap()
    rt.sec("IAS 36 – impairment test summary", "የንብረት ዋጋ መቀነስ ፈተና")
    rt.scalar("IMPN", "Years with impairment indicator", lambda: XT("26_Assets", "IMPF"), fmt="int", unit="no.", name="N_IMPAIR")
    rt.scalar("IMPH", "Lowest headroom (value in use − carrying amount)", lambda: XT("26_Assets", "HEAD"), name="MIN_HEADROOM")

    # =================================================================== 51 PAYBACK
    pb = Sheet(b, "51_Payback", "51 · Payback – simple and discounted", "የመመለሻ ጊዜ", TAB_AN, banner=True)
    g = lambda k, i: pb.c(k, i)
    gp = lambda k, i: (pb.c(k, i - 1) if i > 1 else "0")
    yfp = lambda i: (T("T_YF", i - 1) if i > 1 else "0")

    def pb_block(prefix, src, label):
        pb.sec(label)
        pb.line(prefix + "CF", "Cash flow", lambda i: X("50_Returns", src, i), fmt="k", unit=U)
        pb.line(prefix + "CUM", "Cumulative", lambda i: f"{gp(prefix + 'CUM', i)}+{g(prefix + 'CF', i)}", fmt="k", unit=U,
                tot="last")
        pb.line(prefix + "X", "Crossing point (years from start)",
                lambda i: f"IF(AND({gp(prefix + 'CUM', i)}<0,{g(prefix + 'CUM', i)}>=0,{g(prefix + 'CF', i)}>0),{yfp(i)}+(-{gp(prefix + 'CUM', i)}/{g(prefix + 'CF', i)})*{T('T_LEN', i)},0)",
                fmt="n2", unit="years", tot="max")
    pb_block("P", "PCF", "Project – simple")
    pb_block("PD", "PDCF", "Project – discounted")
    pb_block("E", "ECF", "Equity – simple")
    pb_block("ED", "EDCF", "Equity – discounted")
    pb.gap()
    pb.sec("Payback results (years from construction start; fractional-year method)", "ውጤት")
    for k, lab, nm in (("P", "Project payback – simple", "PB_P"), ("PD", "Project payback – discounted", "PB_PD"),
                       ("E", "Equity payback – simple", "PB_E"), ("ED", "Equity payback – discounted", "PB_ED")):
        pb.scalar(nm, lab, f"$D${pb.rows[k + 'X']}", fmt="n1", unit="years", name=nm, style="key",
                  note="0 = not reached within the horizon")

    return pl, cf, bs, ra, rt, pb


def build_sens(b):
    """52 Sensitivity & break-even – live what-if dials, first-order tornado and closed-form break-evens.

    No Excel Data Tables: everything recalculates in Excel, LibreOffice and Google Sheets alike."""
    X = lambda sh, key, i: b.sheets[sh].c(key, i, sheet=True)
    T = lambda key, i: f"'20_Timeline'!{pc(i)}{b.TL[key]}"
    C23 = "23_Construction"
    cmr = b.sheets[C23].rows
    SUMC = lambda key, i: f"SUMIFS('{C23}'!$E${cmr[key]}:${pc(MON_C)}${cmr[key]},'{C23}'!$E${cmr['PIDX']}:${pc(MON_C)}${cmr['PIDX']},{T('T_P', i)})"
    se = Sheet(b, "52_Sensitivity", "52 · Sensitivity & break-even (LIVE)", "የስጋት ትንተናና የሚዛን ነጥብ", TAB_AN,
               kind="free", n=0, width_label=58)
    ws = se.ws
    se.sec("What-if dials – change one and every sheet updates (100% = as input). Put them back to 100% / 0 afterwards.",
           "የማስተካከያ ሕዋሶች")
    flex = [("FLX_PRICE", "Sale price", 1.0, "pct"), ("FLX_COST", "Construction cost", 1.0, "pct"),
            ("FLX_RENT", "Rent and hotel rate", 1.0, "pct"), ("FLX_OCC", "Occupancy", 1.0, "pct"),
            ("FLX_RATE", "Interest rate", 1.0, "pct"), ("FLX_OPEX", "Operating costs", 1.0, "pct"),
            ("FLX_DELAY", "Extra construction delay (months)", 0, "int")]
    b.FLEX = {}
    for k, lab, v, fmt in flex:
        r = se.scalar(k, lab, v, fmt=fmt, unit="x" if fmt == "pct" else "months", style="input", name=k)
        b.FLEX[k] = f"$D${r}"
        b.DV.append(("52_Sensitivity", "dec", ("dec", 0, 5) if fmt == "pct" else ("dec", 0, 24), f"D{r}"))
    se.scalar("FLXON", "Dials changed from neutral? (1 = yes – results are NOT the base case)",
              "IF(AND(FLX_PRICE=1,FLX_COST=1,FLX_RENT=1,FLX_OCC=1,FLX_RATE=1,FLX_OPEX=1,FLX_DELAY=0),0,1)", fmt="int", unit="flag",
              name="FLEX_ON", style="key")
    se.gap()
    se.sec("Live outputs", "ውጤቶች")
    outs = [("O_EIRR", "Equity IRR", "EIRR", "pct"), ("O_PIRR", "Project IRR", "PIRR", "pct"), ("O_PNPV", "Project NPV", "PNPV", "n0"),
            ("O_POC", "Profit on cost", "POC", "pct"), ("O_DSCR", "Minimum DSCR", "MIN_DSCR", "x"),
            ("O_PEAK", "Peak debt", "PEAK_DEBT", "n0"), ("O_TDC", "Total development cost", "TDC", "n0"),
            ("O_VERD", "Verdict", "VERDICT", "gen")]
    for k, lab, f, fmt in outs:
        se.scalar(k, lab, f, fmt=fmt, unit="ETB" if fmt == "n0" else "", style="key")
    se.gap()
    # ------------------------------------------------------------------ helpers (annual)
    se.sec("Present values used by the tornado and break-evens (annual, @ project discount rate)", "ረዳት ስሌቶች")
    rr_ = lambda sh, k, i: X(sh, k, i)
    se.line("DF", "Discount factor @ project rate", lambda i: T("T_DF", i), fmt="n2", unit="x", tot=None, n=ANN_N)
    se.line("SPCF", "Space rent less %-based costs", lambda i: f"{rr_('24_Revenue', 'RENT_SP', i)}*(1-OPEX_PCT*(OPEX_METHOD=1)-MGMT_PCT-BAD_DEBT)",
            fmt="n0", unit="ETB", n=ANN_N)
    se.line("HOTCF", "Hotel rent less %-based costs", lambda i: f"{rr_('24_Revenue', 'HOTEL', i)}*(1-MGMT_PCT-BAD_DEBT)",
            fmt="n0", unit="ETB", n=ANN_N)
    se.line("CAPX", "Construction cost driven by the cost dial",
            lambda i: "+".join(SUMC(k, i) for k in ("HARD", "CONT", "DESIGN", "SUPERV", "CONSULT")), fmt="n0", unit="ETB", n=ANN_N)
    se.line("OPXL", "Operating costs driven by the opex dial",
            lambda i: "+".join(rr_("25_Opex", k, i) for k in ("INS", "MAINT", "OVH")) + f"+IF(OPEX_METHOD=2,{rr_('25_Opex', 'PROP', i)},0)",
            fmt="n0", unit="ETB", n=ANN_N)
    ln = b.sheets["27_Loan"].rows
    se.line("INTR", "Loan interest (paid + capitalised)",
            lambda i: f"'27_Loan'!{pc(i)}{ln['A_INTP']}+'27_Loan'!{pc(i)}{ln['A_CAPC']}+'27_Loan'!{pc(i)}{ln['A_CAPG']}", fmt="n0",
            unit="ETB", n=ANN_N)
    se.line("OPSIN", "Operating inflows (NOI + sales collections)",
            lambda i: f"{rr_('24_Revenue', 'NOI', i)}+{rr_('24_Revenue', 'SALES_COLL', i)}", fmt="n0", unit="ETB", n=ANN_N)
    sp = lambda k: f"$E${se.rows[k]}:$R${se.rows[k]}"
    tvshare = lambda k: f"IF(NOI_O10>0,TV_PV*$R${se.rows[k]}/NOI_O10,0)"
    se.scalar("TVPV", "PV of terminal value", f"TV_NET*$R${se.rows['DF']}", name="TV_PV")
    ann = lambda sh, k: f"'{sh}'!$E${b.sheets[sh].rows[k]}:$R${b.sheets[sh].rows[k]}"
    se.scalar("PVS", "PV of sales collections net of commission & marketing",
              f"SUMPRODUCT({ann('24_Revenue', 'SALES_COLL')},{sp('DF')})*(1-COMM_PCT-MKT_PCT)", name="PV_SALESNET")
    se.scalar("PVSP", "PV of space rent (incl. its share of terminal value)", f"SUMPRODUCT({sp('SPCF')},{sp('DF')})+{tvshare('SPCF')}",
              name="PV_SPRENT")
    se.scalar("PVH", "PV of hotel rent (incl. its share of terminal value)", f"SUMPRODUCT({sp('HOTCF')},{sp('DF')})+{tvshare('HOTCF')}",
              name="PV_HOTRENT")
    se.scalar("PVC", "PV of construction cost (cost dial)", f"SUMPRODUCT({sp('CAPX')},{sp('DF')})", name="PV_CAPX")
    se.scalar("PVO", "PV of operating costs (opex dial, incl. terminal value effect)",
              f"SUMPRODUCT({sp('OPXL')},{sp('DF')})+{tvshare('OPXL')}", name="PV_OPX")
    se.scalar("PVI", "PV of loan interest", f"SUMPRODUCT({sp('INTR')},{sp('DF')})", name="PV_INTR")
    se.scalar("PVIN", "PV of operating inflows incl. terminal value", f"SUMPRODUCT({sp('OPSIN')},{sp('DF')})+TV_PV", name="PV_OPSIN")
    se.scalar("DLY", "Value lost from a 6-month delay (first-order)",
              "PV_OPSIN*(1-(1+DISC_EFF)^(-0.5))+PV_CAPX*0.5*((1+ESC_EFF)^0.5-1)", name="DELAY6_PV")
    se.gap()
    # ------------------------------------------------------------------ tornado
    se.sec("Tornado – first-order change in project value for ±20% (delay: +6 months), ETB million", "የቶርኔዶ ሰንጠረዥ")
    se.note("Pre-tax estimate from the present values above. For the exact effect of any driver, change the matching dial at the top.")
    tr0 = se.r
    drivers = [("Sale price", "0.2*PV_SALESNET", "-0.2*PV_SALESNET"),
               ("Rent – spaces", "0.2*PV_SPRENT", "-0.2*PV_SPRENT"),
               ("Hotel trading (rate × occupancy)", "0.2*PV_HOTRENT", "-0.2*PV_HOTRENT"),
               ("Construction cost", "-0.2*PV_CAPX", "0.2*PV_CAPX"),
               ("Operating costs", "-0.2*PV_OPX", "0.2*PV_OPX"),
               ("Interest rate (value to owner)", "-0.2*PV_INTR", "0.2*PV_INTR"),
               ("Construction delay +6 months", "-DELAY6_PV", "0")]
    se.r += 1 + len(drivers) + 2 + len(drivers) + 1
    b.TORNADO = {}

    def _torn(_r):
        for j, h in enumerate(["Driver", "", "", "Adverse case", "Favourable case", "Swing", "Rank key"]):
            if h:
                ws.cell(tr0, 2 + j, h).font = font(9, True, NAVY)
        for k, (lab, hi, lo) in enumerate(drivers):
            rr = tr0 + 1 + k
            ws.cell(rr, 2, lab).font = F_BODY
            for col, f in ((5, f"MIN({hi},{lo})/1000000"), (6, f"MAX({hi},{lo})/1000000"), (7, f"F{rr}-E{rr}"),
                           (8, f"G{rr}+{k + 1}/1000000")):
                cc = ws.cell(rr, col, "=" + f)
                cc.number_format = NF["n1"] if col < 8 else NF["n2"]
                cc.font = F_BODY
        s0 = tr0 + 1 + len(drivers) + 1
        ws.cell(s0, 2, "Sorted (largest swing first) – chart source").font = font(9, True, NAVY)
        keyr = f"$H${tr0 + 1}:$H${tr0 + len(drivers)}"
        for k in range(len(drivers)):
            rr = s0 + 1 + k
            mt = f"MATCH(LARGE({keyr},{k + 1}),{keyr},0)"
            ws.cell(rr, 2, f"=INDEX($B${tr0 + 1}:$B${tr0 + len(drivers)},{mt})").font = F_BODY
            for col, srccol in ((5, "E"), (6, "F")):
                cc = ws.cell(rr, col, f"=INDEX(${srccol}${tr0 + 1}:${srccol}${tr0 + len(drivers)},{mt})")
                cc.number_format = NF["n1"]
        b.name("TOP_RISK", f"'52_Sensitivity'!$B${s0 + 1}")
        b.TORNADO = {"s0": s0, "n": len(drivers)}
    se.raw(_torn)
    # ------------------------------------------------------------------ break-even
    se.sec("Break-even analysis (live, closed-form – each holds the other drivers at their inputs)", "የሚዛን ነጥብ ትንተና")
    be0 = se.r
    be = [("BE_PRICE_NPV", "Sale price at which project NPV = 0", "IF(PV_SALESNET>0,BLEND_PRICE*(1-PNPV/PV_SALESNET),0)",
           "n0", "ETB/m²", "BLEND_PRICE", "IF(AND(BLEND_PRICE>0,E{r}>0),1-E{r}/BLEND_PRICE,0)"),
          ("BE_PRICE_POC", "Sale price at which profit on cost = 0", "IF(SALES_NOM>0,BLEND_PRICE*(TDC-RET_VALUE)/SALES_NOM,0)",
           "n0", "ETB/m²", "BLEND_PRICE", "IF(AND(BLEND_PRICE>0,E{r}>0),1-E{r}/BLEND_PRICE,0)"),
          ("BE_RENT_NPV", "Rent level at which project NPV = 0 (% of today's input)",
           "IF(PV_SPRENT+PV_HOTRENT>0,1-PNPV/(PV_SPRENT+PV_HOTRENT),0)", "pct", "% of rent", "1", "IF(E{r}>0,1-E{r},0)"),
          ("BE_OCC", "Stabilised occupancy at which min DSCR = covenant", "OCC_STAB*BE_OCC_MULT", "pct", "% occ.", "OCC_STAB",
           "IF(G{r}>0,1-E{r}/G{r},0)"),
          ("BE_COST", "Cost overrun before equity IRR falls below the hurdle", "IF(PV_HARD_H>0,ENPV_H/PV_HARD_H,0)", "pct",
           "overrun", "0", "E{r}"),
          ("BE_RATE", "Interest rate at which min DSCR = covenant", "BE_RATE_CF", "pct2", "% p.a.", "RATE_EFF",
           "IF(AND(G{r}>0,E{r}>0),E{r}/G{r}-1,0)"),
          ("BE_DELAY", "Delay before equity IRR falls below the hurdle",
           "IF(AND(PV_EIN_H>0,PV_EOUT_H>0),12*LN(PV_EIN_H/PV_EOUT_H)/LN(1+HURDLE),0)", "n1", "months", "0", "E{r}/12"),
          ("BE_CAP", "Exit cap rate at which project NPV = 0",
           f"IF(AND(TV_PV>0,TV_PV-PNPV>0),EXIT_CAP_EFF*TV_PV/(TV_PV-PNPV),0)", "pct2", "% yield", "EXIT_CAP_EFF",
           "IF(G{r}>0,E{r}/G{r}-1,0)")]
    se.r += 1 + len(be) + 2

    def _be(_r):
        for j, h in enumerate(["Break-even", "", "", "Value", "Unit", "Current", "Headroom"]):
            if h:
                ws.cell(be0, 2 + j, h).font = font(9, True, NAVY)
        for k, (key, lab, f, fmt, unit, cur, head) in enumerate(be):
            rr = be0 + 1 + k
            ws.cell(rr, 2, lab).font = F_BODY
            c1 = ws.cell(rr, 5, "=" + f)
            c1.number_format = NF[fmt]
            c1.fill = FILL_KEY
            c1.font = F_BOLD
            ws.cell(rr, 6, unit).font = F_UNIT
            c3 = ws.cell(rr, 7, "=" + cur)
            c3.number_format = NF[fmt]
            c4 = ws.cell(rr, 8, "=" + head.format(r=rr))
            c4.number_format = NF["pct"] if key != "BE_DELAY" else NF["yr"]
            b.name(key, f"'52_Sensitivity'!$E${rr}")
        ws.cell(be0 + 1 + len(be), 2, "0 = not applicable for this project (e.g. no sales, no debt service). Headroom = how far the "
                                        "input can move before the break-even is reached.").font = F_NOTE
        ws.column_dimensions["E"].width = 14
        ws.column_dimensions["F"].width = 14
        ws.column_dimensions["G"].width = 14
        ws.column_dimensions["H"].width = 12
    se.raw(_be)
    se.gap()
    se.sec("Closed-form helpers for break-evens", "ረዳት ስሌቶች")
    RVr = b.sheets["24_Revenue"].rows
    rt = b.sheets["50_Returns"].rows
    ra = b.sheets["43_Ratios"].rows
    se.line("DFH_ROW", "Discount factor @ hurdle", lambda i: f"(1+HURDLE)^(-'20_Timeline'!{pc(i)}{b.TL['T_YF']})", fmt="n2",
            unit="x", tot=None, n=ANN_N)
    dfh = f"$E${se.rows['DFH_ROW']}:$R${se.rows['DFH_ROW']}"
    se.scalar("ENPVH", "Equity NPV @ hurdle", f"SUMPRODUCT({ann('50_Returns', 'ECF')},{dfh})", name="ENPV_H")
    se.scalar("PVHH", "PV @ hurdle of construction cost (cost dial)", f"SUMPRODUCT({sp('CAPX')},{dfh})", name="PV_HARD_H")
    se.line("EIN", "Equity inflows (positive)", lambda i: f"MAX(0,'50_Returns'!{pc(i)}{rt['ECF']})", fmt="n0", unit="ETB", n=ANN_N)
    se.scalar("PVEI", "PV @ hurdle of equity inflows", f"SUMPRODUCT($E${se.rows['EIN']}:$R${se.rows['EIN']},{dfh})", name="PV_EIN_H")
    se.scalar("PVEO", "PV @ hurdle of equity outflows", "PV_EIN_H-ENPV_H", name="PV_EOUT_H")
    se.line("RENTCF", "Rent-driven cash (spaces + hotel, after %-based costs)",
            lambda i: f"{pc(i)}{se.rows['SPCF']}+{pc(i)}{se.rows['HOTCF']}", fmt="n0", unit="ETB", n=ANN_N)
    se.line("OCCM", "Occupancy multiplier needed for DSCR = covenant",
            lambda i: (0 if i <= 4 else
                       f"IF(AND('43_Ratios'!{pc(i)}{ra['DS']}>0,{pc(i)}{se.rows['RENTCF']}>0),(DSCR_COV*'43_Ratios'!{pc(i)}{ra['DS']}-'43_Ratios'!{pc(i)}{ra['CFADS']}+{pc(i)}{se.rows['RENTCF']})/{pc(i)}{se.rows['RENTCF']},0)"),
            fmt="n2", unit="x", tot="max", n=ANN_N)
    se.scalar("OCCMX", "Binding occupancy multiplier", f"MAX(0,$D${se.rows['OCCM']})", fmt="n2", unit="x", name="BE_OCC_MULT")
    se.line("RATEM", "Rate at which DSCR = covenant (by year)",
            lambda i: (99 if i <= 4 else
                       f"IF(AND('27_Loan'!{pc(i)}{ln['A_INTP']}>0,'43_Ratios'!{pc(i)}{ra['DS']}>0),RATE_EFF*('43_Ratios'!{pc(i)}{ra['CFADS']}/DSCR_COV-'27_Loan'!{pc(i)}{ln['A_PRIN']})/'27_Loan'!{pc(i)}{ln['A_INTP']},99)"),
            fmt="pct2", unit="%", tot="min", n=ANN_N)
    se.scalar("RATEX", "Binding break-even rate", f"IF($D${se.rows['RATEM']}=99,0,MAX(0,$D${se.rows['RATEM']}))", fmt="pct2",
              unit="%", name="BE_RATE_CF")
    for i in range(1, ANN_N + 1):
        se.ws.column_dimensions[pc(i)].width = 12.5
    return se


def build_checks(b):
    ck = Sheet(b, "60_Checks", "60 · Checks", "ማረጋገጫዎች", "B42318", kind="free", n=0, width_label=58)
    ws = ck.ws
    ck.sec("Model integrity checks – all must PASS; warnings are advisory", "የሞዴል ማረጋገጫዎች")
    hdr = ck.r
    ck.r += 1
    bsr = b.sheets["41_BS"].rows
    cfr = b.sheets["42_CF"].rows
    ann = lambda sh, k: f"'{sh}'!$E${b.sheets[sh].rows[k]}:$R${b.sheets[sh].rows[k]}"
    scan_sheets = ["10_Inputs", "11_Scenarios", "12_TaxLaw", "13_Spaces", "20_Timeline", "21_Area", "22_Capex", "23_Construction",
                   "24_Revenue", "25_Opex", "26_Assets", "27_Loan", "28_Tax", "29_VAT", "30_WorkingCap", "40_PL", "41_BS",
                   "42_CF", "43_Ratios", "50_Returns", "51_Payback", "52_Sensitivity"]
    def scan():
        parts = []
        for s in scan_sheets:
            shx = b.sheets[s]
            lastc = CL(max(20, 4 + (shx.n or 0)))
            parts.append(f"SUMPRODUCT(--ISERROR('{s}'!$A$1:${lastc}${shx.r + 5}))")
        return "+".join(parts)
    checks = [
        ("C01", "Balance sheet balances every year", f"MAX(MAX({ann('41_BS', 'CHK')}),-MIN({ann('41_BS', 'CHK')}))", 1, "E"),
        ("C02", "Cash-flow closing cash = balance-sheet cash", f"SUMPRODUCT(ABS({ann('42_CF', 'CLOSE')}-{ann('41_BS', 'CASH')}))", 1, "E"),
        ("C03", "Sources = uses (monthly funding waterfall)", "ABS(SRC_DIFF)", 1, "E"),
        ("C04", "TDC = monthly capex accrual; cost categories sum to TDC", "ABS(TDC_DIFF)+ABS(CAT_TOT-TDC)", 1, "E"),
        ("C05", "Capex fully allocated across IAS 2 / 40 / 16", "ABS(ALLOC_DIFF)+ABS(SHARE_SUM-1)*1000*IF(NET_AREA>0,1,0)", 1, "E"),
        ("C06", "Loan fully repaid within tenor", "ABS(LN_BAL_TENOR)", 1, "E"),
        ("C07", "Σ principal repaid = debt drawn + capitalised interest", "ABS(LN_PRIN_TOT-LN_DRAWS-LN_CAPINT)", 1, "E"),
        ("C08", "Inventory fully released by final sale (end O10)", lambda: f"ABS('26_Assets'!$R${b.sheets['26_Assets'].rows['INV']})", 1, "E"),
        ("C09", "Contract liabilities / assets unwind to zero (end O10)",
         lambda: f"ABS('24_Revenue'!$R${b.sheets['24_Revenue'].rows['CBAL']})", 1, "E"),
        ("C10", "Land not depreciated (cum. depreciation ≤ depreciable base)",
         lambda: f"MAX(0,'26_Assets'!$R${b.sheets['26_Assets'].rows['DEPC']}-(BASE_STR+BASE_SYS+BASE_FFE))", 1, "E"),
        ("C11", "Net area of the spaces ≤ gross floor area", "ABS(AREA_DIFF)", 0.5, "E"),
        ("C12", "Tax losses never negative", lambda: f"'28_Tax'!$D${b.sheets['28_Tax'].rows['LNEG']}", 0.5, "E"),
        ("C13", "Track, project type, data set and scenario switches valid",
         "IF(AND(TRACK>=1,TRACK<=2,SCEN>=1,SCEN<=4,PROJECT_MODE>=1,PROJECT_MODE<=3,DATA_SET>=1,DATA_SET<=3),0,1)", 0.5, "E"),
        ("C14", "No formula errors anywhere (REF, DIV/0, N/A, VALUE – ISERROR scan)", scan, 0.5, "E"),
        ("C15", "Payment plan and presales ≤ 100%", "MAX(0,PAY_DP+PAY_INST-1)+MAX(0,PRESALE_1+PRESALE_2+PRESALE_3+PRESALE_4-1)", 0.0001, "E"),
        ("C16", "VAT account rolls forward", lambda: f"ABS('29_VAT'!$D${b.sheets['29_VAT'].rows['CHK']})", 1, "E"),
        ("C17", "What-if dials at neutral (100% / 0) – otherwise results are not the base case", "FLEX_ON", 0.5, "W"),
        ("C18", "DSCR ≥ covenant in every repayment year", "N_BREACH", 0.5, "W"),
        ("C19", "Cash never below minimum in operating years", "N_LOWCASH", 0.5, "W"),
        ("C20", "Funding gap = 0 (equity + debt + presales cover the cost)", "FUND_GAP", 1, "W"),
        ("C21", "FAR within permitted limit", "IF(AND(FAR>0,FAR_ACT>FAR),1,0)", 0.5, "W"),
        ("C22", "Net-to-gross ratio plausible (50%–95%)", "IF(OR(GFA=0,NET_AREA=0),0,IF(AND(NTG>=0.5,NTG<=0.95),0,1))", 0.5, "W"),
        ("C23", "Hotel inputs consistent (hotel space ⇔ keys and room rate)", "HOTEL_INCONS", 0.5, "W"),
        ("C24", "At least one space entered on sheet 13", "IF(N_SPACES_USED>0,0,1)", 0.5, "W"),
        ("C25", "No IAS 36 impairment indicator", "N_IMPAIR", 0.5, "W"),
    ]
    b.CHECK_ROWS = []

    def _ck(_r):
        heads = ["", "Check", "Type", "Value", "Tolerance", "Status"]
        for j, h in enumerate(heads):
            if h:
                c = ws.cell(hdr, 1 + j, h)
                c.font = font(9, True, NAVY)
        for k, (cid, lab, f, tol, typ) in enumerate(checks):
            rr = hdr + 1 + k
            ws.cell(rr, 2, f"{cid}  {lab}").font = F_BODY
            ws.cell(rr, 3, "Error" if typ == "E" else "Warning").font = F_UNIT
            fv = f() if callable(f) else f
            vc = ws.cell(rr, 4, "=" + fv)
            vc.number_format = NF["n2"]
            ws.cell(rr, 5, tol).number_format = NF["n2"]
            st = ws.cell(rr, 6, f'=IF(ABS(D{rr})<=E{rr},"PASS",IF(C{rr}="Warning","WARN","FAIL"))')
            st.font = F_BOLD
            st.alignment = Alignment(horizontal="center")
            b.CHECK_ROWS.append(rr)
        first, last = hdr + 1, hdr + len(checks)
        rng = f"F{first}:F{last}"
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"PASS"'], fill=fill(GOOD_BG), font=Font(color=GOOD_FG, bold=True)))
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"WARN"'], fill=fill(WARN_BG), font=Font(color=WARN_FG, bold=True)))
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"FAIL"'], fill=fill(CRIT_BG), font=Font(color=CRIT_FG, bold=True)))
        s = last + 2
        for j, (lab, f, nm) in enumerate((("Checks passed", f'COUNTIF({rng},"PASS")', "CHK_PASS"),
                                          ("Errors (FAIL)", f'COUNTIF({rng},"FAIL")', "CHK_FAIL"),
                                          ("Warnings", f'COUNTIF({rng},"WARN")', "CHK_WARN"),
                                          ("Total checks", f"ROWS({rng})", "CHK_TOTAL"))):
            ws.cell(s + j, 2, lab).font = F_BOLD
            c = ws.cell(s + j, 4, "=" + f)
            c.font = F_BOLD
            c.fill = FILL_KEY
            b.name(nm, f"'60_Checks'!$D${s + j}")
        ws.cell(s + 5, 2, "Status").font = F_BOLD
        c = ws.cell(s + 5, 4, '=IF(CHK_FAIL>0,"ERRORS – see list",IF(CHK_WARN>0,"OK with warnings","ALL CHECKS PASS"))')
        c.font = F_BOLD
        b.name("CHK_STATUS", f"'60_Checks'!$D${s + 5}")
        ws.cell(s + 7, 2, "How to fix common failures: C01/C02 – never type over black formula cells; C06 – increase tenor or "
                          "reduce grace; C11 – net areas on sheet 13 exceed the floors on Inputs block 2; C15 – payment plan or "
                          "presales above 100%; C17 – set the dials on sheet 52 back to 100% / 0.").font = F_NOTE
        ws.column_dimensions["C"].width = 10
        ws.column_dimensions["D"].width = 16
        ws.column_dimensions["E"].width = 11
        ws.column_dimensions["F"].width = 12
    ck.raw(_ck)
    ck.r += len(checks) + 12
    return ck
