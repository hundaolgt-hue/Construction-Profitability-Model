"""Private residence track – separate inputs, calculations, one-page estimate and checks.

P1_Inputs   – all assumptions (My house / Sample), own cost library
P2_Calc     – building summary, element costs, 48-month cost & funding waterfall, home-loan schedule (vertical, 240 months)
P3_Estimate – one-page home estimate (print-ready)
P4_Checks   – integrity checks for the private track
"""
from openpyxl.styles import Alignment, Font, Border, Side
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.chart import BarChart, LineChart, Reference
from core import (axis_font, Sheet, pc, CL, NF, font, fill, F_BODY, F_BOLD, F_UNIT, F_NOTE, F_INPUT, FILL_INPUT, BOX,
                  NAVY, AMBER, BLUE, TEAL, VIOLET, ROSE, CANVAS, CARD, BORDER, INK, SEC, GOOD_FG, GOOD_BG, WARN_FG, WARN_BG,
                  CRIT_FG, CRIT_BG, FILL_KEY, AMFONT, BOTTOM_AMBER)
from private_def import P_INPUTS, P_EL, P_STAGES, P_FLOORS
from inputs_def import SAMPLE_NOTE

TAB_P = "2A9D8F"
NM = 48          # private monthly columns
NA = 240         # amortization months (20 years)


def _card(ws, r1, c1, r2, c2, bg=CARD, border=BORDER):
    s = Side(style="thin", color=border)
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(r, c)
            cell.fill = fill(bg)
            cell.border = Border(left=s if c == c1 else None, right=s if c == c2 else None,
                                 top=s if r == r1 else None, bottom=s if r == r2 else None)


# ====================================================================== P1 INPUTS
def build_p_inputs(b, blank=False):
    sh = Sheet(b, "P1_Inputs", "P1 · My house – inputs", "የግል መኖሪያ ቤት – ግብዓቶች", TAB_P, kind="free", n=0, width_label=48)
    ws = sh.ws
    b.P_INPUT_ROWS = {}
    r = 6
    heads = [("B", "Input"), ("C", "ግብዓት"), ("D", "Unit"), ("E", "MY HOUSE\n(your values)"),
             ("F", "SAMPLE\nG+1 house" if not blank else "SAMPLE\n(cleared)"), ("G", "ACTIVE\n(used)"), ("H", "Range name"),
             ("I", "Notes")]
    for col, h in heads:
        c = ws[f"{col}{r}"]
        c.value = h
        c.font = font(9, True, "FFFFFF")
        c.fill = fill(TAB_P if col == "E" else NAVY)
        c.alignment = Alignment(horizontal="center" if col in "EFG" else "left", vertical="center", wrap_text=True)
    ws.row_dimensions[r].height = 30
    ws["B7"] = ("Type your own values in column E (blue on amber). Column G uses the data set chosen on the Start sheet "
                "(1 = My house, 2 = Sample). The sample is an example only – \"Sample value — not market advice\".")
    ws["B7"].font = F_NOTE
    r = 9
    for item in P_INPUTS:
        if item[0] == "sec":
            ws.cell(r, 2, item[1]).font = font(11, True, NAVY)
            ws.cell(r, 3, item[2]).font = font(10, False, SEC, name=AMFONT)
            for col in range(2, 10):
                ws.cell(r, col).border = BOTTOM_AMBER
            ws.row_dimensions[r].height = 20
            r += 1
            continue
        key, en, am, unit, fmt, uv, sv, note, dv = item
        if blank:
            sv = uv
        ws.cell(r, 2, en).font = F_BODY
        ws.cell(r, 3, am).font = font(9, False, SEC, name=AMFONT)
        ws.cell(r, 4, unit).font = F_UNIT
        for col, v in ((5, uv), (6, sv)):
            c = ws.cell(r, col, v)
            c.font = F_INPUT
            c.fill = FILL_INPUT
            c.border = BOX
            c.number_format = NF.get(fmt, fmt)
            c.alignment = Alignment(horizontal="right" if fmt not in ("gen", "txt") else "left")
            if col == 6 and not blank and v not in (None, "", 0) and fmt != "txt":
                c.comment = Comment(SAMPLE_NOTE, "Rapha")
        ch = f"CHOOSE(P_DATA,E{r},F{r})"
        if isinstance(dv, str) and dv.startswith("L_"):
            f, hf = f"IFERROR(MATCH({ch},{dv},0),1)", "int"
        elif dv == "yn":
            f, hf = f'IF({ch}="Yes",1,0)', "int"
        elif dv == "txt":
            f, hf = f'{ch}&""', "gen"
        else:
            f, hf = ch, fmt
        c = ws.cell(r, 7, "=" + f)
        c.number_format = NF.get(hf, hf)
        c.font = F_BOLD
        c.fill = fill("EEF1F5")
        c.border = BOX
        ws.cell(r, 8, key).font = font(8, False, SEC)
        ws.cell(r, 9, note).font = F_NOTE
        b.name(key, f"'P1_Inputs'!$G${r}")
        b.P_INPUT_ROWS[key] = r
        rng = f"E{r}:F{r}"
        if isinstance(dv, str) and dv.startswith("L_"):
            b.DV.append(("P1_Inputs", "list", dv, rng))
        elif dv == "yn":
            b.DV.append(("P1_Inputs", "list", "L_YN", rng))
        elif isinstance(dv, tuple):
            b.DV.append(("P1_Inputs", dv[0], dv, rng))
        r += 1
    # cost library
    r += 1
    ws.cell(r, 2, "8. Cost library for houses – ETB per m² gross, budget-date prices").font = font(11, True, NAVY)
    ws.cell(r, 3, "8. የቤት ግንባታ ወጪ መረጃ").font = font(10, False, SEC, name=AMFONT)
    for col in range(2, 10):
        ws.cell(r, col).border = BOTTOM_AMBER
    r += 1
    for col, h in (("E", "Basic"), ("F", "Standard"), ("G", "Premium")):
        cc = ws[f"{col}{r}"]
        cc.value = h
        cc.font = font(9, True, "FFFFFF")
        cc.fill = fill(NAVY)
        cc.alignment = Alignment(horizontal="center")
    ws.cell(r, 9, "Applies to both data sets. Update with current local prices.").font = F_NOTE
    r += 1
    lib0 = r
    b.P_LIB_ROW = {}
    for j, (k, en, am, rates) in enumerate(P_EL):
        ws.cell(r, 2, en).font = F_BODY
        ws.cell(r, 3, am).font = font(9, False, SEC, name=AMFONT)
        ws.cell(r, 4, "ETB/m²").font = F_UNIT
        for jj, v in enumerate(rates):
            cc = ws.cell(r, 5 + jj, v)
            cc.font = F_INPUT
            cc.fill = FILL_INPUT
            cc.border = BOX
            cc.number_format = NF["n0"]
            if not blank:
                cc.comment = Comment(SAMPLE_NOTE, "Rapha")
            b.DV.append(("P1_Inputs", "dec", ("dec", 0, 1e7), f"{CL(5 + jj)}{r}"))
        ws.cell(r, 8, "P_LIB").font = font(8, False, SEC)
        b.P_LIB_ROW[k] = r
        r += 1
    b.name("P_LIB", f"'P1_Inputs'!$E${lib0}:$G${r - 1}")
    ws.cell(r, 2, "Total library rate").font = F_BOLD
    for jj in range(3):
        cc = ws.cell(r, 5 + jj, f"=SUM({CL(5 + jj)}{lib0}:{CL(5 + jj)}{r - 1})")
        cc.number_format = NF["n0"]
        cc.font = F_BOLD
    r += 2
    for key, lab, v, note in (("P_LOCF_ADDIS", "Location factor – Addis Ababa", 1.0, "Base = 1.00"),
                              ("P_LOCF_REG", "Location factor – regional town", 0.92, "Lower labour cost, higher transport")):
        ws.cell(r, 2, lab).font = F_BODY
        ws.cell(r, 4, "x").font = F_UNIT
        cc = ws.cell(r, 5, v)
        cc.font = F_INPUT
        cc.fill = FILL_INPUT
        cc.border = BOX
        cc.number_format = NF["n2"]
        ws.cell(r, 8, key).font = font(8, False, SEC)
        ws.cell(r, 9, note).font = F_NOTE
        b.name(key, f"'P1_Inputs'!$E${r}")
        b.DV.append(("P1_Inputs", "dec", ("dec", 0.1, 5), f"E{r}"))
        r += 1
    for col, w in (("C", 28), ("D", 10), ("E", 16), ("F", 16), ("G", 15), ("H", 14), ("I", 46)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "E9"
    sh.r = r + 1
    return sh


# ====================================================================== P2 CALCULATIONS
def build_p_calc(b):
    sh = Sheet(b, "P2_Calc", "P2 · My house – calculations", "የቤት ወጪ፣ የገንዘብ ምንጭና የብድር ስሌት", TAB_P, kind="free", n=NM,
               width_label=46)
    ws = sh.ws
    m = lambda k, i: sh.c(k, i)
    mp = lambda k, i: (sh.c(k, i - 1) if i > 1 else "0")
    sh.sec("Building summary", "የሕንፃ ማጠቃለያ")
    sh.scalar("GFA", "Gross floor area", "SUM(" + ",".join(f"P_FL{i}" for i in range(1, len(P_FLOORS) + 1)) + ")", unit="m²",
              name="P_GFA", style="key")
    sh.scalar("GB", "Basement area", "P_FL1", unit="m²", name="P_GFA_BASE")
    sh.scalar("NA", "Above-ground floors built", "(P_FL2>0)+(P_FL3>0)+(P_FL4>0)+(P_FL5>0)", fmt="int", unit="no.",
              name="P_N_ABOVE")
    sh.scalar("CFG", "Configuration found", 'IF(P_FL1>0,"B+","")&"G+"&MAX(0,P_N_ABOVE-1)', fmt="gen", unit="",
              name="P_CONFIG")
    sh.scalar("BOK", "Building type matches floors (1 = yes)",
              "IF(AND(P_N_ABOVE=P_BTYPE,IF(P_BASE_YN=1,P_FL1>0,P_FL1=0)),1,0)", fmt="int", unit="flag", name="P_BTYPE_OK")
    sh.scalar("LOCF", "Location factor", "IF(P_LOC_TYPE=1,P_LOCF_ADDIS,P_LOCF_REG)", fmt="n2", unit="x", name="P_LOCF")
    sh.gap()
    sh.sec("Building cost at today's prices (library rate × area × location)", "የሕንፃ ወጪ በዛሬ ዋጋ")
    for j, (k, en, am, _) in enumerate(P_EL):
        sh.scalar(f"EL_{k}", f"{en}  ·  {am}", f"INDEX(P_LIB,{j + 1},P_FINISH)*P_GFA*P_LOCF", name=f"P_EL_{k}")
    sh.scalar("XB", "Basement extra", "P_GFA_BASE*P_BASE_RATE*P_LOCF", name="P_X_BASE")
    sh.scalar("XC", "Compound wall", "P_COMP_YN*P_COMP_LEN*P_COMP_RATE", name="P_X_COMP")
    sh.scalar("XS", "Septic tank", "P_SEPT_YN*P_SEPT_RATE", name="P_X_SEPT")
    sh.scalar("XW", "Water tanks", "P_WTANK_QTY*P_WTANK_RATE", name="P_X_WTANK")
    sh.scalar("XP", "Paving / garden", "P_PAVE_M2*P_PAVE_RATE", name="P_X_PAVE")
    sh.scalar("XEXT", "External works (wall + septic + tanks + paving)", "P_X_COMP+P_X_SEPT+P_X_WTANK+P_X_PAVE", name="P_EL_EXT")
    sh.scalar("BB", "BUILDING COST AT TODAY'S PRICES",
              "+".join(f"P_EL_{k}" for k, *_ in P_EL) + "+P_X_BASE+P_EL_EXT", name="P_BUILD_BASE", style="key")
    sh.gap()
    sh.sec("Month by month (columns = construction months)", "ወር በወር")
    hdr = sh.r
    sh.r += 1

    def _mh(_r):
        ws.cell(hdr, 2, "Month").font = font(9, True, SEC)
        ws.cell(hdr, 4, "Total").font = font(9, True, NAVY)
        for i in range(1, NM + 1):
            c = ws.cell(hdr, 4 + i, f"M{i}")
            c.font = font(9, True, NAVY)
            c.alignment = Alignment(horizontal="right")
            c.border = Border(bottom=Side(style="thin", color=NAVY))
    sh.raw(_mh)
    sh.line("T", "Month number", lambda i: i, fmt="int", unit="no.", tot=None)
    sh.line("DATE", "Month starting", lambda i: f"EDATE(P_START,{i - 1})", fmt="mon", unit="date", tot=None)
    sh.line("CF", "Construction month (1 = yes)", lambda i: f"IF({i}<=P_MONTHS,1,0)", fmt="int", unit="flag")
    sh.line("CUM", "Cumulative progress",
            lambda i: (f"IF({i}>=P_MONTHS,1,IF(P_SCURVE=1,({i}/P_MONTHS)^2/(({i}/P_MONTHS)^2+(1-{i}/P_MONTHS)^2),{i}/P_MONTHS))"),
            fmt="pct", unit="%", tot=None)
    sh.line("S", "Work done this month", lambda i: f"{m('CUM', i)}-{mp('CUM', i)}", fmt="pct", unit="%")
    sh.line("ESC", "Price escalation factor", lambda i: f"(1+P_ESC)^(({i}-0.5)/12)", fmt="n2", unit="x", tot=None)
    sh.line("BUILD", "Building work (incl. price rises)", lambda i: f"P_BUILD_BASE*{m('S', i)}*{m('ESC', i)}")
    sh.line("CONT", "Contingency", lambda i: f"{m('BUILD', i)}*P_CONT")
    sh.line("DES", "Design & drawings", lambda i: "P_DESIGN*P_BUILD_BASE" if i == 1 else 0)
    sh.line("SUP", "Site supervision", lambda i: f"P_SUPERV*{m('BUILD', i)}")
    sh.line("PER", "Permits & connections", lambda i: "P_PERMITS" if i == 1 else 0)
    sh.line("FUR", "Furniture & appliances", lambda i: f"IF({i}=P_MONTHS,P_FURN,0)")
    sh.line("LAND", "Land / lease payment", lambda i: "IF(P_LAND_TYPE=2,P_LAND_COST,0)" if i == 1 else 0)
    sh.line("ARR", "Loan arrangement fee", lambda i: "IF(P_LOAN_MAX>0,P_ARR*P_LOAN_MAX,0)" if i == 1 else 0)
    sh.line("COST", "TOTAL COST THIS MONTH (excl. interest)",
            lambda i: "+".join(m(k, i) for k in ("BUILD", "CONT", "DES", "SUP", "PER", "FUR", "LAND", "ARR")), style="total")
    sh.line("RETW", "Retention held back from contractor", lambda i: f"P_RETENTION*{m('BUILD', i)}")
    sh.line("RETS", "Retention set aside at handover", lambda i: f"IF({i}=P_MONTHS,P_RET_TOT,0)")
    sh.line("NEED", "CASH NEEDED THIS MONTH", lambda i: f"{m('COST', i)}-{m('RETW', i)}+{m('RETS', i)}", style="key")
    sh.sec("Who pays: savings → family → home loan → gap", "ማን ይከፍላል")
    sh.line("SAV", "(1) Own savings", lambda i: f"MIN({m('NEED', i)},MAX(0,P_SAVINGS-{mp('CSAV', i)}))")
    sh.line("FAM", "(2) Family / diaspora", lambda i: f"MIN({m('NEED', i)}-{m('SAV', i)},MAX(0,P_FAMILY-{mp('CFAM', i)}))")
    sh.line("LOAN", "(3) Home loan draw", lambda i: f"MIN({m('NEED', i)}-{m('SAV', i)}-{m('FAM', i)},MAX(0,P_LOAN_MAX-{mp('CLOAN', i)}))")
    sh.line("GAP", "(4) Still to find (funding gap)", lambda i: f"{m('NEED', i)}-{m('SAV', i)}-{m('FAM', i)}-{m('LOAN', i)}")
    sh.line("CSAV", "Cumulative savings used", lambda i: f"{mp('CSAV', i)}+{m('SAV', i)}", tot="last")
    sh.line("CFAM", "Cumulative family money used", lambda i: f"{mp('CFAM', i)}+{m('FAM', i)}", tot="last")
    sh.line("CLOAN", "Cumulative loan drawn", lambda i: f"{mp('CLOAN', i)}+{m('LOAN', i)}", tot="last")
    sh.line("SRC", "Check: sources − cash needed", lambda i: f"{m('SAV', i)}+{m('FAM', i)}+{m('LOAN', i)}+{m('GAP', i)}-{m('NEED', i)}")
    sh.sec("Home loan while building", "በግንባታ ወቅት ብድር")
    sh.line("LOPEN", "Opening loan balance", lambda i: mp("LCLOSE", i), tot=None)
    sh.line("LINT", "Interest for the month", lambda i: f"{m('CF', i)}*({m('LOPEN', i)}+{m('LOAN', i)})*P_RATE/12")
    sh.line("LCAP", "Interest added to the loan", lambda i: f"IF(P_IDC=2,{m('LINT', i)},0)")
    sh.line("LPAID", "Interest you pay from income", lambda i: f"{m('LINT', i)}-{m('LCAP', i)}")
    sh.line("LCLOSE", "Closing loan balance", lambda i: f"{m('LOPEN', i)}+{m('LOAN', i)}+{m('LCAP', i)}", tot="max",
            style="total")
    sh.gap()
    sh.sec("Totals", "ድምር")
    sh.scalar("BE", "Building work incl. price rises", lambda: f"$D${sh.rows['BUILD']}", name="P_BUILD_ESC")
    sh.scalar("CO", "Contingency", lambda: f"$D${sh.rows['CONT']}", name="P_CONT_TOT")
    sh.scalar("DS", "Design", lambda: f"$D${sh.rows['DES']}", name="P_DES_TOT")
    sh.scalar("SU", "Supervision", lambda: f"$D${sh.rows['SUP']}", name="P_SUP_TOT")
    sh.scalar("PE", "Permits", lambda: f"$D${sh.rows['PER']}", name="P_PER_TOT")
    sh.scalar("FU", "Furniture", lambda: f"$D${sh.rows['FUR']}", name="P_FUR_TOT")
    sh.scalar("LA", "Land", lambda: f"$D${sh.rows['LAND']}", name="P_LAND_TOT")
    sh.scalar("AR", "Arrangement fee", lambda: f"$D${sh.rows['ARR']}", name="P_ARR_TOT")
    sh.scalar("IN", "Interest while building (capitalised + paid)", lambda: f"$D${sh.rows['LINT']}", name="P_IDC_TOT")
    sh.scalar("INC", "  of which added to the loan", lambda: f"$D${sh.rows['LCAP']}", name="P_IDC_CAP")
    sh.scalar("INP", "Highest monthly interest paid from income while building", lambda: f"MAX(E{sh.rows['LPAID']}:{pc(NM)}{sh.rows['LPAID']})",
              name="P_IDC_PAID_MAX")
    sh.scalar("RT", "Retention (paid after the defects period)", "P_RETENTION*P_BUILD_ESC", name="P_RET_TOT")
    sh.scalar("CC", "CONSTRUCTION COST (building, contingency, fees, furniture)",
              "P_BUILD_ESC+P_CONT_TOT+P_DES_TOT+P_SUP_TOT+P_PER_TOT+P_FUR_TOT", name="P_CONSTR", style="key")
    sh.scalar("TC", "TOTAL COST OF THE HOUSE (incl. land, loan fees & interest while building)",
              "P_CONSTR+P_LAND_TOT+P_ARR_TOT+P_IDC_TOT", name="P_TOTAL", style="key")
    sh.scalar("CM2", "Construction cost per m² gross", "IF(P_GFA>0,P_CONSTR/P_GFA,0)", unit="ETB/m²", name="P_CC_M2")
    sh.scalar("TLO", "Low estimate (construction cost)", "P_CONSTR*(1+P_EST_LOW)", name="P_LOW")
    sh.scalar("THI", "High estimate (construction cost)", "P_CONSTR*(1+P_EST_HIGH)", name="P_HIGH")
    sh.scalar("NEEDT", "Cash needed in total", lambda: f"$D${sh.rows['NEED']}", name="P_NEED_TOT")
    sh.scalar("SAVT", "Savings used", lambda: f"$D${sh.rows['SAV']}", name="P_SAV_USED")
    sh.scalar("FAMT", "Family money used", lambda: f"$D${sh.rows['FAM']}", name="P_FAM_USED")
    sh.scalar("LT", "Loan drawn", lambda: f"$D${sh.rows['LOAN']}", name="P_LOAN_DRAWN")
    sh.scalar("GT", "Funding gap", lambda: f"$D${sh.rows['GAP']}", name="P_GAP", style="key")
    sh.scalar("PK", "Peak month of spending", lambda: f"MAX(E{sh.rows['NEED']}:{pc(NM)}{sh.rows['NEED']})", name="P_PEAK_M")
    sh.scalar("SRCT", "Check: sources − uses", lambda: f"SUMPRODUCT(ABS(E{sh.rows['SRC']}:{pc(NM)}{sh.rows['SRC']}))", name="P_SRC_DIFF")
    sh.scalar("B0", "Loan balance at completion", lambda: f"INDEX(E{sh.rows['LCLOSE']}:{pc(NM)}{sh.rows['LCLOSE']},P_MONTHS)",
              name="P_B0", style="key")
    sh.gap()
    sh.sec("Cost breakdown (sums to total cost)", "የወጪ ክፍፍል")
    cats = [("C1", "Foundation & basement", "መሠረትና ምድር ቤት", "P_EL_SUBSTR+P_X_BASE"),
            ("C2", "Frame, walls & roof", "ፍሬም፣ ግድግዳና ጣሪያ", "P_EL_SUPER+P_EL_ENVEL"),
            ("C3", "Internal walls, finishes & kitchen", "የውስጥ ሥራ፣ ማጠናቀቂያና ማዕድ ቤት", "P_EL_INTERN+P_EL_FINISH+P_EL_KITCH"),
            ("C4", "Electrical & plumbing", "ኤሌክትሪክና ቧንቧ", "P_EL_MEP"),
            ("C5", "Site set-up & external works", "ቅድመ-ሥራና የውጭ ሥራ", "P_EL_PRELIM+P_EL_EXT"),
            ("C6", "Price rises & contingency", "የዋጋ ጭማሪና መጠባበቂያ", "P_BUILD_ESC-P_BUILD_BASE+P_CONT_TOT"),
            ("C7", "Design, supervision & permits", "ዲዛይን፣ ቁጥጥርና ፈቃድ", "P_DES_TOT+P_SUP_TOT+P_PER_TOT"),
            ("C8", "Furniture & appliances", "የቤት ዕቃዎች", "P_FUR_TOT"),
            ("C9", "Land", "መሬት", "P_LAND_TOT"),
            ("C10", "Loan fee & interest while building", "የብድር ክፍያና ወለድ", "P_ARR_TOT+P_IDC_TOT")]
    b.P_CATS = []
    for k, en, am, f in cats:
        b.P_CATS.append((sh.scalar(k, en, f), am))
    sh.scalar("CT", "Total (= total cost)", f"SUM(D{b.P_CATS[0][0]}:D{b.P_CATS[-1][0]})", name="P_CAT_TOT", style="key")
    sh.gap()
    sh.sec("Paying the contractor – stages (building work incl. price rises & contingency)", "ለኮንትራክተር የክፍያ ደረጃዎች")
    b.P_STAGE_ROWS = []
    for lab, am, keys in P_STAGES:
        parts = "+".join({"BASE": "P_X_BASE", "EXT": "P_EL_EXT"}.get(k, f"P_EL_{k}") for k in keys)
        b.P_STAGE_ROWS.append((sh.scalar("ST" + lab[0], lab, f"IF(P_BUILD_BASE>0,({parts})/P_BUILD_BASE,0)*(P_BUILD_ESC+P_CONT_TOT)"),
                               am))
    sh.scalar("STT", "Total contractor payments", f"SUM(D{b.P_STAGE_ROWS[0][0]}:D{b.P_STAGE_ROWS[-1][0]})", name="P_STAGE_TOT",
              style="key")
    b.P2_CUM = sh.rows  # used by the estimate (row keys)
    sh.gap()
    # ------------------------------------------------------------ loan after completion – vertical schedule
    sh.sec("Home-loan schedule after completion (one row per month)", "ከግንባታ በኋላ የብድር ክፍያ ሰንጠረዥ (በወር)")
    sh.scalar("NPAY", "Monthly payments after the grace period", "MAX(0,P_TENOR*12-P_GRACE)", fmt="int", unit="no.", name="P_NPAY")
    sh.scalar("INST", "Monthly instalment (after grace)", "IF(AND(P_B0>0,P_NPAY>0),PMT(P_RATE/12,P_NPAY,-P_B0),0)",
              name="P_INSTAL", style="key")
    sh.scalar("GINT", "Monthly interest-only payment during grace", "P_B0*P_RATE/12", name="P_GRACE_PAY")
    sh.scalar("PAYD", "Loan repaid by", "IF(P_B0>0,EDATE(P_START,P_MONTHS+P_TENOR*12)-1,0)", fmt="date", unit="date",
              name="P_PAYOFF")
    sh.scalar("DTI", "Debt payments ÷ income", "IF(P_INCOME>0,(P_INSTAL+P_OTHER_DEBT)/P_INCOME,0)", fmt="pct", unit="%",
              name="P_DTI", style="key")
    sh.scalar("AFF", "Affordability", 'IF(P_INCOME<=0,"enter income",IF(P_DTI<=P_AFF_MAX,"Affordable",IF(P_DTI<=P_AFF_MAX*1.25,"Stretch","Not affordable")))',
              fmt="gen", unit="", name="P_AFF_TXT")
    sh.scalar("AFFA", "ውሳኔ", 'IF(P_INCOME<=0,"ገቢ ያስገቡ",IF(P_DTI<=P_AFF_MAX,"ይቻላል",IF(P_DTI<=P_AFF_MAX*1.25,"ከባድ","አይቻልም")))',
              fmt="gen", unit="", name="P_AFF_AM")
    sh.scalar("MAXL", "Largest loan the income supports (at the max %)",
              "IF(AND(P_INCOME>0,P_NPAY>0,P_RATE>0),PV(P_RATE/12,P_NPAY,-MAX(0,P_AFF_MAX*P_INCOME-P_OTHER_DEBT)),0)",
              name="P_LOAN_AFF")
    a0 = sh.r + 1
    b.P_AM0 = a0
    sh.r = a0 + 1 + NA + 1

    def _am(_r):
        heads = ["Month after completion", "Date", "Loan year", "Opening balance", "Interest", "Payment", "Principal",
                 "Closing balance"]
        for j, h in enumerate(heads):
            c = ws.cell(a0, 2 + j, h)
            c.font = font(9, True, "FFFFFF")
            c.fill = fill(NAVY)
            c.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.row_dimensions[a0].height = 28
        for k in range(1, NA + 1):
            r = a0 + k
            vals = {2: k, 3: f"=EDATE(P_START,P_MONTHS+{k - 1})", 4: f"=ROUNDUP(B{r}/12,0)",
                    5: "=P_B0" if k == 1 else f"=I{r - 1}",
                    6: f"=E{r}*P_RATE/12",
                    7: f"=IF(B{r}>P_TENOR*12,0,IF(B{r}<=P_GRACE,F{r},IF(B{r}=P_TENOR*12,E{r}+F{r},MIN(E{r}+F{r},P_INSTAL))))",
                    8: f"=G{r}-F{r}", 9: f"=E{r}-H{r}"}
            for col, v in vals.items():
                c = ws.cell(r, col, v)
                c.number_format = NF["int"] if col in (2, 4) else (NF["mon"] if col == 3 else NF["n0"])
                c.font = font(9)
    sh.raw(_am)
    b.name("P_AM_K", f"'P2_Calc'!$B${a0 + 1}:$B${a0 + NA}")
    b.name("P_AM_Y", f"'P2_Calc'!$D${a0 + 1}:$D${a0 + NA}")
    b.name("P_AM_INT", f"'P2_Calc'!$F${a0 + 1}:$F${a0 + NA}")
    b.name("P_AM_PAY", f"'P2_Calc'!$G${a0 + 1}:$G${a0 + NA}")
    b.name("P_AM_PRIN", f"'P2_Calc'!$H${a0 + 1}:$H${a0 + NA}")
    b.name("P_AM_CLOSE", f"'P2_Calc'!$I${a0 + 1}:$I${a0 + NA}")
    sh.sec("Loan totals", "የብድር ድምር")
    sh.scalar("PRT", "Principal repaid in total", "SUM(P_AM_PRIN)", name="P_PRIN_TOT")
    sh.scalar("INTT", "Interest after completion", "SUM(P_AM_INT)-SUMIFS(P_AM_INT,P_AM_K,\">\"&P_TENOR*12)", name="P_INT_AFTER")
    sh.scalar("INTA", "Interest – total over the life of the loan", "P_INT_AFTER+P_IDC_TOT", name="P_INT_ALL", style="key")
    sh.scalar("BEND", "Balance left at end of the term", "INDEX(P_AM_CLOSE,P_TENOR*12)", name="P_BAL_END")
    sh.gap()
    sh.sec("Optional – renting out part of the house", "አማራጭ – ከፊሉን ማከራየት")
    sh.scalar("RG", "Rent in first year after completion (gross)", "IF(P_RENT_YN=1,P_RENT_M*12*(1+P_RENT_ESC)^(P_MONTHS/12),0)",
              name="P_RENT_Y1")
    sh.scalar("RN", "Rent after vacancy, running costs & tax", "P_RENT_Y1*(1-P_VACANCY)*(1-P_RENT_COST)", name="P_RENT_NET")
    sh.scalar("RA", "Rented area ÷ house area", "IF(P_GFA>0,P_RENT_AREA/P_GFA,0)", fmt="pct", unit="%", name="P_RENT_SH")
    sh.scalar("RCOST", "Cost of the rented part (total cost × rented share of area)", "P_TOTAL*MIN(1,P_RENT_SH)", name="P_RENT_PARTCOST")
    sh.scalar("RY", "Gross yield on the cost of the rented part", "IF(P_RENT_PARTCOST>0,P_RENT_Y1/P_RENT_PARTCOST,0)", fmt="pct", unit="%",
              name="P_RENT_YIELD")
    sh.scalar("RP", "Simple payback of the rented part (years)", "IF(P_RENT_NET>0,P_RENT_PARTCOST/P_RENT_NET,0)", fmt="n1",
              unit="years", name="P_RENT_PB")
    sh.scalar("RC", "Net rent ÷ monthly instalment", "IF(P_INSTAL>0,P_RENT_NET/12/P_INSTAL,0)", fmt="pct", unit="%",
              name="P_RENT_COVER")
    # annual loan summary for the estimate
    sh.gap()
    sh.sec("Loan by year (for the estimate page)", "ብድር በዓመት")
    y0 = sh.r
    b.P_Y0 = y0
    sh.r = y0 + 1 + 20 + 1

    def _yr(_r):
        for j, h in enumerate(["Year", "Interest", "Principal", "Payments", "Closing balance"]):
            c = ws.cell(y0, 2 + j, h)
            c.font = font(9, True, NAVY)
        for y in range(1, 21):
            r = y0 + y
            ws.cell(r, 2, y).number_format = NF["int"]
            for col, rng in ((3, "P_AM_INT"), (4, "P_AM_PRIN"), (5, "P_AM_PAY")):
                c = ws.cell(r, col, f"=SUMIFS({rng},P_AM_Y,B{r})")
                c.number_format = NF["n0"]
            c = ws.cell(r, 6, f"=INDEX(P_AM_CLOSE,12*B{r})")
            c.number_format = NF["n0"]
    sh.raw(_yr)
    for i in range(1, NM + 1):
        ws.column_dimensions[pc(i)].width = 11.5
    return sh


# ====================================================================== P3 ESTIMATE
def build_p_estimate(b):
    sh = Sheet(b, "P3_Estimate", "My house – estimate", "የቤት ግንባታ ወጪ ግምት", "0F2A44", kind="free", n=0, width_label=27)
    ws = sh.ws
    p2 = b.sheets["P2_Calc"]

    def hdr_row(r, cols, labels):
        for c, h in zip(cols, labels):
            x = ws.cell(r, c, h)
            x.font = font(9, True, "FFFFFF")
            x.fill = fill(NAVY)
            x.alignment = Alignment(horizontal="left" if c in (2, 7) else "right")

    def title(r, c, t):
        ws.cell(r, c, t).font = font(11, True, NAVY)

    def _e(_r):
        for col, w in (("A", 2), ("B", 30), ("C", 14), ("D", 10), ("E", 9), ("F", 2), ("G", 9), ("H", 13), ("I", 13),
                       ("J", 13), ("K", 2)):
            ws.column_dimensions[col].width = w
        info = [(6, "Project · ፕሮጀክት", "=P_NAME"),
                (7, "Building · ሕንፃ", '=P_CONFIG&"  ·  "&TEXT(P_GFA,"#,##0")&" m² gross  ·  "&INDEX(L_FINISH,P_FINISH)&" finish"'),
                (8, "Where & how long · ቦታና ጊዜ", '=P_LOCATION&"  ·  "&P_MONTHS&" months from "&TEXT(P_START,"mmm yyyy")')]
        for r, lab, f in info:
            ws.cell(r, 2, lab).font = font(9, True, SEC)
            ws.cell(r, 3, f).font = font(10, True, NAVY)
        tiles = [((2, 2), "CONSTRUCTION COST · የግንባታ ወጪ", "=P_CONSTR", "n0",
                  '="Range "&TEXT(P_LOW/1000000,"#,##0.0")&" – "&TEXT(P_HIGH/1000000,"#,##0.0")&" m"'),
                 ((3, 5), "TOTAL COST · ጠቅላላ ወጪ", "=P_TOTAL", "n0",
                  '="per m² building: "&TEXT(P_CC_M2,"#,##0")'),
                 ((7, 8), "MONTHLY INSTALMENT · ወርሃዊ ክፍያ", "=P_INSTAL", "n0",
                  '=IF(P_INCOME>0,TEXT(P_DTI,"0%")&" of income · "&P_AFF_TXT,"enter income in P1")'),
                 ((9, 10), "FUNDING GAP · የገንዘብ ክፍተት", "=P_GAP", "n0",
                  '=IF(P_GAP>0.5,"find more funds","covered · በቂ ነው")')]
        accents = [BLUE, TEAL, VIOLET, AMBER]
        for k, ((c1, c2), lab, f, fmt, sub) in enumerate(tiles):
            _card(ws, 10, c1, 13, c2, bg=CANVAS)
            for cc in range(c1, c2 + 1):
                ws.cell(10, cc).border = Border(top=Side(style="thick", color=accents[k]))
            if c2 > c1:
                for rr in (10, 11, 13):
                    ws.merge_cells(start_row=rr, start_column=c1, end_row=rr, end_column=c2)
            ws.cell(10, c1, lab).font = font(8, True, SEC)
            v = ws.cell(11, c1, f)
            v.font = font(16, True, NAVY)
            v.number_format = NF[fmt]
            v.alignment = Alignment(horizontal="left")
            ws.cell(13, c1, sub).font = font(8, False, SEC)
        ws.row_dimensions[11].height = 24
        # cost breakdown
        title(15, 2, "COST BREAKDOWN · የወጪ ክፍፍል")
        title(15, 7, "WHERE THE MONEY GOES")
        hdr_row(16, (2, 3, 4), ("Item", "ETB", "% total"))
        for k, (cr, am) in enumerate(b.P_CATS):
            r = 17 + k
            ws.cell(r, 2, f"='P2_Calc'!B{cr}").font = font(9)
            ws.cell(r, 5, am).font = font(7, False, SEC, name=AMFONT)
            c = ws.cell(r, 3, f"='P2_Calc'!D{cr}")
            c.number_format = NF["n0"]
            c.font = font(9)
            c2 = ws.cell(r, 4, f"=IF(P_TOTAL>0,C{r}/P_TOTAL,0)")
            c2.number_format = NF["pct"]
            c2.font = font(9)
        tr = 17 + len(b.P_CATS)
        ws.cell(tr, 2, "TOTAL COST · ጠቅላላ").font = F_BOLD
        c = ws.cell(tr, 3, f"=SUM(C17:C{tr - 1})")
        c.number_format = NF["n0"]
        c.font = F_BOLD
        c.fill = FILL_KEY
        b.name("P_EST_TOTAL", f"'P3_Estimate'!$C${tr}")
        for col in (2, 3, 4, 5):
            ws.cell(tr, col).border = Border(top=Side(style="thin", color=INK))
        ch = BarChart()
        ch.type = "bar"
        ch.style = 10
        ch.legend = None
        ch.y_axis.numFmt = '#,##0.0,,"m"'
        ch.y_axis.majorGridlines = None
        ch.add_data(Reference(ws, min_col=3, min_row=17, max_row=tr - 1), titles_from_data=False)
        ch.set_categories(Reference(ws, min_col=2, min_row=17, max_row=tr - 1))
        ch.series[0].graphicalProperties.solidFill = BLUE
        ch.series[0].graphicalProperties.line.solidFill = BLUE
        ch.x_axis.scaling.orientation = "maxMin"
        ch.height, ch.width = 6.2, 9.6
        ch.gapWidth = 40
        axis_font(ch.x_axis, 650)
        axis_font(ch.y_axis, 650)
        ws.add_chart(ch, "G16")
        # stages
        r = tr + 2
        title(r, 2, "PAYING THE CONTRACTOR · ለኮንትራክተር ክፍያ")
        title(r, 7, "CASH NEEDED EACH MONTH")
        hdr_row(r + 1, (2, 3, 4, 5), ("Stage", "ETB", "Cum. %", "≈ month"))
        st0 = r + 2
        cumrow = f"'P2_Calc'!$E${p2.rows['CUM']}:${pc(NM)}${p2.rows['CUM']}"
        for k, (sr, am) in enumerate(b.P_STAGE_ROWS):
            rr = st0 + k
            ws.cell(rr, 2, f"='P2_Calc'!B{sr}").font = font(9)
            c = ws.cell(rr, 3, f"='P2_Calc'!D{sr}")
            c.number_format = NF["n0"]
            c.font = font(9)
            c2 = ws.cell(rr, 4, f"=IF(P_STAGE_TOT>0,SUM($C${st0}:C{rr})/P_STAGE_TOT,0)")
            c2.number_format = NF["pct"]
            c2.font = font(9)
            c3 = ws.cell(rr, 5, f'=MIN(P_MONTHS,COUNTIF({cumrow},"<"&(D{rr}-0.0001))+1)')
            c3.number_format = NF["int"]
            c3.font = font(9)
        rr = st0 + len(b.P_STAGE_ROWS)
        ws.cell(rr, 2, "Retention paid after defects period").font = F_NOTE
        c = ws.cell(rr, 3, "=P_RET_TOT")
        c.number_format = NF["n0"]
        c.font = font(9)
        c = ws.cell(rr, 5, "=P_MONTHS+P_RET_M")
        c.number_format = NF["int"]
        c.font = font(9)
        ch2 = BarChart()
        ch2.style = 10
        ch2.y_axis.numFmt = '#,##0.0,,"m"'
        ch2.y_axis.majorGridlines = None
        pws = b.wb["P2_Calc"]
        ch2.add_data(Reference(pws, min_col=5, max_col=4 + NM, min_row=p2.rows["NEED"]), from_rows=True, titles_from_data=False)
        ch2.set_categories(Reference(pws, min_col=5, max_col=4 + NM, min_row=p2.rows["T"]))
        ch2.series[0].graphicalProperties.solidFill = AMBER
        ch2.series[0].graphicalProperties.line.solidFill = AMBER
        ch2.legend = None
        ch2.height, ch2.width = 4.2, 9.6
        ch2.gapWidth = 30
        axis_font(ch2.x_axis, 600)
        axis_font(ch2.y_axis, 650)
        ws.add_chart(ch2, f"G{st0 - 1}")
        # funding & loan
        r = rr + 2
        title(r, 2, "FUNDING & HOME LOAN · ገንዘብና ብድር")
        title(r, 7, "HOME-LOAN BY YEAR · ብድር በዓመት")
        fl = [("Own savings used · ቁጠባ", "=P_SAV_USED", "n0"), ("Family / diaspora · ቤተሰብ", "=P_FAM_USED", "n0"),
              ("Home loan drawn · የወጣ ብድር", "=P_LOAN_DRAWN", "n0"), ("Funding gap · ክፍተት", "=P_GAP", "n0"),
              ("Loan at completion (incl. interest added)", "=P_B0", "n0"),
              ("Interest rate · ወለድ", "=P_RATE", "pct"), ("Repayment period (years)", "=P_TENOR", "int"),
              ("Monthly instalment · ወርሃዊ ክፍያ", "=P_INSTAL", "n0"),
              ("Interest paid monthly while building (max)", "=P_IDC_PAID_MAX", "n0"),
              ("Total interest over the loan · ጠቅላላ ወለድ", "=P_INT_ALL", "n0"),
              ("Loan repaid by · የሚያልቅበት", "=P_PAYOFF", "date"),
              ("Debt payments ÷ income · ከገቢ ድርሻ", "=P_DTI", "pct"),
              ("Affordability · የመክፈል አቅም", '=P_AFF_TXT&"  ·  "&P_AFF_AM', "gen"),
              ("Largest loan your income supports", "=P_LOAN_AFF", "n0")]
        for k, (lab, f, fmt) in enumerate(fl):
            rr2 = r + 1 + k
            ws.cell(rr2, 2, lab).font = font(9)
            ws.merge_cells(start_row=rr2, start_column=3, end_row=rr2, end_column=5)
            c = ws.cell(rr2, 3, f)
            c.number_format = NF[fmt]
            c.font = font(9, True, NAVY)
            c.alignment = Alignment(horizontal="right")
        aff_r = r + 13
        gap_r = r + 4
        # local helpers for conditional formatting (Google Sheets: rules may only use cells on the same sheet)
        ws["M5"] = "helpers"
        ws["M5"].font = font(8, False, SEC)
        ws["M6"] = "=P_INCOME"
        ws["M7"] = "=P_DTI"
        ws["M8"] = "=P_AFF_MAX"
        ws["M9"] = "=P_GAP"
        for c in ("M6", "M7", "M8", "M9"):
            ws[c].font = font(8, False, "FFFFFF")
        ws.conditional_formatting.add(f"C{aff_r}:E{aff_r}", FormulaRule(formula=["AND($M$6>0,$M$7<=$M$8)"], fill=fill(GOOD_BG), font=Font(color=GOOD_FG, bold=True)))
        ws.conditional_formatting.add(f"C{aff_r}:E{aff_r}", FormulaRule(formula=["AND($M$6>0,$M$7>$M$8)"], fill=fill(CRIT_BG), font=Font(color=CRIT_FG, bold=True)))
        ws.conditional_formatting.add(f"C{gap_r}:E{gap_r}", FormulaRule(formula=["$M$9>0.5"], fill=fill(CRIT_BG), font=Font(color=CRIT_FG, bold=True)))
        hdr_row(r + 1, (7, 8, 9, 10), ("Year", "Interest", "Principal", "Closing"))
        y0 = b.P_Y0
        for y in range(1, 21):
            rr2 = r + 1 + y
            for col, src in ((7, "B"), (8, "C"), (9, "D"), (10, "F")):
                c = ws.cell(rr2, col, f"='P2_Calc'!{src}{y0 + y}")
                c.number_format = NF["int"] if col == 7 else NF["n0"]
                c.font = font(8)
                if col == 7:
                    c.alignment = Alignment(horizontal="left")
        # rent-out
        ru = r + len(fl) + 2
        title(ru, 2, "OPTIONAL: RENT OUT PART · ከፊሉን ማከራየት")
        rows = [("Switched on? (P1 block 7)", '=IF(P_RENT_YN=1,"Yes · አዎ","No · አይ")', "gen"),
                ("Rent, first year (gross) · ኪራይ", "=P_RENT_Y1", "n0"),
                ("Rent after vacancy, costs & tax", "=P_RENT_NET", "n0"),
                ("Gross yield on the rented part's cost · ገቢ ከወጪ", "=P_RENT_YIELD", "pct"),
                ("Simple payback (years) · የመመለሻ ጊዜ", "=P_RENT_PB", "n1"),
                ("Net rent ÷ monthly instalment · የሚሸፍነው", "=P_RENT_COVER", "pct")]
        for k, (lab, f, fmt) in enumerate(rows):
            rr2 = ru + 1 + k
            ws.cell(rr2, 2, lab).font = font(9)
            ws.merge_cells(start_row=rr2, start_column=3, end_row=rr2, end_column=5)
            c = ws.cell(rr2, 3, f)
            c.number_format = NF[fmt]
            c.font = font(9, True, NAVY)
            c.alignment = Alignment(horizontal="right")
        last = max(ru + len(rows) + 1, r + 23) + 1
        ws.cell(last, 2, "Planning estimate only – get a priced BOQ from a licensed quantity surveyor before signing a contract.").font = F_NOTE
        ws.cell(last + 1, 2, "ይህ ግምት ብቻ ነው፤ ውል ከመፈረምዎ በፊት ፈቃድ ካለው ባለሙያ የተዘጋጀ የሥራ መጠን ዝርዝር (BOQ) ያግኙ።").font = font(9, False, SEC, italic=True, name=AMFONT)
        ws.print_area = f"A1:K{last + 2}"
        ws.page_setup.orientation = "portrait"
        ws.page_setup.fitToHeight = 1
        ws.page_setup.fitToWidth = 1
        ws.freeze_panes = None
    sh.raw(_e)
    sh.r = 100
    return sh


# ====================================================================== P4 CHECKS
def build_p_checks(b):
    ck = Sheet(b, "P4_Checks", "P4 · My house – checks", "የቤት ስሌት ማረጋገጫዎች", "B42318", kind="free", n=0, width_label=60)
    ws = ck.ws
    ck.sec("Private-track integrity checks – errors must PASS; warnings are advice", "ማረጋገጫዎች")
    hdr = ck.r
    ck.r += 1

    def scan():
        parts = []
        for s in ("P1_Inputs", "P2_Calc", "P3_Estimate"):
            shx = b.sheets[s]
            lastc = CL(max(14, 4 + (shx.n or 0)))
            parts.append(f"SUMPRODUCT(--ISERROR('{s}'!$A$1:${lastc}${shx.r + 5}))")
        return "+".join(parts)

    checks = [
        ("P01", "Money in = money out every month (savings + family + loan + gap = cash needed)", "P_SRC_DIFF", 1, "E"),
        ("P02", "Cost breakdown adds up to the total cost", "ABS(P_CAT_TOT-P_TOTAL)+ABS(P_EST_TOTAL-P_TOTAL)", 1, "E"),
        ("P03", "Monthly costs + interest = total cost",
         lambda: f"ABS(SUM('P2_Calc'!$E${b.sheets['P2_Calc'].rows['COST']}:${pc(NM)}${b.sheets['P2_Calc'].rows['COST']})+P_IDC_TOT-P_TOTAL)", 1, "E"),
        ("P04", "Contractor stages = building work + contingency", "ABS(P_STAGE_TOT-P_BUILD_ESC-P_CONT_TOT)", 1, "E"),
        ("P05", "Loan fully repaid at the end of the term", "ABS(P_BAL_END)", 1, "E"),
        ("P06", "Principal repaid = loan balance at completion", "ABS(P_PRIN_TOT-P_B0)", 1, "E"),
        ("P07", "Switches valid (track and data set)", "IF(AND(TRACK>=1,TRACK<=2,P_DATA>=1,P_DATA<=2),0,1)", 0.5, "E"),
        ("P08", "No formula errors on P1–P3", scan, 0.5, "E"),
        ("P09", "Building type matches the floors entered", "IF(P_GFA=0,0,1-P_BTYPE_OK)", 0.5, "W"),
        ("P10", "Funding gap = 0 (savings, family and loan cover the cost)", "P_GAP", 1, "W"),
        ("P11", "Instalment within the affordability limit", "IF(AND(P_INCOME>0,P_DTI>P_AFF_MAX),1,0)", 0.5, "W"),
        ("P12", "Rented area not larger than the house", "IF(AND(P_RENT_YN=1,P_RENT_AREA>P_GFA),1,0)", 0.5, "W"),
    ]

    def _ck(_r):
        for j, h in enumerate(["", "Check", "Type", "Value", "Tolerance", "Status"]):
            if h:
                ws.cell(hdr, 1 + j, h).font = font(9, True, NAVY)
        for k, (cid, lab, f, tol, typ) in enumerate(checks):
            rr = hdr + 1 + k
            ws.cell(rr, 2, f"{cid}  {lab}").font = F_BODY
            ws.cell(rr, 3, "Error" if typ == "E" else "Warning").font = F_UNIT
            fv = f() if callable(f) else f
            ws.cell(rr, 4, "=" + fv).number_format = NF["n2"]
            ws.cell(rr, 5, tol).number_format = NF["n2"]
            st = ws.cell(rr, 6, f'=IF(ABS(D{rr})<=E{rr},"PASS",IF(C{rr}="Warning","WARN","FAIL"))')
            st.font = F_BOLD
            st.alignment = Alignment(horizontal="center")
        first, last = hdr + 1, hdr + len(checks)
        rng = f"F{first}:F{last}"
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"PASS"'], fill=fill(GOOD_BG), font=Font(color=GOOD_FG, bold=True)))
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"WARN"'], fill=fill(WARN_BG), font=Font(color=WARN_FG, bold=True)))
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"FAIL"'], fill=fill(CRIT_BG), font=Font(color=CRIT_FG, bold=True)))
        s = last + 2
        for j, (lab, f, nm) in enumerate((("Checks passed", f'COUNTIF({rng},"PASS")', "PCHK_PASS"),
                                          ("Errors (FAIL)", f'COUNTIF({rng},"FAIL")', "PCHK_FAIL"),
                                          ("Warnings", f'COUNTIF({rng},"WARN")', "PCHK_WARN"),
                                          ("Total checks", f"ROWS({rng})", "PCHK_TOTAL"))):
            ws.cell(s + j, 2, lab).font = F_BOLD
            c = ws.cell(s + j, 4, "=" + f)
            c.font = F_BOLD
            c.fill = FILL_KEY
            b.name(nm, f"'P4_Checks'!$D${s + j}")
        ws.cell(s + 5, 2, "Status").font = F_BOLD
        c = ws.cell(s + 5, 4, '=IF(PCHK_FAIL>0,"ERRORS – see list",IF(PCHK_WARN>0,"OK with warnings","ALL CHECKS PASS"))')
        c.font = F_BOLD
        b.name("PCHK_STATUS", f"'P4_Checks'!$D${s + 5}")
        ws.cell(s + 7, 2, "Fixing warnings: P09 – choose the building type that matches the floors in P1 block 2; P10 – add savings, "
                          "family money or loan; P11 – lower the cost, borrow less or extend the term.").font = F_NOTE
        for col, w in (("C", 10), ("D", 16), ("E", 11), ("F", 12)):
            ws.column_dimensions[col].width = w
    ck.raw(_ck)
    ck.r += len(checks) + 12
    return ck
